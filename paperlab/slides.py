"""발표 자료 만들기 (2026-10-27 사용자 요청): AI가 원고로 만든 슬라이드 JSON → PowerPoint(.pptx).

JSON: {"title": 발표 제목, "subtitle"?: 이름 · 소속, "slides": [{"title", "bullets": [..], "notes"}],
       "questions": [{"q", "a"}]} — 예상 질문은 끝 슬라이드(답변 방향은 아래 단계 글머리)로 넣는다.
글꼴은 한글 '맑은 고딕'(동아시아 글꼴 칸까지), 16:9."""
from __future__ import annotations

import io
import json
import re
from typing import Any

MAX_SLIDES = 40
MAX_BULLETS = 8
MAX_QUESTIONS = 15
FONT = "맑은 고딕"


class SlidesError(ValueError):
    pass


def _s(v: Any, n: int) -> str:
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(v or "")).strip()[:n]


def parse(data: Any) -> dict:
    """AI 답(JSON 글 — 코드 블록 · 앞뒤 글 허용) 또는 dict → 정리한 슬라이드 데이터"""
    if isinstance(data, str):
        a, b = data.find("{"), data.rfind("}")
        if a < 0 or b <= a:
            raise SlidesError("슬라이드 내용을 읽지 못했어요")
        try:
            data = json.loads(data[a:b + 1])
        except json.JSONDecodeError as e:
            raise SlidesError("슬라이드 내용을 읽지 못했어요") from e
    if not isinstance(data, dict):
        raise SlidesError("슬라이드 내용을 읽지 못했어요")
    slides = []
    for s in (data.get("slides") or [])[:MAX_SLIDES]:
        if not isinstance(s, dict):
            continue
        bullets = [_s(b, 300) for b in (s.get("bullets") or []) if _s(b, 300)][:MAX_BULLETS]
        title = _s(s.get("title"), 120)
        if title or bullets:
            slides.append({"title": title, "bullets": bullets, "notes": _s(s.get("notes"), 2000)})
    questions = [{"q": _s(q.get("q"), 300), "a": _s(q.get("a"), 500)} for q in (data.get("questions") or [])[:MAX_QUESTIONS]
                 if isinstance(q, dict) and _s(q.get("q"), 300)]
    if not slides:
        raise SlidesError("슬라이드가 하나도 없어요")
    return {"title": _s(data.get("title"), 200) or slides[0]["title"], "subtitle": _s(data.get("subtitle"), 200),
            "slides": slides, "questions": questions}


def _font(run, size=None, bold=None):
    from pptx.oxml.ns import qn
    from pptx.util import Pt

    run.font.name = FONT
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def _fill(frame, items: list[tuple[str, int]], size: int) -> None:
    """items = [(글, 단계)]"""
    frame.clear()
    frame.word_wrap = True
    for k, (text, level) in enumerate(items):
        p = frame.paragraphs[0] if k == 0 else frame.add_paragraph()
        p.level = level
        r = p.add_run()
        r.text = text
        _font(r, size - 2 * level)


def to_pptx(data: dict) -> bytes:
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches

    d = parse(data)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    title_layout, body_layout = prs.slide_layouts[0], prs.slide_layouts[1]

    s = prs.slides.add_slide(title_layout)
    s.shapes.title.text = d["title"]
    s.shapes.title.left, s.shapes.title.top, s.shapes.title.width, s.shapes.title.height = Inches(1), Inches(2.3), Inches(11.3), Inches(1.6)
    _font(s.shapes.title.text_frame.paragraphs[0].runs[0], 36, True)
    sub = s.placeholders[1]
    sub.left, sub.top, sub.width, sub.height = Inches(1), Inches(4.1), Inches(11.3), Inches(1.0)
    if d["subtitle"]:
        _fill(sub.text_frame, [(d["subtitle"], 0)], 20)
    else:
        sub._element.getparent().remove(sub._element)

    def add(title: str, items: list[tuple[str, int]], notes: str = "") -> None:
        sl = prs.slides.add_slide(body_layout)
        sl.shapes.title.text = title or " "
        _font(sl.shapes.title.text_frame.paragraphs[0].runs[0], 30, True)
        body = sl.placeholders[1]
        body.left, body.top, body.width, body.height = Inches(0.7), Inches(1.6), Inches(11.9), Inches(5.4)
        t = sl.shapes.title
        t.left, t.top, t.width, t.height = Inches(0.7), Inches(0.4), Inches(11.9), Inches(1.0)
        t.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
        if items:
            _fill(body.text_frame, items, 22)
        else:
            body._element.getparent().remove(body._element)
        if notes:
            sl.notes_slide.notes_text_frame.text = notes

    for sl in d["slides"]:
        add(sl["title"], [(b, 0) for b in sl["bullets"]], sl["notes"])
    qs = d["questions"]
    for i in range(0, len(qs), 3):  # 예상 질문: 슬라이드마다 3개 (답변 방향은 아래 단계)
        part = qs[i:i + 3]
        items = []
        for q in part:
            items.append((f"Q. {q['q']}", 0))
            if q["a"]:
                items.append((q["a"], 1))
        n = (len(qs) + 2) // 3
        add(f"예상 질문{f' ({i // 3 + 1}/{n})' if n > 1 else ''}", items, "심사 대비용 — 발표에서는 빼도 돼요")
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
