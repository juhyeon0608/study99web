"""원고를 워드(.docx)·한글(.hwpx) 파일로 쓴다.

화면이 인용·참고문헌을 citeproc-js로 서식화한 뒤 아래 '블록' 목록을 보낸다.

블록: {"type": "title" | "heading" | "paragraph" | "list_item" | "quote" | "code" |
         "bib_heading" | "bib_entry" | "page_break" | "figure" | "table",
       "level": 1~4 (heading), "ordered": bool, "number": int, "depth": int (list_item),
       "hanging": bool (bib_entry), "text": str (code·bib_heading), "runs": [런…]}
그림: {"type": "figure", "src": 그림 이름, "alt": str, "caption": [런…] ("그림 1. …" — 그림 아래),
       "image": {"data", "format", "width", "height"} ← 서버가 저장소에서 읽어 넣는다(figures.attach_images)}
표:   {"type": "table", "rows": [[[런…] 칸…] 줄…], "header": bool(첫 줄이 머리글),
       "align": ["left" | "center" | "right" | None …], "caption": [런…] ("표 1. …" — 표 위)}
런:   {"text": str, "b": bool, "i": bool, "sup": bool, "sub": bool, "code": bool}
      또는 {"footnote": [런…]}  ← 그 자리에 각주 번호가 들어가고 쪽 아래에 내용이 붙는다
"""

from __future__ import annotations

import copy
import io
import re
from datetime import datetime, timezone
from typing import Any, Iterable

from . import doc_formats as df
from . import doc_lists
from . import figures as fig

Block = dict[str, Any]
Run = dict[str, Any]

BODY_FONT_LATIN = "Times New Roman"
BODY_FONT_KO = "바탕"
HEADING_SIZES = {1: 14, 2: 12, 3: 11, 4: 11}


def _clean_text(text: str) -> str:
    # XML에 쓸 수 없는 제어 문자를 지운다
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text or "")


_BULLETS = ("•", "◦", "▪")  # 글머리 기호: 단계(depth)마다 바꾼다


def _list_prefix(block: Block) -> str:
    if block.get("ordered"):
        return f"{block.get('number') or 1}. "
    return _BULLETS[max(0, int(block.get("depth") or 0)) % len(_BULLETS)] + " "


def plain_text(runs: Iterable[Run]) -> str:
    return "".join(r.get("text", "") for r in runs if "footnote" not in r)


# 앞부분 제목: 쪽 번호를 "첫 장부터 1쪽"으로 셀 때 장으로 보지 않는다 (띄어쓰기 · 대소문자 무시)
FRONT_MATTER_TITLES = frozenset({"국문초록", "초록", "국문요약", "요약", "요지", "abstract", "감사의글", "감사의말씀", "사사",
                                 "머리말", "서문", "preface", "acknowledgements", "acknowledgments", "목차", "차례"})


def is_front_matter(block: dict) -> bool:
    text = plain_text(block.get("runs") or []) if block.get("runs") else str(block.get("text") or "")
    return re.sub(r"[\s<>〈〉()\[\]·.:]+", "", text).lower() in FRONT_MATTER_TITLES


# ------------------------------------------------------------------ 그림 · 표 공통
FIGURE_MISSING = "[그림을 찾을 수 없어요]"
MAX_TABLE_ROWS, MAX_TABLE_COLS = 500, 30


def _runs_of(value) -> list[Run]:
    """런 목록만 (각주는 표 칸 · 캡션에 넣지 않고 글자로 둔다)"""
    if not isinstance(value, list):
        return []
    out = []
    for r in value:
        if isinstance(r, dict) and "footnote" in r:
            out.append({"text": plain_text(r["footnote"] if isinstance(r["footnote"], list) else [])})
        elif isinstance(r, dict):
            out.append(r)
    return out


def _table_rows(block: Block) -> list[list[list[Run]]]:
    """표 칸 [[런…]…] — 줄마다 칸 수를 가장 긴 줄에 맞춘다. 비었으면 1×1"""
    rows = [r for r in (block.get("rows") or []) if isinstance(r, list)][:MAX_TABLE_ROWS]
    cols = min(MAX_TABLE_COLS, max((len(r) for r in rows), default=0)) or 1
    return [[_runs_of(r[c]) if c < len(r) else [] for c in range(cols)] for r in rows] or [[[]]]


def _table_align(block: Block, col: int) -> str:
    align = block.get("align") or []
    a = align[col] if isinstance(align, list) and col < len(align) else None
    return a if a in ("left", "center", "right") else "left"


def _cell_align(block: Block, row: int, col: int) -> str:
    """칸 정렬: 머리글 줄은 늘 가운데, 나머지는 열 정렬 (2026-10-26 사용자 요청 — 내용은 왼쪽 · 머리글만 가운데)"""
    return "center" if row == 0 and block.get("header", True) else _table_align(block, col)


def _hwpx_text_run(p, text: str, **kw):
    """HWPX 글 런 — 줄바꿈(표 칸 안 Shift+Enter)은 <hp:lineBreak/>로 (글자 줄바꿈은 한글이 무시함)"""
    parts = text.split("\n")
    run = p.add_run(parts[0], **kw)
    if len(parts) > 1:
        t = run.element.find(f"{_HP}t")
        if t is None:
            t = run.element.makeelement(f"{_HP}t", {})
            run.element.append(t)
        for part in parts[1:]:
            br = t.makeelement(f"{_HP}lineBreak", {})
            br.tail = part
            t.append(br)
    return run


def _cell_runs(block: Block, row: int, runs: list[Run]) -> list[Run]:
    return [dict(r, b=True) for r in runs] if row == 0 and block.get("header", True) else runs


# =================================================================== DOCX
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
FOOTNOTES_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
FOOTNOTES_RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"


class _DocxFootnotes:
    """python-docx에 없는 각주를 OOXML 수준에서 추가한다."""

    def __init__(self, document):
        from docx.opc.packuri import PackURI
        from docx.opc.part import XmlPart
        from docx.oxml import parse_xml

        xml = (
            f'<w:footnotes xmlns:w="{W_NS}">'
            '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" '
            'w:lineRule="auto"/></w:pPr><w:r><w:separator/></w:r></w:p></w:footnote>'
            '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" '
            'w:line="240" w:lineRule="auto"/></w:pPr><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
            "</w:footnotes>"
        )
        self.part = XmlPart(PackURI("/word/footnotes.xml"), FOOTNOTES_CT, parse_xml(xml), document.part.package)
        document.part.relate_to(self.part, FOOTNOTES_RT)
        self.next_id = 1

    def add(self, paragraph, runs: list[Run]) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        nid = self.next_id
        self.next_id += 1
        # 본문 쪽 각주 번호
        ref_run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        va = OxmlElement("w:vertAlign")
        va.set(qn("w:val"), "superscript")
        rpr.append(va)
        ref_run.append(rpr)
        ref = OxmlElement("w:footnoteReference")
        ref.set(qn("w:id"), str(nid))
        ref_run.append(ref)
        paragraph._p.append(ref_run)
        # 쪽 아래 각주 내용
        note = OxmlElement("w:footnote")
        note.set(qn("w:id"), str(nid))
        p = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "0")
        spacing.set(qn("w:line"), "240")
        spacing.set(qn("w:lineRule"), "auto")
        ppr.append(spacing)
        p.append(ppr)
        mark = OxmlElement("w:r")
        mark.append(copy.deepcopy(rpr))
        mark.append(OxmlElement("w:footnoteRef"))
        p.append(mark)
        for r in [{"text": " "}] + list(runs):
            if "footnote" in r:
                continue
            p.append(_docx_run_element(r, size_half_points=18))
        note.append(p)
        self.part.element.append(note)


def _docx_run_element(r: Run, size_half_points: int | None = None):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    if r.get("b"):
        rpr.append(OxmlElement("w:b"))
    if r.get("i"):
        rpr.append(OxmlElement("w:i"))
    if r.get("sup") or r.get("sub"):
        va = OxmlElement("w:vertAlign")
        va.set(qn("w:val"), "superscript" if r.get("sup") else "subscript")
        rpr.append(va)
    if size_half_points:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(size_half_points))
        rpr.append(sz)
    if len(rpr):
        run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = _clean_text(r.get("text", ""))
    run.append(t)
    return run


def _docx_set_fonts(style, latin: str, east_asia: str, size_pt: float | None = None, bold=None, color_black=False):
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    style.font.name = latin
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    if size_pt:
        style.font.size = Pt(size_pt)
    if bold is not None:
        style.font.bold = bold
    if color_black:
        style.font.color.rgb = RGBColor(0, 0, 0)


FIGURE_UNREADABLE_WARNING = "figure-unreadable"


def _docx_picture(paragraph, image: dict, max_w_mm: float, max_h_mm: float) -> bool:
    """그림을 넣으면 True. python-docx가 못 읽는 그림(JFIF · Exif 머리글 없는 JPEG 등)이면 아무것도 넣지 않고 False
    — 부른 쪽이 '[그림을 찾을 수 없어요]'로 적고 경고를 남긴다(내보내기 전체가 500으로 실패하지 않게)"""
    from docx.shared import Mm

    w, h = fig.fit_mm(image["width"], image["height"], max_w_mm, max_h_mm)
    run = paragraph.add_run()
    try:
        run.add_picture(io.BytesIO(image["data"]), width=Mm(w), height=Mm(h))
    except Exception:  # noqa: BLE001 - docx.image.exceptions.UnrecognizedImageError 등 (ValueError가 아님)
        run._r.getparent().remove(run._r)
        return False
    return True


def _figure_warn(warnings: list[str] | None) -> None:
    if warnings is not None and FIGURE_UNREADABLE_WARNING not in warnings:
        warnings.append(FIGURE_UNREADABLE_WARNING)


def _docx_single_spacing(paragraph) -> None:
    """그림 문단은 줄간격을 '한 줄'로 (고정 줄간격이면 그림이 잘려 보인다)"""
    from docx.oxml.ns import qn

    sp = paragraph._p.get_or_add_pPr().get_or_add_spacing()
    sp.set(qn("w:line"), "240")
    sp.set(qn("w:lineRule"), "auto")


def _docx_table(document, block: Block, fill):
    """표 (Table Grid 테두리 · 가운데 정렬). fill(문단, 런, 정렬)로 칸을 채운다. 머리글 줄은 쪽마다 되풀이"""
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement

    rows = _table_rows(block)
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    try:
        table.style = document.styles["Table Grid"]
    except KeyError:
        pass
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        cells = table.rows[r].cells
        for c, runs in enumerate(row):
            fill(cells[c].paragraphs[0], _cell_runs(block, r, runs), _cell_align(block, r, c))
    if block.get("header", True) and len(rows) > 1:
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    return table


def to_docx(blocks: list[Block], meta: dict | None = None, doc_format: dict | None = None,
            cover: dict | None = None, warnings: list[str] | None = None) -> bytes:
    """doc_format(검증된 양식 데이터)이 없으면 지금 서식 그대로(기본 A4) 만든다."""
    if doc_format is not None:
        return _render_formatted(_DocxBuilder(doc_format, meta or {}), blocks, doc_format, meta or {}, cover, warnings)
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Mm, Pt

    meta = meta or {}
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)  # A4
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Mm(30))
    section.top_margin = section.bottom_margin = Mm(25)

    styles = document.styles
    normal = styles["Normal"]
    _docx_set_fonts(normal, BODY_FONT_LATIN, BODY_FONT_KO, 11)
    normal.paragraph_format.line_spacing = 1.6
    normal.paragraph_format.space_after = Pt(6)
    for level in range(1, 5):
        st = styles[f"Heading {level}"]
        _docx_set_fonts(st, BODY_FONT_LATIN, BODY_FONT_KO, HEADING_SIZES[level], bold=True, color_black=True)
        st.font.italic = False
        st.paragraph_format.space_before = Pt(14 if level == 1 else 10)
        st.paragraph_format.space_after = Pt(6)
    document.core_properties.title = meta.get("title") or ""
    document.core_properties.author = meta.get("author") or ""

    notes: _DocxFootnotes | None = None

    def add_runs(paragraph, runs: list[Run]):
        nonlocal notes
        for r in runs:
            if "footnote" in r:
                if notes is None:
                    notes = _DocxFootnotes(document)
                notes.add(paragraph, r["footnote"])
                continue
            run = paragraph.add_run(_clean_text(r.get("text", "")))
            run.bold = bool(r.get("b")) or None
            run.italic = bool(r.get("i")) or None
            if r.get("sup"):
                run.font.superscript = True
            if r.get("sub"):
                run.font.subscript = True
            if r.get("code"):
                run.font.name = "Consolas"

    for block in blocks:
        kind = block.get("type")
        runs = block.get("runs") or []
        if kind == "title":
            # 워드 기본 '제목' 스타일은 파란 밑줄이 있어 논문에 맞지 않으므로 직접 꾸민다
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(18)
            add_runs(p, [dict(r, b=True) if "footnote" not in r else r for r in runs])
            for run in p.runs:
                run.font.size = Pt(18)
        elif kind == "heading":
            level = max(1, min(4, int(block.get("level") or 1)))
            add_runs(document.add_paragraph(style=f"Heading {level}"), runs)
        elif kind == "list_item":
            p = document.add_paragraph()
            depth = int(block.get("depth") or 0)
            p.paragraph_format.left_indent = Mm(8 + 6 * depth)
            p.paragraph_format.first_line_indent = Mm(-5)
            add_runs(p, [{"text": _list_prefix(block)}] + runs)
        elif kind == "quote":
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Mm(10)
            p.paragraph_format.right_indent = Mm(10)
            add_runs(p, runs)
        elif kind == "code":
            p = document.add_paragraph()
            run = p.add_run(_clean_text(block.get("text", "")))
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif kind == "bib_heading":
            document.add_paragraph(block.get("text") or "참고문헌", style="Heading 1")
        elif kind == "bib_entry":
            p = document.add_paragraph()
            if block.get("hanging", True):
                p.paragraph_format.left_indent = Mm(10)
                p.paragraph_format.first_line_indent = Mm(-10)
            add_runs(p, runs)
        elif kind == "page_break":
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "figure":
            # 그림(가운데) → 아래에 "그림 n. 캡션"
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            image = block.get("image")
            if image and _docx_picture(p, image, 150, 200):  # A4 본문 폭 150mm
                _docx_single_spacing(p)
            else:
                if image:
                    _figure_warn(warnings)
                add_runs(p, [{"text": FIGURE_MISSING}])
            caption = _runs_of(block.get("caption"))
            if caption:
                cp = document.add_paragraph()
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_runs(cp, caption)
        elif kind == "table":
            # "표 n. 캡션"(표 위, 다음과 함께) → 표
            caption = _runs_of(block.get("caption"))
            if caption:
                cp = document.add_paragraph()
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cp.paragraph_format.keep_with_next = True
                add_runs(cp, caption)

            def fill(p, cell_runs, align):
                p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}.get(align)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                add_runs(p, cell_runs)

            _docx_table(document, block, fill)
        else:
            p = document.add_paragraph()
            p.paragraph_format.first_line_indent = Mm(meta.get("first_line_indent_mm", 0) or 0) or None
            add_runs(p, runs)

    body_el = document.element.body
    if len(body_el) > 1 and body_el[-2].tag.endswith("}tbl"):
        document.add_paragraph()  # 표로 끝나는 문서는 뒤에 빈 문단(워드 규칙)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# =================================================================== HWPX
class _HwpxWriter:
    def __init__(self):
        from hwpx import HwpxDocument

        self.doc = HwpxDocument.new()
        self._char_cache: dict[tuple, str] = {}
        self._first_paragraph = self.doc.paragraphs[0] if self.doc.paragraphs else None

    def char(self, b=False, i=False, script=None, size=None) -> str | None:
        key = (bool(b), bool(i), script, size)
        if key == (False, False, None, None):
            return None
        if key not in self._char_cache:
            self._char_cache[key] = self.doc.oxml.ensure_run_style(
                bold=bool(b), italic=bool(i), script=script, size=size)
        return self._char_cache[key]

    def paragraph(self):
        # 빈 문서의 첫 문단(구역 설정을 담고 있음)을 먼저 쓴다
        if self._first_paragraph is not None:
            p, self._first_paragraph = self._first_paragraph, None
            return p
        return self.doc.add_paragraph("", include_run=False, inherit_style=False)

    def runs(self, paragraph, runs: list[Run], size=None, bold=False):
        for r in runs:
            if "footnote" in r:
                text = plain_text(r["footnote"]).strip()
                paragraph.add_footnote(_clean_text(text) or " ")
                continue
            text = _clean_text(r.get("text", ""))
            if not text:
                continue
            script = "sup" if r.get("sup") else ("sub" if r.get("sub") else None)
            cid = self.char(b=bold or r.get("b"), i=r.get("i"), script=script, size=size)
            if cid is None:
                _hwpx_text_run(paragraph, text)
            else:
                _hwpx_text_run(paragraph, text, char_pr_id_ref=cid)

    def fmt(self, paragraph, **kw):
        from hwpx._document import layout

        layout.set_paragraph_format(self.doc, paragraphs=[paragraph], **kw)

    def picture(self, paragraph, image: dict, max_w_mm: float, max_h_mm: float) -> None:
        _hwpx_picture(self.doc, paragraph, image, max_w_mm, max_h_mm)

    def table(self, block: Block, width_mm: float) -> None:
        rows = _table_rows(block)
        t = _hwpx_table(self.paragraph(), block, rows, width_mm, 7.0)
        for r, row in enumerate(rows):
            for c, runs in enumerate(row):
                p = t.cell(r, c).paragraphs[0]
                self.runs(p, _cell_runs(block, r, runs))
                self.fmt(p, alignment=_HWPX_ALIGN[_cell_align(block, r, c)], first_line_indent_mm=0,
                         spacing_after_pt=0)


def _hwpx_table(anchor, block: Block, rows: list, width_mm: float, row_mm: float, char_pr_id_ref=None):
    """문단에 표를 놓는다. 줄 높이는 최소값(글이 길면 한글이 늘림) — 0에 가까우면 다른 프로그램에서 납작해진다.
    머리글 줄은 쪽마다 되풀이"""
    kw = {"char_pr_id_ref": char_pr_id_ref} if char_pr_id_ref is not None else {}
    t = anchor.add_table(len(rows), len(rows[0]), width=df.mm_to_hwp(width_mm),
                         height=df.mm_to_hwp(row_mm * len(rows)), **kw)
    header = bool(block.get("header", True) and len(rows) > 1)
    if header:
        t.element.set("repeatHeader", "1")
    for tc in t.element.iter(f"{_HP}tc"):
        tc.set("hasMargin", "1")  # 칸 안 여백(cellMargin)을 쓰게 — 안 쓰면 글이 테두리에 붙는 프로그램이 있음
        addr = tc.find(f"{_HP}cellAddr")
        if header and addr is not None and addr.get("rowAddr") == "0":
            tc.set("header", "1")
    t.mark_dirty()
    return t


def _hwpx_picture(doc, paragraph, image: dict, max_w_mm: float, max_h_mm: float, char_pr_id_ref=None) -> None:
    """그림을 문서에 넣고(BinData) 문단에 글자처럼 취급하는 그림으로 놓는다"""
    w, h = fig.fit_mm(image["width"], image["height"], max_w_mm, max_h_mm)
    # 같은 그림(attach_images가 같은 바이트를 돌려줌)을 여러 번 넣어도 BinData에는 한 번만
    added = doc.__dict__.setdefault("_paperlab_media", {})
    item_id = added.get(id(image["data"]))
    if item_id is None:
        item_id = added[id(image["data"])] = doc.media.add_image(image["data"], image["format"]).item_id
    kw = {"char_pr_id_ref": char_pr_id_ref} if char_pr_id_ref is not None else {}
    paragraph.add_picture(item_id, width=df.mm_to_hwp(w), height=df.mm_to_hwp(h), **kw)


def to_hwpx(blocks: list[Block], meta: dict | None = None, doc_format: dict | None = None,
            cover: dict | None = None, warnings: list[str] | None = None) -> bytes:
    """doc_format(검증된 양식 데이터)이 없으면 지금 서식 그대로(한글 기본 서식) 만든다."""
    if doc_format is not None:
        return _render_formatted(_HwpxBuilder(doc_format), blocks, doc_format, meta or {}, cover, warnings)
    meta = meta or {}
    w = _HwpxWriter()
    for block in blocks:
        kind = block.get("type")
        runs = block.get("runs") or []
        if kind == "title":
            p = w.paragraph()
            w.runs(p, runs, size=16, bold=True)
            w.fmt(p, alignment="CENTER", spacing_after_pt=12)
        elif kind == "heading":
            level = max(1, min(4, int(block.get("level") or 1)))
            p = w.paragraph()
            w.runs(p, runs, size=HEADING_SIZES[level], bold=True)
            w.fmt(p, spacing_before_pt=12 if level == 1 else 8, spacing_after_pt=4, keep_with_next=True)
        elif kind == "list_item":
            p = w.paragraph()
            w.runs(p, [{"text": _list_prefix(block)}] + runs)
            depth = int(block.get("depth") or 0)
            w.fmt(p, indent_left_mm=8 + 6 * depth, first_line_indent_mm=-5)
        elif kind == "quote":
            p = w.paragraph()
            w.runs(p, runs)
            w.fmt(p, indent_left_mm=10, indent_right_mm=10)
        elif kind == "code":
            p = w.paragraph()
            w.runs(p, [{"text": block.get("text", "")}])
        elif kind == "bib_heading":
            p = w.paragraph()
            w.runs(p, [{"text": block.get("text") or "참고문헌"}], size=HEADING_SIZES[1], bold=True)
            w.fmt(p, spacing_before_pt=12, spacing_after_pt=4, keep_with_next=True)
        elif kind == "bib_entry":
            p = w.paragraph()
            w.runs(p, runs)
            if block.get("hanging", True):
                w.fmt(p, indent_left_mm=10, first_line_indent_mm=-10)
        elif kind == "page_break":
            p = w.paragraph()
            w.fmt(p, page_break_before=True)
        elif kind == "figure":
            p = w.paragraph()
            image = block.get("image")
            if image:
                w.picture(p, image, 150, 200)  # 한글 기본 쪽 본문 폭 150mm
            else:
                w.runs(p, [{"text": FIGURE_MISSING}])
            w.fmt(p, alignment="CENTER", keep_with_next=True, line_spacing_percent=100, spacing_before_pt=6)
            caption = _runs_of(block.get("caption"))
            if caption:
                p = w.paragraph()
                w.runs(p, caption)
                w.fmt(p, alignment="CENTER", spacing_after_pt=6)
        elif kind == "table":
            caption = _runs_of(block.get("caption"))
            if caption:
                p = w.paragraph()
                w.runs(p, caption)
                w.fmt(p, alignment="CENTER", keep_with_next=True, spacing_before_pt=6)
            w.table(block, 150)
        else:
            p = w.paragraph()
            w.runs(p, runs)
    report = w.doc.validate()
    if getattr(report, "issues", ()):
        raise ValueError(f"한글 문서 검증 실패: {report.issues[:3]}")
    return w.doc.to_bytes()


def to_markdown(blocks: list[Block]) -> str:
    out, notes = [], []

    def md(runs):
        s = ""
        for r in runs:
            if "footnote" in r:
                notes.append(md(r["footnote"]))
                s += f"[^{len(notes)}]"
                continue
            t = r.get("text", "")
            if r.get("b"):
                t = f"**{t}**"
            if r.get("i"):
                t = f"*{t}*"
            if r.get("sup"):
                t = f"<sup>{t}</sup>"
            s += t
        return s

    list_cols: list[int] = []  # 단계별 바로 앞 항목의 내용 시작 칸 (하위 항목은 부모 기호 폭만큼 들여써야 하위 목록이 됨)
    for b in blocks:
        kind = b.get("type")
        if kind != "list_item":
            list_cols = []
        if kind == "title":
            out.append("# " + md(b.get("runs") or []))
        elif kind == "heading":
            out.append("#" * (int(b.get("level") or 1) + 1) + " " + md(b.get("runs") or []))
        elif kind == "list_item":
            marker = _list_prefix(b) if b.get("ordered") else "- "
            depth = max(0, int(b.get("depth") or 0))
            # "10. 항목"의 하위는 4칸, "- 항목"의 하위는 2칸 (CommonMark: 부모 내용 칸 이상 들여써야 하위)
            indent = (list_cols[depth - 1] if depth <= len(list_cols) else list_cols[-1]) if depth and list_cols else 0
            list_cols = list_cols[:depth] + [indent] * (depth - len(list_cols)) + [indent + len(marker)]
            out.append(" " * indent + marker + md(b.get("runs") or []))
        elif kind == "figure":
            src = b.get("src") if isinstance(b.get("src"), str) and fig.NAME_RE.fullmatch(b["src"]) else ""
            alt = str(b.get("alt") or "").replace("]", "\\]")
            out.append((f"![{alt}](fig:{src})\n\n" if src else "") + md(_runs_of(b.get("caption"))))
        elif kind == "table":
            rows = _table_rows(b)
            cell = lambda runs: md(runs).replace("|", "\\|").replace("\n", "<br>")  # noqa: E731 — 칸 안 줄바꿈
            sep = {"left": ":--", "center": ":-:", "right": "--:"}
            lines = ["| " + " | ".join(cell(c) for c in rows[0]) + " |",
                     "| " + " | ".join(sep[_table_align(b, i)] for i in range(len(rows[0]))) + " |"]
            lines += ["| " + " | ".join(cell(c) for c in r) + " |" for r in rows[1:]]
            caption = md(_runs_of(b.get("caption")))
            out.append((caption + "\n\n" if caption else "") + "\n".join(lines))
        elif kind == "quote":
            out.append("> " + md(b.get("runs") or []))
        elif kind == "code":
            out.append("```\n" + b.get("text", "") + "\n```")
        elif kind == "bib_heading":
            out.append("## " + (b.get("text") or "참고문헌"))
        elif kind == "page_break":
            out.append("\n---\n")
        else:
            out.append(md(b.get("runs") or []))
    text = "\n\n".join(out)
    if notes:
        text += "\n\n" + "\n".join(f"[^{i}]: {n}" for i, n in enumerate(notes, 1))
    return text + "\n"


# =================================================================== 양식 적용 (docs/specs/doc-formats.md)
# 기본 (A4)·양식 없음은 위의 지금 코드 경로를 그대로 쓰고, 그 밖의 양식만 아래 경로로 만든다.
# 순서: [표지 → 속표지 → 인정서/인준서] → [첫 장 앞 부분(쪽 번호 없음)] → [본문(1쪽부터)]
_HP = "{http://www.hancom.co.kr/hwpml/2011/paragraph}"
_HH = "{http://www.hancom.co.kr/hwpml/2011/head}"
_HWPX_ALIGN = {"left": "LEFT", "center": "CENTER", "right": "RIGHT", "justify": "JUSTIFY"}
_HWPX_LANGS = ("HANGUL", "LATIN", "HANJA", "JAPANESE", "OTHER", "SYMBOL", "USER")
# 쪽 번호 위치 → (머리말/꼬리말, 정렬, 한글 위치 이름)
_PAGE_NUMBER_PLACES = {"footer-center": ("footer", "center", "BOTTOM_CENTER"),
                       "footer-right": ("footer", "right", "BOTTOM_RIGHT"),
                       "header-center": ("header", "center", "TOP_CENTER"),
                       "header-right": ("header", "right", "TOP_RIGHT")}
COVER_BOTTOM_MARGIN_MM = 10  # 표지류 구역의 실제 아래 여백(줄 위치는 '아래 N'으로 계산하므로 여유만 둔다)
# 한글 목록 줄의 탭 채움: (hp:tab 숫자, hh:tabItem 이름) — 한컴이 만드는 차례와 같은 3번(DASH, 촘촘한 점선으로 보임)
_HWPX_TOC_LEADER = (3, "DASH")
CODE_FONT, CODE_SIZE = "Consolas", 9.5


def _pspec(st: dict, **over) -> dict:
    """양식 글자·문단 모양(실제 적용값) → 문단 사양."""
    spec = {"size": st["size_pt"], "align": st["align"], "line_pct": st["line_spacing_pct"],
            "before_pt": st["space_before_pt"], "after_pt": st["space_after_pt"],
            "left": st["indent_left"], "right": st["indent_right"], "first": st["first_line_indent"],
            "page_break": st["page_break_before"], "keep_next": False}
    spec.update(over)
    return spec


def _rspec(st: dict, **over) -> dict:
    spec = {"fonts": st["fonts"], "size": st["size_pt"], "bold": st["bold"], "italic": st["italic"]}
    spec.update(over)
    return spec


def _neg(length: dict) -> dict:
    return {"value": -float(length["value"]), "unit": length["unit"]}


def _render_formatted(builder, blocks: list[Block], fmt: dict, meta: dict, cover: dict | None,
                      warnings: list[str] | None) -> bytes:
    warnings = warnings if warnings is not None else []
    builder.warnings = warnings  # 그림을 못 넣으면 builder가 경고를 남긴다
    page = fmt["page"]
    title_text = (meta.get("title") or "").strip() or next(
        (plain_text(b.get("runs") or []).strip() for b in blocks if b.get("type") == "title"), "")
    kind = fmt["cover"]["kind"]
    if kind != "none":
        pages = df.cover_pages(kind, cover, title_text)
        missing = df.missing_cover_fields(df.cover_with_defaults(cover), kind)
        if pages and missing:
            warnings.append("cover-missing:" + ",".join(missing))
        for pg in pages:
            lay = df.layout_cover_page(pg, page["width_mm"], page["height_mm"])
            if lay["overflow"]:
                warnings.append(f"cover-overflow:{pg['key']}")
            builder.cover_page(pg, lay)

    # 목록 블록(type=contents)은 서버만 만든다 — 화면이 보낸 것은 버린다
    body = [b for b in blocks if not (b.get("type") == "title" and not fmt["title"]["show"])
            and b.get("type") != "contents"]
    # 목차 · 그림 목차 · 표 목차 (doc_lists): 표지류 바로 뒤, 첫 장 앞 부분보다 앞 (표지 없는 양식은 빈 목록)
    lists = doc_lists.build(body, fmt, cover) if kind != "none" else []
    pn = fmt["page_number"]
    parts = [(lists + body, pn["show"])]
    if pn["show"] and pn["start"] == "first_chapter":
        # 첫 장 = 앞부분 제목(국문 초록 · ABSTRACT · 감사의 글 · 머리말 등)이 아닌 첫 '##' 제목 (2026-10-17)
        idx = next((i for i, b in enumerate(body)
                    if b.get("type") == "heading" and int(b.get("level") or 1) == 1 and not is_front_matter(b)), None)
        if idx:
            before = body[:idx]
            # 첫 장 앞 부분(목차 · 국문초록 등)은 쪽 번호 없는 구역. 쪽 나눔만 있으면 버린다
            parts = ([(lists + before, False)] if lists or any(b.get("type") != "page_break" for b in before)
                     else []) + [(body[idx:], True)]
        elif idx == 0 and lists:
            parts = [(lists, False), (body, True)]
    if lists:  # 목록 항목의 쪽 번호(어림값). 쪽 번호를 안 넣는 양식이면 본문 처음부터 센다
        doc_lists.fill_pages(parts if pn["show"] else [(p, True) for p, _ in parts], fmt)
    for part, numbered in parts:
        builder.body_section(numbered)
        for block in part:
            builder.block(block)
    return builder.finish()


class _FormattedBuilder:
    """docx·hwpx 공통: 블록 → 양식 모양 고르기, 쪽 나눔 처리."""

    def __init__(self, fmt: dict):
        self.fmt = fmt
        self.styles = {k: df.resolve_style(fmt, k) for k in
                       ("body", "title", "headings.1", "headings.2", "headings.3", "quote", "footnote", "bibliography")}
        self.warnings: list[str] | None = None  # _render_formatted가 넣어 준다
        self.section_start = True   # 다음 문단이 구역 첫 문단이면 '새 쪽에서 시작'을 끈다(빈 쪽 방지)
        self.pending_break = False  # page_break 블록은 다음 문단의 '새 쪽에서 시작'으로 바꾼다

    def block(self, block: Block) -> None:
        kind = block.get("type")
        runs = block.get("runs") or []
        S = self.styles
        override_indent = False
        if kind == "page_break":
            self.pending_break = True
            return
        if kind == "contents":  # doc_lists가 넣은 목차 · 그림 목차 · 표 목차
            self.contents(block)
            return
        if kind in ("figure", "table"):
            # 첫 문단(그림 문단 · 표 캡션)이 쪽 나눔을 받고 다음 것과 함께 둔다
            spec = self.caption_spec(keep_next=True)
            spec["page_break"] = bool(self.pending_break and not self.section_start)
            self.pending_break = self.section_start = False
            if kind == "figure":
                self.figure(block, spec)
            else:
                self.table(block, spec)
            return
        if kind == "title":
            key, spec = "title", _pspec(S["title"], keep_next=True)
        elif kind == "heading":
            key = f"headings.{max(1, min(3, int(block.get('level') or 1)))}"
            spec = _pspec(S[key], keep_next=True)
        elif kind == "list_item":
            depth = int(block.get("depth") or 0)
            key, override_indent = "body", True
            spec = _pspec(S["body"], left=df.mm(8 + 6 * depth), right=df.mm(0), first=df.mm(-5))
            runs = [{"text": _list_prefix(block)}] + runs
        elif kind == "quote":
            key, spec = "quote", _pspec(S["quote"])
        elif kind == "code":
            key, override_indent = "body", True
            spec = _pspec(S["body"], left=df.mm(0), right=df.mm(0), first=df.mm(0))
            runs = [{"text": block.get("text", ""), "code": True}]
        elif kind == "bib_heading":
            key = "headings.1"
            spec = _pspec(S[key], keep_next=True, page_break=self.fmt["bibliography"]["new_page"])
            runs = [{"text": block.get("text") or "참고문헌"}]
        elif kind == "bib_entry":
            key = "bibliography"
            if block.get("hanging", True):
                spec = self.bib_spec()
            else:  # 번호식 스타일: 내어쓰기 안 함
                override_indent = True
                spec = _pspec(S["bibliography"], first=df.mm(0))
        else:
            key, spec = "body", _pspec(S["body"])
        spec["page_break"] = bool((spec["page_break"] or self.pending_break) and not self.section_start)
        self.pending_break = self.section_start = False
        self.write(key, spec, runs, override_indent)

    def caption_spec(self, keep_next: bool = False) -> dict:
        """캡션 · 그림 문단: 본문 모양, 가운데, 들여쓰기 없음. 그림 · 표 위(keep_next)나 아래에 6pt 이상 띄운다"""
        body = self.styles["body"]
        before = max(float(body["space_before_pt"]), 6.0) if keep_next else float(body["space_before_pt"])
        after = float(body["space_after_pt"]) if keep_next else max(float(body["space_after_pt"]), 6.0)
        return _pspec(body, align="center", left=df.mm(0), right=df.mm(0), first=df.mm(0), before_pt=before,
                      after_pt=after, page_break=False, keep_next=keep_next)

    def cell_spec(self, align: str) -> dict:
        """표 칸 문단: 본문 글자 크기, 줄간격은 130%까지, 위아래 간격 · 들여쓰기 없음"""
        body = self.styles["body"]
        return _pspec(body, align=align, left=df.mm(0), right=df.mm(0), first=df.mm(0), before_pt=0, after_pt=0,
                      line_pct=min(float(body["line_spacing_pct"]), 130.0), page_break=False)

    def text_area_mm(self) -> tuple[float, float]:
        """본문 영역 (폭, 높이) mm — 그림 크기 상한"""
        p, m = self.fmt["page"], self.body_margins()
        width = p["width_mm"] - m["left"] - m["right"]
        height = p["height_mm"] - m["top"] - m["bottom"] - m["header"] - m["footer"]
        return max(20.0, width), max(20.0, height * 0.8)

    def figure(self, block: Block, spec: dict) -> None:
        """그림 문단(spec, 줄간격 100%) → 아래에 캡션 문단. 그림 문단 쓰기는 picture_paragraph(docx · hwpx 따로)"""
        self.picture_paragraph(block.get("image"), dict(spec, line_pct=100))
        caption = _runs_of(block.get("caption"))
        if caption:
            self.write_spec(self.caption_spec(), caption, role="figure")

    def table(self, block: Block, spec: dict) -> None:
        """위에 캡션 문단(spec) → 표. 표 쓰기는 table_grid(docx · hwpx 따로)"""
        caption = _runs_of(block.get("caption"))
        if caption or spec["page_break"]:
            self.write_spec(spec, caption, role="table" if caption else None)
        self.table_grid(block)

    def write_spec(self, spec: dict, runs: list[Run], role: str | None = None) -> None:
        """본문 모양 + spec의 정렬 · 간격 · 함께 두기를 모두 적용한 문단.
        role(figure · table) = 캡션 문단 — 워드는 그림 · 표 목차 필드가 모으는 캡션 스타일을 붙인다"""
        self.write("body", spec, runs, True)

    def contents(self, block: Block) -> None:
        """목차 · 그림 목차 · 표 목차 쪽 (doc_lists): 가운데 제목(제목 1수준 글자 모양 — 제목 스타일은 쓰지 않아
        목차에 다시 잡히지 않음) → 항목 줄(글 + 점선 + 쪽 번호). 줄 쓰기는 list_title · list_entries(docx · hwpx 따로)"""
        h1 = self.styles["headings.1"]
        spec = _pspec(h1, align="center", left=df.mm(0), right=df.mm(0), first=df.mm(0), keep_next=True,
                      page_break=bool(self.pending_break and not self.section_start))
        spec["after_pt"] = max(float(spec["after_pt"]), 12.0)
        self.pending_break = self.section_start = False
        self.list_title(spec, _rspec(h1), block["title"])
        self.list_entries(block)

    def entry_spec(self, level: int) -> dict:
        """목록 항목 문단: 본문 글자 · 줄간격, 왼쪽 정렬, 단계마다 2글자 들여쓰기, 위아래 간격 없음"""
        body = self.styles["body"]
        return _pspec(body, align="left", left=df.ch(2 * (max(1, level) - 1)), right=df.mm(0), first=df.mm(0),
                      before_pt=0, after_pt=0, page_break=False)

    def bib_spec(self) -> dict:
        """참고문헌 항목 문단 사양: 왼쪽 = 들여쓰기 + 내어쓰기 폭, 첫 줄 = −내어쓰기 폭."""
        st = self.styles["bibliography"]
        hang = self.fmt["bibliography"]["hanging_indent"]
        left = st["indent_left"]
        if float(left["value"]) and left["unit"] != hang["unit"]:
            total = df.length_pt(left, st["size_pt"]) + df.length_pt(hang, st["size_pt"])
            left = df.mm(round(total / df.PT_PER_MM, 3))
        elif float(left["value"]):
            left = {"value": float(left["value"]) + float(hang["value"]), "unit": hang["unit"]}
        else:
            left = hang
        return _pspec(st, left=left, first=_neg(hang))

    def cover_spec(self, ln: dict) -> tuple[dict, dict]:
        cover = self.fmt["cover"]
        pspec = {"size": ln["size"], "align": "center", "line_pct": df.COVER_LINE_HEIGHT * 100,
                 "before_pt": ln["space_before_mm"] * df.PT_PER_MM, "after_pt": 0,
                 "left": df.mm(0), "right": df.mm(0), "first": df.mm(0), "page_break": False, "keep_next": False}
        rspec = {"fonts": cover["fonts"] or self.fmt["fonts"], "size": ln["size"], "bold": cover["bold"],
                 "italic": False}
        return pspec, rspec

    def body_margins(self) -> dict:
        p = self.fmt["page"]
        return dict(p["margin_mm"], header=p["header_mm"], footer=p["footer_mm"])


# ------------------------------------------------------------------ docx
_DOCX_STYLE_NAMES = {"body": "Normal", "title": "Title", "headings.1": "Heading 1", "headings.2": "Heading 2",
                     "headings.3": "Heading 3", "quote": "Quote", "bibliography": "Bibliography"}


def _docx_ppr(ppr, spec: dict, *, explicit_page_break: bool = False) -> None:
    """문단 사양 → w:pPr (정렬·간격·고정 줄간격·들여쓰기)."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    ppr.keepNext_val = True if spec["keep_next"] else None
    if spec["page_break"] or explicit_page_break:
        ppr.pageBreakBefore_val = bool(spec["page_break"])
    sp = ppr.get_or_add_spacing()
    sp.set(qn("w:before"), str(df.pt_to_twip(spec["before_pt"])))
    sp.set(qn("w:after"), str(df.pt_to_twip(spec["after_pt"])))
    # 줄간격: 고정값 = 글자 크기 × % (사용자 답변 Q1, 11pt × 180% = 19.8pt)
    sp.set(qn("w:line"), str(df.pt_to_twip(spec["size"] * spec["line_pct"] / 100)))
    sp.set(qn("w:lineRule"), "exact")
    _docx_ind(ppr, spec)
    ppr.jc_val = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                  "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[spec["align"]]


def _docx_ind(ppr, spec: dict) -> None:
    """들여쓰기. ch(글자 수)는 *Chars(글자 수 × 100)와 twip을 함께 쓴다. 상속된 *Chars는 0으로 끊는다."""
    from docx.oxml.ns import qn

    ind = ppr.get_or_add_ind()
    for attr in list(ind.attrib):
        del ind.attrib[attr]
    size = spec["size"]

    def chars(length):
        return str(round(float(length["value"]) * 100)) if length["unit"] == "ch" else "0"

    ind.set(qn("w:leftChars"), chars(spec["left"]))
    ind.set(qn("w:left"), str(df.pt_to_twip(df.length_pt(spec["left"], size))))
    ind.set(qn("w:rightChars"), chars(spec["right"]))
    ind.set(qn("w:right"), str(df.pt_to_twip(df.length_pt(spec["right"], size))))
    first = df.pt_to_twip(df.length_pt(spec["first"], size))
    if first < 0:
        ind.set(qn("w:hangingChars"), chars(_neg(spec["first"])))
        ind.set(qn("w:hanging"), str(-first))
        ind.set(qn("w:firstLineChars"), "0")
    else:
        ind.set(qn("w:firstLineChars"), chars(spec["first"]))
        ind.set(qn("w:firstLine"), str(first))


def _docx_font(font, rpr, rspec: dict, *, style: bool = False) -> None:
    """글자 사양 → 글꼴(영문 ascii·hAnsi, 한글 eastAsia)·크기·굵게·기울임."""
    from docx.oxml.ns import qn
    from docx.shared import Pt

    fonts = rspec.get("fonts")
    if fonts:
        font.name = fonts["latin"]
        rfonts = rpr.get_or_add_rFonts()
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)
        rfonts.set(qn("w:cs"), fonts["latin"])
        # 워드는 한자 글꼴을 따로 지정할 수 없어 한글 글꼴을 쓴다(4.3)
        rfonts.set(qn("w:eastAsia"), fonts["hangul"])
    if rspec.get("size"):
        font.size = Pt(rspec["size"])
    if style:
        font.bold = True if rspec.get("bold") else None
        font.italic = True if rspec.get("italic") else None
    else:
        if rspec.get("bold"):
            font.bold = True
        if rspec.get("italic"):
            font.italic = True


class _DocxBuilder(_FormattedBuilder):
    def __init__(self, fmt: dict, meta: dict):
        from docx import Document

        super().__init__(fmt)
        self.document = Document()
        self.document.core_properties.title = meta.get("title") or ""
        self.document.core_properties.author = meta.get("author") or ""
        self.notes = None
        self.last_p = None      # 지금 구역의 마지막 문단(구역 나눔을 여기에 붙인다)
        self.last_is_cover = False
        self.started = False    # 첫 구역(문서 끝 sectPr)을 썼는지
        self.style_page_break: dict[str, bool] = {}
        self.update_fields = False  # 목차 필드를 넣었으면 워드가 열 때 필드를 새로 계산하게(settings updateFields)
        self._setup_styles()

    def _setup_styles(self) -> None:
        from docx.enum.style import WD_STYLE_TYPE
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        styles = self.document.styles
        for key, name in _DOCX_STYLE_NAMES.items():
            try:
                style = styles[name]
            except KeyError:
                style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            if name != "Normal":
                style.base_style = styles["Normal"]
            el = style.element
            # 기본 서식 파일의 꾸밈(제목 밑줄·색·테마 글꼴 등)을 지우고 양식 값만 넣는다
            for tag in ("w:pPr", "w:rPr"):
                old = el.find(qn(tag))
                if old is not None:
                    el.remove(old)
            st = self.styles[key]
            if key == "bibliography":
                spec = self.bib_spec()
            else:
                spec = _pspec(st, keep_next=key.startswith(("title", "headings")))
            self.style_page_break[name] = bool(spec["page_break"])
            ppr = el.get_or_add_pPr()
            _docx_ppr(ppr, spec)
            if key.startswith("headings"):
                lvl = OxmlElement("w:outlineLvl")
                lvl.set(qn("w:val"), str(int(key[-1]) - 1))
                ppr.append(lvl)
            _docx_font(style.font, el.get_or_add_rPr(), _rspec(st), style=True)

    # --- 구역
    def _start_section(self, *, top, bottom, left, right, header, footer) -> None:
        from docx.section import Section
        from docx.shared import Twips

        body = self.document.element.body
        sentinel = body.get_or_add_sectPr()
        if self.started:
            if self.last_p is None or not self.last_is_cover:
                # 워드는 구역 나눔이 붙은 양쪽 정렬 문단의 끝줄까지 벌리므로, 본문 구역은
                # 1pt 높이의 빈 문단을 따로 두고 거기에 구역 나눔을 붙인다(표지 줄은 가운데 정렬이라 그대로)
                self.last_p = self._section_break_paragraph()
            # 지금까지의 구역 설정을 그 구역 마지막 문단에 붙이고, 문서 끝 sectPr은 새 구역이 된다
            self.last_p.set_sectPr(sentinel.clone())
            for el in sentinel.xpath("w:headerReference|w:footerReference|w:pgNumType|w:titlePg"):
                sentinel.remove(el)
        self.started = True
        sec = Section(sentinel, self.document.part)
        page = self.fmt["page"]
        sec.page_width = Twips(df.mm_to_twip(page["width_mm"]))
        sec.page_height = Twips(df.mm_to_twip(page["height_mm"]))
        sec.top_margin, sec.bottom_margin = Twips(df.mm_to_twip(top)), Twips(df.mm_to_twip(bottom))
        sec.left_margin, sec.right_margin = Twips(df.mm_to_twip(left)), Twips(df.mm_to_twip(right))
        sec.header_distance, sec.footer_distance = Twips(df.mm_to_twip(header)), Twips(df.mm_to_twip(footer))
        sec.gutter = Twips(0)
        self.section = sec
        self.section_start = True
        self.last_p = None
        self.last_is_cover = False

    def _section_break_paragraph(self):
        from docx.shared import Pt

        p = self.document.add_paragraph()
        ppr = p._p.get_or_add_pPr()
        _docx_ppr(ppr, {"size": 1, "align": "left", "line_pct": 100, "before_pt": 0, "after_pt": 0,
                        "left": df.mm(0), "right": df.mm(0), "first": df.mm(0), "page_break": False,
                        "keep_next": False}, explicit_page_break=True)
        # 문단 기호 글자 크기도 1pt로 (빈 문단 높이 최소화)
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        mark = ppr.find(qn("w:rPr"))
        if mark is None:
            mark = OxmlElement("w:rPr")
            ppr.append(mark)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(Pt(1).pt * 2)))
        mark.append(sz)
        return p._p

    def cover_page(self, page: dict, lay: dict) -> None:
        m = page["margin_lr_mm"]
        self._start_section(top=lay["top_mm"], bottom=COVER_BOTTOM_MARGIN_MM, left=m, right=m, header=0, footer=0)
        for ln in lay["lines"]:
            pspec, rspec = self.cover_spec(ln)
            p = self.document.add_paragraph()
            _docx_ppr(p._p.get_or_add_pPr(), pspec)
            run = p.add_run(_clean_text(ln["text"]))
            _docx_font(run.font, run._r.get_or_add_rPr(), rspec)
            self.last_p = p._p
            self.last_is_cover = True

    def body_section(self, numbered: bool) -> None:
        # 워드 여백 환산(Q2): 위 여백 = 위쪽 + 머리말, 머리글 거리 = 위쪽 (아래도 같음)
        m = self.body_margins()
        header, footer = m["top"], m["bottom"]
        distance = self.fmt["page_number"]["distance_mm"]
        if numbered and distance is not None:
            # 쪽 번호 위치(distance_mm): 머리글/바닥글 거리만 바꾸고 본문 위·아래 여백은 그대로
            if _PAGE_NUMBER_PLACES[self.fmt["page_number"]["position"]][0] == "footer":
                footer = distance
            else:
                header = distance
        self._start_section(top=m["top"] + m["header"], bottom=m["bottom"] + m["footer"], left=m["left"],
                            right=m["right"], header=header, footer=footer)
        if numbered:
            self._page_number()

    def _page_number(self) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.text.run import Run as DocxRun

        pn = self.fmt["page_number"]
        where, align, _ = _PAGE_NUMBER_PLACES[pn["position"]]
        sect = self.section._sectPr
        num = OxmlElement("w:pgNumType")
        num.set(qn("w:start"), "1")
        anchor = sect.find(qn("w:cols"))
        if anchor is None:
            anchor = sect.find(qn("w:docGrid"))
        if anchor is not None:
            anchor.addprevious(num)
        else:
            sect.append(num)
        hf = self.section.footer if where == "footer" else self.section.header
        hf.is_linked_to_previous = False
        p = hf.paragraphs[0] if hf.paragraphs else hf.add_paragraph()
        body = self.styles["body"]
        _docx_ppr(p._p.get_or_add_pPr(), _pspec(body, align=align, line_pct=100, before_pt=0, after_pt=0,
                                                left=df.mm(0), right=df.mm(0), first=df.mm(0), page_break=False))
        rspec = _rspec(body, bold=False, italic=False)

        def text_run(text):
            run = p.add_run(text)
            _docx_font(run.font, run._r.get_or_add_rPr(), rspec)

        if pn["dashes"]:
            text_run("- ")
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), " PAGE ")
        r = OxmlElement("w:r")
        run = DocxRun(r, p)
        _docx_font(run.font, r.get_or_add_rPr(), rspec)
        run.text = "1"
        fld.append(r)
        p._p.append(fld)
        if pn["dashes"]:
            text_run(" -")

    # --- 본문 블록
    def write(self, key: str, spec: dict, runs: list[Run], override_indent: bool) -> None:
        name = _DOCX_STYLE_NAMES[key]
        p = self.document.add_paragraph(style=name)
        ppr = p._p.get_or_add_pPr()
        if bool(spec["page_break"]) != self.style_page_break[name]:
            ppr.pageBreakBefore_val = bool(spec["page_break"])
        if override_indent:
            _docx_ind(ppr, spec)
        self._runs(p, runs)
        self.last_p = p._p

    def write_spec(self, spec: dict, runs: list[Run], role: str | None = None) -> None:
        p = self.document.add_paragraph(style=self._plain_style(_DOCX_CAPTION_STYLES[role]) if role else "Normal")
        _docx_ppr(p._p.get_or_add_pPr(), spec)
        self._runs(p, runs)
        self.last_p = p._p

    # --- 목차 · 그림 목차 · 표 목차 (doc_lists): 진짜 TOC 필드 + 어림 쪽 번호로 채운 결과
    def _plain_style(self, name: str, *, setup=None) -> str:
        """Normal을 따르는 문단 스타일(없으면 만듦). setup(style)로 처음 만들 때만 꾸민다"""
        from docx.enum.style import WD_STYLE_TYPE

        styles = self.document.styles
        try:
            styles[name]
        except KeyError:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = styles["Normal"]
            style.hidden = False
            style.quick_style = False
            if setup:
                setup(style)
        return name

    def _entry_style(self, kind: str, level: int) -> str:
        """목차 항목은 워드 기본 이름 'toc 1~3', 그림 · 표 목차는 'table of figures' — 워드가 필드를 새로 계산해도
        같은 스타일을 쓴다. 오른쪽 끝 점선 탭 · 단계별 들여쓰기"""
        from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
        from docx.shared import Twips

        page, m = self.fmt["page"], self.fmt["page"]["margin_mm"]
        width = df.mm_to_twip(page["width_mm"] - m["left"] - m["right"])

        def setup(style):
            _docx_ppr(style.element.get_or_add_pPr(), self.entry_spec(level))
            style.paragraph_format.tab_stops.add_tab_stop(Twips(width), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

        return self._plain_style(f"toc {level}" if kind == "toc" else "table of figures", setup=setup)

    def list_title(self, spec: dict, rspec: dict, text: str) -> None:
        p = self.document.add_paragraph(style="Normal")
        _docx_ppr(p._p.get_or_add_pPr(), spec)
        run = p.add_run(_clean_text(text))
        _docx_font(run.font, run._r.get_or_add_rPr(), rspec)
        self.last_p = p._p

    def list_entries(self, block: Block) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        def field(p, kind: str, instr: str | None = None) -> None:
            r = OxmlElement("w:r")
            if instr is None:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            else:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = instr
            r.append(el)
            p._p.append(r)

        kind, items = block["kind"], block["entries"]
        # \o 제목 1~N수준 · \u 개요 수준 문단 · \h 항목을 하이퍼링크로 · \z 웹 보기에서만 쪽 번호 숨김
        if kind == "toc":
            instr = f' TOC \\o "1-{int(block.get("depth") or 2)}" \\h \\z \\u '
        else:  # 그림 · 표 캡션 스타일 문단을 모은다
            instr = f' TOC \\h \\z \\t "{_DOCX_CAPTION_STYLES["figure" if kind == "lof" else "table"]},1" '
        for i, e in enumerate(items):
            p = self.document.add_paragraph(style=self._entry_style(kind, e["level"]))
            if i == 0:
                field(p, "begin")
                field(p, "", instr)
                field(p, "separate")
            p.add_run(_clean_text(e["text"]))
            p.add_run("\t" + (str(e["page"]) if e.get("page") else ""))
            if i == len(items) - 1:
                field(p, "end")
            self.last_p = p._p
        self.update_fields = True

    def picture_paragraph(self, image: dict | None, spec: dict) -> None:
        p = self.document.add_paragraph(style="Normal")
        _docx_ppr(p._p.get_or_add_pPr(), spec)
        if image and _docx_picture(p, image, *self.text_area_mm()):
            _docx_single_spacing(p)
        else:
            if image:
                _figure_warn(self.warnings)
            self._runs(p, [{"text": FIGURE_MISSING}])
        self.last_p = p._p

    def table_grid(self, block: Block) -> None:
        def fill(p, runs, align):
            _docx_ppr(p._p.get_or_add_pPr(), self.cell_spec(align))
            self._runs(p, runs)

        _docx_table(self.document, block, fill)
        self.last_p = None  # 표 뒤에 문단이 없으면 finish가 빈 문단을 붙인다 (구역 나눔은 따로 빈 문단을 만듦)

    def _runs(self, paragraph, runs: list[Run]) -> None:
        from docx.shared import Pt

        for r in runs:
            if "footnote" in r:
                if self.notes is None:
                    self.notes = _DocxFootnotes(self.document)
                fn = self.styles["footnote"]
                _footnote_styled(self.notes, paragraph, r["footnote"], _pspec(fn, page_break=False), _rspec(fn))
                continue
            run = paragraph.add_run(_clean_text(r.get("text", "")))
            run.bold = bool(r.get("b")) or None
            run.italic = bool(r.get("i")) or None
            if r.get("sup"):
                run.font.superscript = True
            if r.get("sub"):
                run.font.subscript = True
            if r.get("code"):
                run.font.name = CODE_FONT
                run.font.size = Pt(CODE_SIZE)

    def finish(self) -> bytes:
        if self.last_p is None:
            self.document.add_paragraph()
        if self.update_fields:
            _docx_update_fields_on_open(self.document)
        buf = io.BytesIO()
        self.document.save(buf)
        return buf.getvalue()


_DOCX_CAPTION_STYLES = {"figure": "Figure Caption", "table": "Table Caption"}  # 그림 · 표 목차 필드(\t)가 모으는 스타일
# w:settings 안에서 updateFields 뒤에 와야 하는 요소들(OOXML CT_Settings 순서 — 순서가 틀리면 워드가 파일을 고치려 함)
_SETTINGS_AFTER_UPDATE_FIELDS = (
    "hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr", "attachedSchema",
    "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
    "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags",
    "decimalSymbol", "listSeparator")


def _docx_update_fields_on_open(document) -> None:
    """워드가 문서를 열 때 필드(목차)를 새로 계산하게 한다 — '이 문서에는 다른 파일을 참조하는 필드가…' 물음에
    [예]를 누르면 목차 쪽 번호가 실제 쪽으로 바뀐다(LibreOffice · 한컴은 저장된 어림값을 그대로 보여 줌)"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    settings = document.settings.element
    if settings.find(qn("w:updateFields")) is not None:
        return
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    after = next((c for c in settings if isinstance(c.tag, str) and c.tag.rsplit("}", 1)[-1]
                  in _SETTINGS_AFTER_UPDATE_FIELDS), None)
    if after is not None:
        after.addprevious(el)
    else:
        settings.append(el)


def _footnote_styled(notes: _DocxFootnotes, paragraph, runs: list[Run], pspec: dict, rspec: dict) -> None:
    """양식의 각주 모양(크기·줄간격·글꼴)으로 각주를 단다."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.text.run import Run as DocxRun

    nid = notes.next_id
    notes.next_id += 1
    ref = OxmlElement("w:r")
    DocxRun(ref, paragraph).font.superscript = True
    el = OxmlElement("w:footnoteReference")
    el.set(qn("w:id"), str(nid))
    ref.append(el)
    paragraph._p.append(ref)

    note = OxmlElement("w:footnote")
    note.set(qn("w:id"), str(nid))
    p = OxmlElement("w:p")
    _docx_ppr(p.get_or_add_pPr(), pspec)
    note.append(p)
    mark = OxmlElement("w:r")
    mark_run = DocxRun(mark, paragraph)
    _docx_font(mark_run.font, mark.get_or_add_rPr(), rspec)
    mark_run.font.superscript = True
    mark.append(OxmlElement("w:footnoteRef"))
    p.append(mark)
    for r in [{"text": " "}] + list(runs):
        if "footnote" in r:
            continue
        r_el = OxmlElement("w:r")
        run = DocxRun(r_el, paragraph)
        _docx_font(run.font, r_el.get_or_add_rPr(), dict(rspec, bold=bool(r.get("b")), italic=bool(r.get("i"))))
        if r.get("sup"):
            run.font.superscript = True
        if r.get("sub"):
            run.font.subscript = True
        run.text = _clean_text(r.get("text", ""))
        p.append(r_el)
    notes.part.element.append(note)


# ------------------------------------------------------------------ hwpx
class _HwpxBuilder(_FormattedBuilder):
    def __init__(self, fmt: dict):
        from hwpx import HwpxDocument

        super().__init__(fmt)
        self.doc = HwpxDocument.new()
        self.header = self.doc.oxml.headers[0]
        self._first = self.doc.paragraphs[0] if self.doc.paragraphs else None
        self.sec_index = 0
        self.started = False
        self.numbered_sections: list[int] = []
        self._para_cache: dict = {}
        self._char_cache: dict = {}
        self._font_ids: dict = {}
        # '바탕글'·'각주' 스타일이 양식의 본문·각주 모양을 가리키게 한다(각주 내용은 '각주' 스타일을 쓴다)
        body, fn = self.styles["body"], self.styles["footnote"]
        self._set_style_refs("Normal", self.para_id(_pspec(body)), self.char_id(_rspec(body)))
        self._set_style_refs("Footnote", self.para_id(_pspec(fn, page_break=False)), self.char_id(_rspec(fn)))

    def _set_style_refs(self, eng_name: str, para_id: str, char_id: str) -> None:
        for style in self.header.element.iter(f"{_HH}style"):
            if style.get("engName") == eng_name:
                style.set("paraPrIDRef", str(para_id))
                style.set("charPrIDRef", str(char_id))
        self.header.mark_dirty()

    def _font_id(self, face: str, lang: str) -> str:
        if face not in self._font_ids:
            self.doc.oxml.ensure_font(face)
            ids = {}
            for ff in self.header.element.iter(f"{_HH}fontface"):
                for font in ff.findall(f"{_HH}font"):
                    if font.get("face") == face:
                        ids[ff.get("lang")] = font.get("id")
            self._font_ids[face] = ids
        return self._font_ids[face][lang]

    def para_id(self, spec: dict) -> str:
        key = tuple(sorted((k, repr(v)) for k, v in spec.items()))
        if key not in self._para_cache:
            size = spec["size"]

            def hwp(length):
                return df.pt_to_hwp(df.length_pt(length, size))

            margins = {"intent": hwp(spec["first"]), "left": hwp(spec["left"]), "right": hwp(spec["right"]),
                       "prev": df.pt_to_hwp(spec["before_pt"]), "next": df.pt_to_hwp(spec["after_pt"])}
            tab = {}
            if spec.get("tab_right_mm"):  # 목록 항목: 오른쪽 끝 탭 + 점선(한컴 차례와 같은 채움 3번 = DASH)
                tab["tab_pr_id_ref"] = self.header.ensure_tab_definition(tab_stops=[
                    {"pos": df.mm_to_hwp(spec["tab_right_mm"]), "type": "RIGHT", "leader": _HWPX_TOC_LEADER[1]}])
            self._para_cache[key] = self.header.ensure_paragraph_format(
                base_para_pr_id="0", alignment=_HWPX_ALIGN[spec["align"]],
                line_spacing_percent=round(spec["line_pct"]), margins=margins,
                break_setting={"page_break_before": bool(spec["page_break"]),
                               "keep_with_next": bool(spec["keep_next"])}, **tab)
        return self._para_cache[key]

    def char_id(self, rspec: dict, script: str | None = None) -> str:
        fonts = rspec["fonts"]
        key = (fonts["hangul"], fonts["latin"], fonts["hanja"], float(rspec["size"]), bool(rspec["bold"]),
               bool(rspec["italic"]), script)
        if key not in self._char_cache:
            base = self.doc.oxml.ensure_run_style(bold=bool(rspec["bold"]), italic=bool(rspec["italic"]),
                                                  script=script, size=rspec["size"])
            face = {"HANGUL": fonts["hangul"], "LATIN": fonts["latin"], "HANJA": fonts["hanja"]}
            # 일본어·기타·기호·사용자 글꼴은 한글 글꼴을 쓴다
            ref = {lang.lower(): self._font_id(face.get(lang, fonts["hangul"]), lang) for lang in _HWPX_LANGS}

            def set_fonts(el):
                font_ref = el.find(f"{_HH}fontRef")
                if font_ref is None:
                    font_ref = el.makeelement(f"{_HH}fontRef", {})
                    el.insert(0, font_ref)
                for name, value in ref.items():
                    font_ref.set(name, value)

            el = self.header.ensure_char_property(modifier=set_fonts, base_char_pr_id=base)
            self._char_cache[key] = el.get("id")
        return self._char_cache[key]

    # --- 구역
    def _start_section(self, margins: dict) -> None:
        from hwpx._document import layout

        if self.started:
            sec = self.doc.add_section()
            self.sec_index = len(self.doc.sections) - 1
            self._first = sec.paragraphs[0] if sec.paragraphs else None
        self.started = True
        page = self.fmt["page"]
        layout.set_page_setup(self.doc, width_mm=page["width_mm"], height_mm=page["height_mm"],
                              margins_mm=margins, section_index=self.sec_index)
        self.section_start = True

    def _paragraph(self):
        if self._first is not None:
            p, self._first = self._first, None
            return p
        return self.doc.add_paragraph("", section_index=self.sec_index, include_run=False, inherit_style=False)

    def cover_page(self, page: dict, lay: dict) -> None:
        m = page["margin_lr_mm"]
        self._start_section({"top": lay["top_mm"], "bottom": COVER_BOTTOM_MARGIN_MM, "left": m, "right": m,
                             "header": 0, "footer": 0})
        for ln in lay["lines"]:
            pspec, rspec = self.cover_spec(ln)
            p = self._paragraph()
            p.para_pr_id_ref = self.para_id(pspec)
            p.add_run(_clean_text(ln["text"]), char_pr_id_ref=self.char_id(rspec))

    def body_section(self, numbered: bool) -> None:
        from hwpx._document import layout

        pn = self.fmt["page_number"]
        where, align, position = _PAGE_NUMBER_PLACES[pn["position"]]
        margins = self.body_margins()
        if numbered and pn["distance_mm"] is not None:
            # 쪽 번호 위치(distance_mm): 꼬리말(머리말) 글은 영역의 용지 쪽 끝에 붙으므로
            # 아래쪽(위쪽) = 거리, 꼬리말(머리말) = 나머지로 나눠 본문 영역은 그대로 둔다
            edge, band = ("bottom", "footer") if where == "footer" else ("top", "header")
            total = margins[edge] + margins[band]
            d = min(pn["distance_mm"], total)
            margins[edge], margins[band] = d, total - d
        self._start_section(margins)
        if not numbered:
            return
        self.doc.sections[self.sec_index].properties.set_start_numbering(page=1)
        layout.set_page_number(self.doc, target=where, align=align.upper(), position=position,
                               prefix="- " if pn["dashes"] else "", suffix=" -" if pn["dashes"] else "",
                               section_index=self.sec_index)
        self.numbered_sections.append(self.sec_index)

    def write(self, key: str, spec: dict, runs: list[Run], override_indent: bool) -> None:
        p = self._paragraph()
        p.para_pr_id_ref = self.para_id(spec)
        self._add_runs(p, self.styles[key], runs)

    def list_title(self, spec: dict, rspec: dict, text: str) -> None:
        p = self._paragraph()
        p.para_pr_id_ref = self.para_id(spec)
        p.add_run(_clean_text(text), char_pr_id_ref=self.char_id(rspec))

    def list_entries(self, block: Block) -> None:
        """정적 목록: 한 줄 = [단계 들여쓰기(전각 공백) + 글][탭(점선)][쪽 번호(어림값)] (doc_lists 참고).
        탭 위치는 문단 들여쓰기와 무관하게 본문 오른쪽 끝이 되게 단계 들여쓰기는 문단 여백 대신 전각 공백으로 둔다"""
        width = self.text_area_mm()[0]
        body = self.styles["body"]
        cid = self.char_id(_rspec(body, bold=False, italic=False))
        para = self.para_id(dict(self.entry_spec(1), tab_right_mm=width))
        size_hwp = df.pt_to_hwp(float(body["size_pt"]))
        for e in block["entries"]:
            text = df.FULL_SPACE * 2 * (max(1, e["level"]) - 1) + _clean_text(e["text"])
            p = self._paragraph()
            p.para_pr_id_ref = para
            t = p.add_run(text, char_pr_id_ref=cid).element.find(f"{_HP}t")
            # hp:tab 속성은 한글이 다시 배치하는 캐시값 — 글 길이로 어림한 남은 폭을 넣어 둔다
            full = df.mm_to_hwp(width)
            used = int(sum(1.0 if ord(c) >= 0x1100 else 0.3 if c == " " else 0.55 for c in text) * size_hwp) % full
            tab = t.makeelement(f"{_HP}tab", {"width": str(full - used), "leader": str(_HWPX_TOC_LEADER[0]),
                                              "type": "2"})
            tab.tail = str(e["page"]) if e.get("page") else ""
            t.append(tab)

    def picture_paragraph(self, image: dict | None, spec: dict) -> None:
        p = self._paragraph()
        p.para_pr_id_ref = self.para_id(spec)
        if image:
            _hwpx_picture(self.doc, p, image, *self.text_area_mm(),
                          char_pr_id_ref=self.char_id(_rspec(self.styles["body"], bold=False, italic=False)))
        else:
            self._add_runs(p, self.styles["body"], [{"text": FIGURE_MISSING}])

    def table_grid(self, block: Block) -> None:
        rows = _table_rows(block)
        anchor = self._paragraph()
        anchor.para_pr_id_ref = self.para_id(dict(self.caption_spec(), line_pct=100))
        body = self.styles["body"]
        row_mm = body["size_pt"] * min(float(body["line_spacing_pct"]), 130.0) / 100 / df.PT_PER_MM + 1.5
        t = _hwpx_table(anchor, block, rows, self.text_area_mm()[0], row_mm,
                        self.char_id(_rspec(body, bold=False, italic=False)))
        for r, row in enumerate(rows):
            for c, runs in enumerate(row):
                p = t.cell(r, c).paragraphs[0]
                p.para_pr_id_ref = self.para_id(self.cell_spec(_cell_align(block, r, c)))
                self._add_runs(p, self.styles["body"], _cell_runs(block, r, runs))

    def _add_runs(self, p, st: dict, runs: list[Run]) -> None:
        for r in runs:
            if "footnote" in r:
                text = plain_text(r["footnote"]).strip()
                p.add_footnote(_clean_text(text) or " ")
                continue
            text = _clean_text(r.get("text", ""))
            if not text:
                continue
            script = "sup" if r.get("sup") else ("sub" if r.get("sub") else None)
            rspec = _rspec(st, bold=st["bold"] or bool(r.get("b")), italic=st["italic"] or bool(r.get("i")))
            if r.get("code"):
                rspec.update(fonts=dict(st["fonts"], latin=CODE_FONT), size=CODE_SIZE)
            _hwpx_text_run(p, text, char_pr_id_ref=self.char_id(rspec, script))

    def finish(self) -> bytes:
        # 쪽 번호 글자는 본문 글꼴·크기
        body_char = self.char_id(_rspec(self.styles["body"], bold=False, italic=False))
        for index in self.numbered_sections:
            section = self.doc.sections[index]
            for hf in list(section.element.iter(f"{_HP}footer")) + list(section.element.iter(f"{_HP}header")):
                for run in hf.iter(f"{_HP}run"):
                    run.set("charPrIDRef", str(body_char))
            section.mark_dirty()
        report = self.doc.validate()
        if getattr(report, "issues", ()):
            raise ValueError(f"한글 문서 검증 실패: {report.issues[:3]}")
        return self.doc.to_bytes()


# ------------------------------------------------------------------ 수정 내용 표시 (2026-10-27)
# 화면(writeredline.js)이 넣은 글을 U+E000…U+E001, 지운 글을 U+E002…U+E003으로 감싸 보낸다. 워드 파일을 다 만든 뒤
# 글자 런을 그 자리에서 나눠 w:ins · w:del(변경 내용 추적)로 감싼다 — 양식 · 기본 서식 · 표 · 각주 어디서 만든 런이든 같게
REDLINE_MARKS = ""
_REDLINE_RE = re.compile("([-])")
REDLINE_AUTHOR = "PaperLab"


def strip_redline(blocks):
    """표시 문자 지우기 (한글 · 마크다운 · 수정 표시 없는 내보내기)"""
    if isinstance(blocks, str):
        return re.sub("[-]", "", blocks)
    if isinstance(blocks, list):
        return [strip_redline(b) for b in blocks]
    if isinstance(blocks, dict):
        return {k: strip_redline(v) for k, v in blocks.items()}
    return blocks


def redline_docx(data: bytes, when: str = "") -> bytes:
    """워드 파일의 표시 문자 → 변경 내용 추적. 표시는 문단마다 새로 시작(문단을 넘는 표시는 문단 끝에서 닫힘)"""
    import copy
    import io
    import zipfile

    from lxml import etree

    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    q = lambda t: f"{{{W}}}{t}"  # noqa: E731
    when = when or datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    next_id = [9000]

    def fix(xml: bytes) -> bytes:
        root = etree.fromstring(xml)
        for p in root.iter(q("p")):
            state = None
            for r in [r for r in p.iter(q("r")) if r.getparent().tag != q("del") and r.getparent().tag != q("ins")]:
                ts = r.findall(q("t"))
                if len(ts) != 1 or not _REDLINE_RE.search(ts[0].text or ""):
                    if state and ts:  # 표시 안쪽의 다른 런(굵게 · 위첨자 등)
                        wrap(r, state)
                    continue
                parent, idx = r.getparent(), r.getparent().index(r)
                parent.remove(r)
                pieces = []
                for part in _REDLINE_RE.split(ts[0].text or ""):
                    if part in ("", ""):
                        state = "ins" if part == "" else "del"
                    elif part in ("", ""):
                        state = None
                    elif part:
                        nr = copy.deepcopy(r)
                        t = nr.find(q("t"))
                        t.text = part
                        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                        pieces.append((nr, state))
                for k, (nr, st) in enumerate(pieces):
                    parent.insert(idx + k, nr)
                    if st:
                        wrap(nr, st)
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    def wrap(r, st):
        box = etree.Element(q(st))
        box.set(q("id"), str(next_id[0]))
        next_id[0] += 1
        box.set(q("author"), REDLINE_AUTHOR)
        box.set(q("date"), when)
        r.addprevious(box)
        box.append(r)
        if st == "del":
            for t in r.findall(q("t")):
                t.tag = q("delText")

    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            body = zin.read(item.filename)
            if item.filename in ("word/document.xml", "word/footnotes.xml") and any(m.encode() in body for m in REDLINE_MARKS):
                body = fix(body)
            zout.writestr(item, body)
    return out.getvalue()
