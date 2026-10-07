"""워드·한글 연동: 문서 속 인용 표시([@인용키])를 서식 있는 인용으로 바꾸고 참고문헌을 넣는다.

워드(.docx)나 한글(.hwpx)에서 평소처럼 쓰다가 인용할 자리에 [@vaswani2017attention] 처럼 적고,
참고문헌을 넣을 자리에 [참고문헌] 을 한 줄로 적는다. 파일을 올리면
1) scan  : 인용 표시를 문서 순서대로 찾는다
2) 화면이 citeproc-js로 선택한 스타일에 맞춰 인용 문구·참고문헌을 만든다
3) apply : 원래 글꼴·문단 서식은 그대로 두고 표시만 바꾼 파일을 돌려준다

인용 표시 문법은 Pandoc과 같다: [@a] · [@a; @b] · [@a, p. 12] · [@a, 34쪽]
"""

from __future__ import annotations

import copy
import io
import re
import zipfile
from dataclasses import dataclass

# 앞에 \가 붙은 [는 인용 표시가 아님(마크다운 이스케이프 — 참고 패널이 넣은 PDF 글 속 [@…], refquote.js와 같은 규칙)
CITE_RE = re.compile(r"(?<!\\)\[(?=[^\[\]]*@)([^\[\]]{1,400})\]")
KEY_RE = re.compile(r"(-?)@([\w][\w:.#$%&\-+?<>~/]*)")
BIB_MARKERS = ("[참고문헌]", "[References]", "[Bibliography]", "[REFERENCES]")
LOCATOR_RE = re.compile(
    r"^\s*,?\s*(?:(p|pp|page|pages|쪽|면)\.?\s*)?([\divxlcIVXLC][\w\-–,\s]*?)\s*(쪽|면)?\s*$")


@dataclass
class Citation:
    raw: str
    items: list[dict]


def parse_citation(inner: str) -> list[dict] | None:
    """'@a, p. 12; @b' → [{key, locator, label, suppress_author}]. 인용이 아니면 None."""
    items = []
    for part in inner.split(";"):
        m = KEY_RE.search(part)
        if not m:
            return None
        prefix = part[:m.start()].strip()
        rest = part[m.end():]
        item: dict = {"key": m.group(2).rstrip(".,"), "suppress_author": m.group(1) == "-"}
        if prefix:
            item["prefix"] = prefix
        rest = rest.strip()
        if rest:
            lm = LOCATOR_RE.match(rest)
            if lm and (lm.group(1) or lm.group(3) or re.match(r"^\s*,?\s*\d", rest)):
                item["locator"] = lm.group(2).strip()
                item["label"] = "page"
            else:
                item["suffix"] = rest.lstrip(", ")
        items.append(item)
    return items or None


def find_citations(text: str) -> list[Citation]:
    out = []
    for m in CITE_RE.finditer(text):
        items = parse_citation(m.group(1))
        if items:
            out.append(Citation(m.group(0), items))
    return out


# ------------------------------------------------------------------ DOCX
def _docx_paragraphs(document):
    """본문 문단을 문서 순서대로 (표 안 문단 포함)."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    def walk(container):
        for child in container.iter_inner_content():
            if isinstance(child, Paragraph):
                yield child
            elif isinstance(child, Table):
                for row in child.rows:
                    seen = set()
                    for cell in row.cells:
                        if id(cell._tc) in seen:  # 병합된 칸은 한 번만
                            continue
                        seen.add(id(cell._tc))
                        yield from walk(cell)

    yield from walk(document)


def _docx_replace_span(paragraph, start: int, end: int, new_text: str):
    """문단 글자 [start, end)를 new_text로 바꾼다. 첫 글자가 있던 런의 서식을 따른다. 바뀐 런을 돌려준다."""
    pos = 0
    first_run = None
    for run in paragraph.runs:
        text = run.text
        r_start, r_end = pos, pos + len(text)
        pos = r_end
        if r_end <= start or r_start >= end:
            continue
        a = max(start, r_start) - r_start
        b = min(end, r_end) - r_start
        if first_run is None:
            run.text = text[:a] + new_text + text[b:]
            first_run = (run, a + len(new_text))
        else:
            run.text = text[:a] + text[b:]
    return first_run


def docx_scan(data: bytes) -> dict:
    from docx import Document

    document = Document(io.BytesIO(data))
    cites, has_bib = [], False
    for p in _docx_paragraphs(document):
        text = p.text
        cites += find_citations(text)
        if text.strip() in BIB_MARKERS:
            has_bib = True
    return {"citations": cites, "has_bib_marker": has_bib}


def docx_apply(data: bytes, rendered: list[dict], bibliography: list[list[dict]], bib_title: str,
               note_style: bool) -> bytes:
    """rendered[i] = {"runs": [...]} — i번째 인용 표시(문서 순서)를 바꿀 내용."""
    from docx import Document
    from docx.shared import Mm

    from .writer import _docx_run_element, _DocxFootnotes

    document = Document(io.BytesIO(data))
    notes = None
    idx = 0
    bib_paragraph = None
    for p in list(_docx_paragraphs(document)):
        text = p.text
        if text.strip() in BIB_MARKERS and bib_paragraph is None:
            bib_paragraph = p
            continue
        # 앞에서부터 차례로 바꾼다 (각주 번호가 문서 순서대로 붙도록). 바꿀 때마다 글자를 다시 읽는다
        cursor = 0
        while True:
            text = p.text
            m = next((m for m in CITE_RE.finditer(text, cursor) if parse_citation(m.group(1))), None)
            if m is None:
                break
            runs = (rendered[idx].get("runs") if idx < len(rendered) else None) or [{"text": m.group(0)}]
            idx += 1
            if note_style:
                hit = _docx_replace_span(p, m.start(), m.end(), "")
                if hit:
                    run, offset = hit
                    if notes is None:
                        notes = _docx_existing_or_new_footnotes(document)
                    _docx_insert_footnote_at(run, offset, notes, runs)
                cursor = m.start()
            else:
                hit = _docx_replace_span(p, m.start(), m.end(), "\ue000")
                if hit:
                    run, offset = hit
                    _docx_insert_runs_at(run, offset - 1, runs, _docx_run_element)
                cursor = m.start() + len("".join(r.get("text", "") for r in runs if "footnote" not in r))
            if not hit:
                break

    if bibliography:
        entries = []
        for runs in bibliography:
            para = document.add_paragraph()
            para.paragraph_format.left_indent = Mm(10)
            para.paragraph_format.first_line_indent = Mm(-10)
            for r in runs:
                para._p.append(_docx_run_element(r))
            entries.append(para)
        if bib_paragraph is not None:
            # [참고문헌] 줄을 제목으로 바꾸고 그 뒤에 목록을 옮긴다
            for run in bib_paragraph.runs[1:]:
                run._r.getparent().remove(run._r)
            if bib_paragraph.runs:
                bib_paragraph.runs[0].text = bib_title
            anchor = bib_paragraph._p
        else:
            heading = document.add_paragraph(bib_title)
            heading.runs[0].bold = True
            anchor = heading._p
        for para in entries:
            anchor.addnext(para._p)
            anchor = para._p
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _docx_insert_runs_at(run, offset: int, runs: list[dict], make_run):
    """run 글자의 offset 위치(자리표시 U+E000)를 서식 있는 런들로 바꾼다. 원래 런의 서식을 바탕으로 한다."""
    text = run.text
    before, after = text[:offset], text[offset + 1:]
    run.text = before
    anchor = run._r
    base_rpr = run._r.rPr
    for r in runs:
        el = make_run(r)
        if base_rpr is not None:
            merged = copy.deepcopy(base_rpr)
            new_rpr = el.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr")
            if new_rpr is not None:
                for child in list(new_rpr):
                    for old in merged.findall(child.tag):
                        merged.remove(old)
                    merged.append(child)
                el.remove(new_rpr)
            el.insert(0, merged)
        anchor.addnext(el)
        anchor = el
    if after:
        tail = copy.deepcopy(run._r)
        for t in tail.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"):
            t.getparent().remove(t)
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = after
        tail.append(t)
        anchor.addnext(tail)
    if not before:
        run._r.getparent().remove(run._r)


def _docx_existing_or_new_footnotes(document):
    from .writer import FOOTNOTES_RT, _DocxFootnotes

    for rel in document.part.rels.values():
        if rel.reltype == FOOTNOTES_RT:
            notes = _DocxFootnotes.__new__(_DocxFootnotes)
            notes.part = rel.target_part
            # 문서에 이미 있는 각주 번호 다음부터
            if not hasattr(notes.part, "element"):
                from docx.opc.part import XmlPart
                from docx.oxml import parse_xml
                xml_part = XmlPart(notes.part.partname, notes.part.content_type, parse_xml(notes.part.blob),
                                   notes.part.package)
                document.part.rels[rel.rId]._target = xml_part
                notes.part = xml_part
            w = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            ids = [int(x) for x in notes.part.element.xpath("//w:footnote/@w:id", namespaces=w)
                   if x.lstrip("-").isdigit()]
            notes.next_id = max([0] + ids) + 1
            return notes
    return _DocxFootnotes(document)


def _docx_insert_footnote_at(run, offset: int, notes, runs: list[dict]):
    from docx.text.paragraph import Paragraph

    text = run.text
    run.text = text[:offset]
    tail = None
    if text[offset:]:
        tail = copy.deepcopy(run._r)
    # 각주 번호 런을 run 바로 뒤에 넣기 위해 임시 문단에 만든 뒤 옮긴다
    paragraph = Paragraph(run._r.getparent(), None)
    before = len(paragraph._p)
    notes.add(paragraph, runs)
    ref = paragraph._p[before:][0] if len(paragraph._p) > before else paragraph._p[-1]
    run._r.addnext(ref)
    if tail is not None:
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        for t in tail.findall(ns + "t"):
            tail.remove(t)
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text[offset:]
        tail.append(t)
        ref.addnext(tail)


# ------------------------------------------------------------------ HWPX
def hwpx_scan(data: bytes) -> dict:
    from hwpx import HwpxDocument

    doc = HwpxDocument.open(io.BytesIO(data))
    text = doc.text.plain()
    lines = [ln.strip() for ln in text.splitlines()]
    return {"citations": find_citations(text), "has_bib_marker": any(ln in BIB_MARKERS for ln in lines)}


def hwpx_apply(data: bytes, rendered: list[dict], bibliography: list[list[dict]], bib_title: str) -> bytes:
    from hwpx import HwpxDocument

    from .writer import _clean_text, plain_text

    doc = HwpxDocument.open(io.BytesIO(data))
    text = doc.text.plain()
    cites = find_citations(text)
    # 같은 표시는 같은 문구가 된다 (저자-연도·번호 스타일)
    done = set()
    for c, r in zip(cites, rendered):
        if c.raw in done:
            continue
        done.add(c.raw)
        doc.text.replace(c.raw, _clean_text(plain_text(r.get("runs") or [])), everywhere=True)

    if bibliography:
        from hwpx._document import layout

        marker_p = None
        for p in doc.paragraphs:
            if (p.text or "").strip() in BIB_MARKERS:
                marker_p = p
                break
        italic_id = None
        new_paragraphs = []
        if marker_p is None:
            head = doc.add_paragraph("", include_run=False, inherit_style=False)
            head.add_run(bib_title, char_pr_id_ref=doc.oxml.ensure_run_style(bold=True))
        for runs in bibliography:
            para = doc.add_paragraph("", include_run=False, inherit_style=False)
            for r in runs:
                t = _clean_text(r.get("text", ""))
                if not t:
                    continue
                if r.get("i"):
                    italic_id = italic_id or doc.oxml.ensure_run_style(italic=True)
                    para.add_run(t, char_pr_id_ref=italic_id)
                else:
                    para.add_run(t)
            layout.set_paragraph_format(doc, paragraphs=[para], indent_left_mm=10, first_line_indent_mm=-10)
            new_paragraphs.append(para)
        if marker_p is not None:
            doc.text.replace(marker_p.text.strip(), bib_title, everywhere=True)
            anchor = marker_p.element
            for para in new_paragraphs:
                el = para.element
                el.getparent().remove(el)
                anchor.addnext(el)
                anchor = el
            marker_p.section.mark_dirty()
    report = doc.validate()
    if getattr(report, "issues", ()):
        raise ValueError(f"한글 문서 검증 실패: {report.issues[:3]}")
    return doc.to_bytes()


def detect_kind(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if not zipfile.is_zipfile(io.BytesIO(data)):
        if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
            raise ValueError("예전 형식(.doc · .hwp)이에요. 워드에서는 .docx, 한글에서는 .hwpx로 다른 이름으로 저장해 주세요.")
        raise ValueError("워드(.docx)나 한글(.hwpx) 파일만 쓸 수 있어요")
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = set(z.namelist())
    if "word/document.xml" in names:
        return "docx"
    if any(n.startswith("Contents/") for n in names) or name.endswith(".hwpx"):
        return "hwpx"
    raise ValueError("워드(.docx)나 한글(.hwpx) 파일만 쓸 수 있어요")
