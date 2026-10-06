"""논문 양식(서식 프리셋) — docs/specs/doc-formats.md 13장 수용 기준(자동 확인 항목).

사용자 답변으로 명세 수치가 바뀐 항목:
- Q1 워드 줄간격 = 고정값(글자 크기 × %, lineRule=exact)  → AC-18·20·21·22·37·41·43
- Q5 학위논문 인용문 한 탭 = 40pt(14.11mm)                → AC-20
- Q12 신명조·신명 세명조 = 한양신명조                    → AC-37·38
"""
import io
import json
import re
import sqlite3
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from urllib.parse import unquote

import pytest
from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from hwpx import HwpxDocument

from paperlab import doc_formats as df
from paperlab import format_import, writer
from paperlab.server import create_app

GOLDEN = Path(__file__).parent / "golden"
BLOCKS = json.loads((GOLDEN / "blocks.json").read_text(encoding="utf-8"))
META = {"title": "스마트 제조 연구"}
COVER = {"degree": "master", "degree_field": "공학", "title_en": "A Study", "graduation": "2027-02",
         "department": "스마트제조공학과", "name": "김인하", "spaced_name": True, "advisors": ["홍길동"],
         "committee": ["", "", ""]}
HP = "{http://www.hancom.co.kr/hwpml/2011/paragraph}"
HH = "{http://www.hancom.co.kr/hwpml/2011/head}"
HC = "{http://www.hancom.co.kr/hwpml/2011/core}"
FULL = "　"


@pytest.fixture
def client(tmp_path):
    return TestClient(create_app(tmp_path), headers={"X-PaperLab": "1"})


def fmt(fid: str, **changes) -> dict:
    data = df.builtin_data(fid)
    for path, value in changes.items():
        cur = data
        keys = path.split("__")
        for k in keys[:-1]:
            cur = cur[k]
        cur[keys[-1]] = value
    return df.normalize(data, fid)


def strip_hwpx(xml: bytes) -> bytes:
    s = xml.decode("utf-8")
    s = re.sub(r'(<hp:p\b[^>]*?)\sid="[^"]*"', r"\1", s)
    s = re.sub(r'\sinstId="[^"]*"', "", s)
    return s.encode("utf-8")


def zread(data: bytes, name: str) -> bytes:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return z.read(name)


# ------------------------------------------------------------------ docx 읽기 도우미
def eff_p(p, getter):
    """문단의 실제 적용값: 직접 서식 → 스타일 → basedOn 사슬."""
    ppr = p._p.pPr
    if ppr is not None:
        v = getter(ppr)
        if v is not None:
            return v
    style = p.style
    while style is not None:
        ppr = style.element.pPr
        if ppr is not None:
            v = getter(ppr)
            if v is not None:
                return v
        style = style.base_style
    return None


def eff_r(p, run, getter):
    rpr = run._r.rPr
    if rpr is not None:
        v = getter(rpr)
        if v is not None:
            return v
    style = p.style
    while style is not None:
        rpr = style.element.rPr
        if rpr is not None:
            v = getter(rpr)
            if v is not None:
                return v
        style = style.base_style
    return None


def attr(tag, name):
    def get(el):
        child = el.find(qn(tag))
        return None if child is None else child.get(qn(name))
    return get


def flag(tag):
    def get(el):
        child = el.find(qn(tag))
        if child is None:
            return None
        return child.get(qn("w:val"), "1") not in ("0", "false")
    return get


def docx_sections(doc):
    """[(구역 sectPr, [문단…])] — 문단에 붙은 sectPr이 그 구역의 끝."""
    from docx.text.paragraph import Paragraph

    out, cur = [], []
    for el in doc.element.body:
        if el.tag == qn("w:p"):
            cur.append(Paragraph(el, doc._body))
            ppr = el.pPr
            if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
                out.append((ppr.find(qn("w:sectPr")), cur))
                cur = []
    out.append((doc.element.body.find(qn("w:sectPr")), cur))
    return out


def footer_xml(doc, sect, kind="footer"):
    for ref in sect.findall(qn(f"w:{kind}Reference")):
        part = doc.part.related_parts[ref.get(qn("r:id"))]
        return part.blob.decode("utf-8") if hasattr(part, "blob") and part.blob else part._element.xml
    return ""


def para_by_text(doc, text):
    return next(p for p in doc.paragraphs if p.text == text)


def tw_mm(v):
    return int(v) * 25.4 / 1440


# ------------------------------------------------------------------ hwpx 읽기 도우미
class Hwpx:
    def __init__(self, data: bytes):
        z = zipfile.ZipFile(io.BytesIO(data))
        self.header = ET.fromstring(z.read("Contents/header.xml"))
        names = sorted((n for n in z.namelist() if re.fullmatch(r"Contents/section\d+\.xml", n)),
                       key=lambda n: int(re.findall(r"\d+", n)[0]))
        self.sections = [ET.fromstring(z.read(n)) for n in names]
        self.fonts = {ff.get("lang"): {f.get("id"): f.get("face") for f in ff.findall(f"{HH}font")}
                      for ff in self.header.iter(f"{HH}fontface")}
        self.chars = {c.get("id"): c for c in self.header.iter(f"{HH}charPr")}
        self.paras = {c.get("id"): c for c in self.header.iter(f"{HH}paraPr")}

    def paragraphs(self, index=None):
        secs = self.sections if index is None else [self.sections[index]]
        return [p for sec in secs for p in sec.findall(f"{HP}p")]

    @staticmethod
    def text(p):
        out = []
        for run in p.findall(f"{HP}run"):
            out += [t.text or "" for t in run.findall(f"{HP}t")]
        return "".join(out)

    def para_with(self, text):
        return next(p for p in self.paragraphs() if self.text(p) == text)

    def char_of(self, p):
        # 구역 첫 문단에는 구역 설정(secPr)만 담은 빈 런이 먼저 있다
        run = next(r for r in p.findall(f"{HP}run") if "".join(t.text or "" for t in r.findall(f"{HP}t")))
        return self.chars[run.get("charPrIDRef")]

    def font_names(self, char):
        ref = char.find(f"{HH}fontRef")
        return {k: self.fonts[k.upper()][ref.get(k)] for k in ("hangul", "latin", "hanja")}

    def para_pr(self, p):
        return self.paras[p.get("paraPrIDRef")]

    @staticmethod
    def case(pp, tag):
        el = pp.find(f"{HH}{tag}")
        return el if el is not None else pp.find(f"{HP}switch/{HP}case/{HH}{tag}")

    def line(self, p):
        ls = self.case(self.para_pr(p), "lineSpacing")
        return ls.get("type"), int(ls.get("value"))

    def margin(self, p, name):
        return int(self.case(self.para_pr(p), "margin").find(f"{HC}{name}").get("value"))

    def page_margin(self, index):
        page = self.sections[index].find(f".//{HP}pagePr")
        return page, {k: int(v) for k, v in page.find(f"{HP}margin").attrib.items()}


def thesis_docx(**kw):
    return writer.to_docx(BLOCKS, META, fmt("inha-mie-thesis"), kw.get("cover", COVER), kw.get("warnings"))


def thesis_hwpx():
    return writer.to_hwpx(BLOCKS, META, fmt("inha-mie-thesis"), COVER)


# =================================================================== A. 기본 (A4) 호환
def _check_golden(docx_bytes, hwpx_bytes):
    for part, name in (("word/document.xml", "docx_document.xml"), ("word/styles.xml", "docx_styles.xml"),
                       ("word/footnotes.xml", "docx_footnotes.xml")):
        assert zread(docx_bytes, part) == (GOLDEN / name).read_bytes(), part
    assert zread(hwpx_bytes, "Contents/header.xml") == (GOLDEN / "hwpx_header.xml").read_bytes()
    assert strip_hwpx(zread(hwpx_bytes, "Contents/section0.xml")) == (GOLDEN / "hwpx_section0.xml").read_bytes()


def test_ac01_ac02_default_export_matches_golden():
    _check_golden(writer.to_docx(BLOCKS, META), writer.to_hwpx(BLOCKS, META))


def test_ac03_default_format_and_markdown_via_api(client):
    def export(fmt_name, **extra):
        r = client.post("/api/export-document", json={"format": fmt_name, "blocks": BLOCKS, "meta": META, **extra})
        assert r.status_code == 200, r.text
        return r.content

    _check_golden(export("docx"), export("hwpx"))
    _check_golden(export("docx", doc_format="default"), export("hwpx", doc_format="default"))
    md = export("md")
    for fid in df.BUILTINS:
        assert export("md", doc_format=fid, cover=COVER) == md


# =================================================================== B. 양식 API
def test_ac06_list_builtins(client):
    items = client.get("/api/doc-formats").json()
    assert [i["id"] for i in items] == ["default", "apa7-student", "inha-mie-thesis", "inha-mie-report"]
    assert all(i["builtin"] and i["base"] is None and i["updated_at"] is None for i in items)
    assert [i["cover_kind"] for i in items] == ["none", "none", "thesis", "report"]
    assert all(i["used_by"] == 0 for i in items)


def test_ac07_ac08_create_and_patch(client):
    made = client.post("/api/doc-formats", json={"base": "inha-mie-thesis", "name": "내 학위논문"}).json()
    assert made["id"].startswith("user-") and made["builtin"] is False and made["base"] == "inha-mie-thesis"
    assert made["data"] == client.get("/api/doc-formats/inha-mie-thesis").json()["data"]
    assert made["created_at"] and made["cover_kind"] == "thesis"
    assert made["id"] in [i["id"] for i in client.get("/api/doc-formats").json()]

    data = dict(made["data"])
    data["body"] = dict(data["body"], size_pt=12)
    r = client.patch(f"/api/doc-formats/{made['id']}", json={"data": data})
    assert r.status_code == 200
    got = client.get(f"/api/doc-formats/{made['id']}").json()
    assert got["data"]["body"]["size_pt"] == 12 and got["updated_at"] != made["updated_at"]
    client.patch(f"/api/doc-formats/{made['id']}", json={"name": "새 이름"})
    again = client.get(f"/api/doc-formats/{made['id']}").json()
    assert again["name"] == "새 이름" and again["data"] == got["data"]
    # data는 통째로 바꾸고 빠진 항목은 base 값으로 채운다
    client.patch(f"/api/doc-formats/{made['id']}", json={"data": {"body": {"size_pt": 10}}})
    partial = client.get(f"/api/doc-formats/{made['id']}").json()["data"]
    assert partial["body"]["size_pt"] == 10 and partial["page"]["width_mm"] == 188


def test_ac09_builtin_readonly_and_missing(client):
    assert client.patch("/api/doc-formats/default", json={"name": "x"}).status_code == 403
    assert client.delete("/api/doc-formats/inha-mie-thesis").status_code == 403
    r = client.get("/api/doc-formats/user-999")
    assert r.status_code == 404 and r.json()["detail"] == "양식을 찾을 수 없어요"
    assert client.get("/api/doc-formats/nope").status_code == 404
    assert client.post("/api/doc-formats", json={"base": "user-1", "name": "x"}).status_code == 400
    assert client.post("/api/doc-formats", json={"base": "default", "name": ""}).status_code == 400
    assert client.post("/api/doc-formats", json={"base": "default", "name": "가" * 61}).status_code == 400


@pytest.mark.parametrize("data, path", [
    ({"body": {"size_pt": 0}}, "body.size_pt"),
    ({"body": {"align": "middle"}}, "body.align"),
    ({"page": {"width_mm": 1000}}, "page.width_mm"),
    ({"body": {"first_line_indent": {"value": 3, "unit": "cm"}}}, "body.first_line_indent.unit"),
    ({"page_number": {"distance_mm": 200}}, "page_number.distance_mm"),
    ({"page_number": {"distance_mm": "11"}}, "page_number.distance_mm"),
    ({"quote": {"indent_left": {"value": 300, "unit": "pt"}}}, "quote.indent_left"),
    # 팀장 결정: 여백·머리말·꼬리말·문단 간격 0~100mm, 들여쓰기·내어쓰기 -50~100mm (ch·pt는 mm 환산)
    ({"page": {"margin_mm": {"top": 120}}}, "page.margin_mm.top"),
    ({"page": {"header_mm": -1}}, "page.header_mm"),
    ({"body": {"space_before_pt": 300}}, "body.space_before_pt"),
    ({"quote": {"indent_left": {"value": -60, "unit": "mm"}}}, "quote.indent_left"),
    ({"body": {"first_line_indent": {"value": 30, "unit": "ch"}}}, "body.first_line_indent"),
    ({"bibliography": {"hanging_indent": {"value": 101, "unit": "mm"}}}, "bibliography.hanging_indent"),
])
def test_ac10_validation_reports_path(client, data, path):
    r = client.post("/api/doc-formats", json={"base": "default", "name": "x", "data": data})
    assert r.status_code == 400 and path in r.json()["detail"], r.json()


def test_ac10_pt_length_and_distance_accepted(client):
    r = client.post("/api/doc-formats", json={"base": "default", "name": "x", "data": {
        "quote": {"indent_left": {"value": 40, "unit": "pt"}}, "page_number": {"distance_mm": 11}}})
    assert r.status_code == 200, r.text
    assert r.json()["data"]["quote"]["indent_left"] == {"value": 40, "unit": "pt"}
    assert r.json()["data"]["page_number"]["distance_mm"] == 11
    assert client.get("/api/doc-formats/inha-mie-report").json()["data"]["page_number"]["distance_mm"] is None


def test_length_ranges_accept_edges():
    ok = df.normalize({"body": {"first_line_indent": {"value": 20, "unit": "ch"}, "space_after_pt": 283},
                       "quote": {"indent_left": {"value": -50, "unit": "mm"}},
                       "page": {"margin_mm": {"top": 100}, "height_mm": 400}})
    assert ok["body"]["first_line_indent"] == {"value": 20, "unit": "ch"}  # 20 × 11pt ≈ 77.6mm


def test_ac11_delete_resets_manuscripts_and_default(client):
    fid = client.post("/api/doc-formats", json={"base": "inha-mie-report", "name": "보고서"}).json()["id"]
    ids = [client.post("/api/manuscripts", json={}).json()["id"] for _ in range(3)]
    for mid in ids[:2]:
        assert client.patch(f"/api/manuscripts/{mid}", json={"doc_format": fid}).status_code == 200
    assert client.put("/api/settings", json={"doc_format_default": fid}).json()["doc_format_default"] == fid
    used = {i["id"]: i["used_by"] for i in client.get("/api/doc-formats").json()}
    assert used[fid] == 2 and used["default"] == 1 and used["inha-mie-thesis"] == 0
    assert client.get(f"/api/doc-formats/{fid}").json()["used_by"] == 2
    assert client.delete(f"/api/doc-formats/{fid}").json() == {"ok": True, "reset_manuscripts": 2}
    assert all(client.get(f"/api/manuscripts/{m}").json()["doc_format"] == "default" for m in ids)
    assert client.get("/api/settings").json()["doc_format_default"] == "default"
    assert client.get("/api/doc-formats").json()[0]["used_by"] == 3


def test_deleted_format_id_is_not_reused(client, tmp_path):
    first = client.post("/api/doc-formats", json={"base": "default", "name": "A"}).json()["id"]
    mid = client.post("/api/manuscripts", json={}).json()["id"]
    client.patch(f"/api/manuscripts/{mid}", json={"doc_format": first})
    client.put("/api/settings", json={"doc_format_default": first})
    assert client.delete(f"/api/doc-formats/{first}").json()["reset_manuscripts"] == 1
    second = client.post("/api/doc-formats", json={"base": "default", "name": "B"}).json()["id"]
    assert second != first
    assert client.get(f"/api/doc-formats/{first}").status_code == 404
    assert client.get(f"/api/manuscripts/{mid}").json()["doc_format"] == "default"
    assert client.get("/api/settings").json()["doc_format_default"] == "default"
    # 설정 파일에 남은 옛 id(지워진 양식)도 기본 양식으로 돌려준다
    settings = json.loads((tmp_path / "settings.json").read_text(encoding="utf-8"))
    settings["doc_format_default"] = "user-999"
    (tmp_path / "settings.json").write_text(json.dumps(settings), encoding="utf-8")
    c2 = TestClient(create_app(tmp_path), headers={"X-PaperLab": "1"})
    assert c2.get("/api/settings").json()["doc_format_default"] == "default"
    assert c2.post("/api/manuscripts", json={}).json()["doc_format"] == "default"


def test_old_doc_formats_table_upgraded(tmp_path):
    conn = sqlite3.connect(tmp_path / "library.db")
    conn.execute("CREATE TABLE doc_formats (id INTEGER PRIMARY KEY, name TEXT NOT NULL, base TEXT NOT NULL "
                 "DEFAULT 'default', data TEXT NOT NULL DEFAULT '{}', created_at TEXT NOT NULL, updated_at TEXT NOT NULL)")
    conn.execute("INSERT INTO doc_formats (id, name, created_at, updated_at) VALUES (5, '옛 양식', 'x', 'x')")
    conn.commit()
    conn.close()
    c = TestClient(create_app(tmp_path), headers={"X-PaperLab": "1"})
    assert c.get("/api/doc-formats/user-5").json()["name"] == "옛 양식"
    c.delete("/api/doc-formats/user-5")
    assert c.post("/api/doc-formats", json={"base": "default", "name": "새"}).json()["id"] == "user-6"
    sql = sqlite3.connect(tmp_path / "library.db").execute(
        "SELECT sql FROM sqlite_master WHERE name = 'doc_formats'").fetchone()[0]
    assert "AUTOINCREMENT" in sql


def test_ac12_write_requests_need_header(tmp_path):
    bare = TestClient(create_app(tmp_path))
    assert bare.post("/api/doc-formats", json={"base": "default", "name": "x"}).status_code == 403
    assert bare.patch("/api/doc-formats/user-1", json={"name": "x"}).status_code == 403
    assert bare.delete("/api/doc-formats/user-1").status_code == 403
    files = {"file": ("a.docx", _docx_fixture(), "application/octet-stream")}
    assert bare.post("/api/doc-formats/import", files=files).status_code == 403
    assert bare.get("/api/doc-formats").status_code == 200


# =================================================================== C. 원고 설정
def test_ac13_ac14_manuscript_format_and_cover(client):
    mid = client.post("/api/manuscripts", json={}).json()["id"]
    m = client.get(f"/api/manuscripts/{mid}").json()
    assert m["doc_format"] == "default" and m["cover"] == {}
    assert client.patch(f"/api/manuscripts/{mid}", json={"doc_format": "inha-mie-thesis"}).json()["ok"]
    assert client.get(f"/api/manuscripts/{mid}").json()["doc_format"] == "inha-mie-thesis"
    assert client.get("/api/manuscripts").json()[0]["doc_format"] == "inha-mie-thesis"
    assert client.patch(f"/api/manuscripts/{mid}", json={"doc_format": "user-77"}).status_code == 400

    cover = dict(COVER, title_ko="", subtitle="", approval="", school="인하대학교 제조혁신전문대학원",
                 include={"front": True, "inner": True, "approval": False})
    assert client.patch(f"/api/manuscripts/{mid}", json={"cover": cover}).status_code == 200
    assert client.get(f"/api/manuscripts/{mid}").json()["cover"] == cover
    for bad in ({"graduation": "2027-13"}, {"advisors": ["a", "b", "c"]},
                {"degree": "master", "committee": [""] * 6}, {"degree": "master", "committee": [""] * 4},
                {"degree": "bachelor"}, {"approval": "2027/02"}, {"spaced_name": "yes"}):
        r = client.patch(f"/api/manuscripts/{mid}", json={"cover": dict(COVER, **bad)})
        assert r.status_code == 400 and r.json()["detail"].startswith("cover."), bad
    assert client.patch(f"/api/manuscripts/{mid}", json={"cover": dict(COVER, degree="doctor",
                                                                       committee=[""] * 5)}).status_code == 200


def test_ac15_old_library_gets_new_columns(tmp_path):
    conn = sqlite3.connect(tmp_path / "library.db")
    conn.execute("CREATE TABLE manuscripts (id INTEGER PRIMARY KEY, title TEXT NOT NULL DEFAULT '', "
                 "content TEXT NOT NULL DEFAULT '', template TEXT NOT NULL DEFAULT '', style TEXT NOT NULL DEFAULT '', "
                 "created_at TEXT NOT NULL, updated_at TEXT NOT NULL)")
    conn.execute("INSERT INTO manuscripts (title, content, created_at, updated_at) VALUES ('옛 원고', '# 옛', 'x', 'x')")
    conn.commit()
    conn.close()
    c = TestClient(create_app(tmp_path), headers={"X-PaperLab": "1"})
    m = c.get("/api/manuscripts/1").json()
    assert m["title"] == "옛 원고" and m["doc_format"] == "default" and m["cover"] == {}
    cols = {r[1] for r in sqlite3.connect(tmp_path / "library.db").execute("PRAGMA table_info(manuscripts)")}
    assert {"doc_format", "cover"} <= cols


def test_ac16_new_manuscript_uses_setting(client):
    assert client.put("/api/settings", json={"doc_format_default": "nope"}).status_code == 400
    assert client.put("/api/settings", json={"doc_format_default": "inha-mie-report"}).status_code == 200
    assert client.post("/api/manuscripts", json={}).json()["doc_format"] == "inha-mie-report"
    assert client.post("/api/manuscripts", json={"doc_format": "apa7-student"}).json()["doc_format"] == "apa7-student"
    assert client.post("/api/manuscripts", json={"doc_format": "zzz"}).status_code == 400


# =================================================================== D. 학위논문 — 워드
def test_ac17_thesis_docx_body_section():
    doc = Document(io.BytesIO(thesis_docx()))
    s = doc.sections[-1]
    assert round(s.page_width.mm) == 188 and round(s.page_height.mm) == 257
    assert abs(s.left_margin.mm - 24) < 0.1 and abs(s.right_margin.mm - 21) < 0.1
    assert abs(s.top_margin.mm - 20) < 0.1 and abs(s.header_distance.mm - 15) < 0.1
    assert abs(s.bottom_margin.mm - 29) < 0.1 and abs(s.footer_distance.mm - 21) < 0.1


def test_ac18_thesis_docx_body_paragraph():
    doc = Document(io.BytesIO(thesis_docx()))
    p = para_by_text(doc, "둘째 장 본문.")
    run = p.runs[0]
    assert eff_r(p, run, attr("w:rFonts", "w:eastAsia")) == "휴먼명조"
    assert eff_r(p, run, attr("w:rFonts", "w:ascii")) == "휴먼명조"
    assert eff_r(p, run, attr("w:sz", "w:val")) == "22"
    # Q1: 고정 줄간격 = 11pt × 180% = 19.8pt = 396twip
    assert eff_p(p, attr("w:spacing", "w:line")) == "396" and eff_p(p, attr("w:spacing", "w:lineRule")) == "exact"
    assert eff_p(p, attr("w:ind", "w:firstLineChars")) == "300" and eff_p(p, attr("w:ind", "w:firstLine")) == "660"
    assert eff_p(p, attr("w:jc", "w:val")) == "both"
    assert eff_p(p, attr("w:spacing", "w:before")) == "0" and eff_p(p, attr("w:spacing", "w:after")) == "0"


def test_ac19_thesis_docx_headings():
    doc = Document(io.BytesIO(thesis_docx()))
    h1 = para_by_text(doc, "2. 본론")
    assert eff_r(h1, h1.runs[0], attr("w:sz", "w:val")) == "32" and eff_r(h1, h1.runs[0], flag("w:b"))
    assert eff_r(h1, h1.runs[0], attr("w:rFonts", "w:eastAsia")) == "휴먼명조"
    assert eff_p(h1, attr("w:jc", "w:val")) == "center" and eff_p(h1, flag("w:pageBreakBefore"))
    h2 = para_by_text(doc, "1.1 연구 배경")
    assert eff_r(h2, h2.runs[0], attr("w:sz", "w:val")) == "26" and eff_r(h2, h2.runs[0], flag("w:b"))
    assert eff_p(h2, attr("w:jc", "w:val")) == "left"
    h3 = para_by_text(doc, "1.1.1 세부 항목")
    assert eff_r(h3, h3.runs[0], attr("w:sz", "w:val")) == "22" and eff_r(h3, h3.runs[0], flag("w:b"))
    # 워드 목차(참조 → 목차)가 잡을 수 있게 제목 스타일을 쓴다
    assert h1.style.name == "Heading 1" and h2.style.name == "Heading 2"


def test_ac20_thesis_docx_quote():
    doc = Document(io.BytesIO(thesis_docx()))
    p = para_by_text(doc, "인용문 문단입니다.")
    assert eff_r(p, p.runs[0], attr("w:sz", "w:val")) == "22"
    assert eff_p(p, attr("w:spacing", "w:line")) == "352"  # 11pt × 160% = 17.6pt (고정)
    assert eff_p(p, attr("w:ind", "w:left")) == "800"      # Q5: 한 탭 40pt = 800twip
    assert eff_p(p, attr("w:ind", "w:firstLine")) == "0" and eff_p(p, attr("w:ind", "w:firstLineChars")) == "0"


def test_ac21_thesis_docx_footnotes():
    notes = zread(thesis_docx(), "word/footnotes.xml").decode()
    body = notes.split('w:id="1">', 1)[1]
    assert 'w:line="280" w:lineRule="exact"' in body  # 10pt × 140% = 14pt
    assert re.search(r'<w:sz w:val="20"/>.*?<w:t xml:space="preserve">각주 </w:t>', body, re.S)


def test_ac22_thesis_docx_bibliography():
    doc = Document(io.BytesIO(thesis_docx()))
    p = next(p for p in doc.paragraphs if p.text.startswith("Vaswani"))
    assert eff_r(p, p.runs[0], attr("w:sz", "w:val")) == "22"
    assert eff_p(p, attr("w:spacing", "w:line")) == "396"
    assert eff_p(p, attr("w:ind", "w:left")) == "567" and eff_p(p, attr("w:ind", "w:hanging")) == "567"
    assert p.paragraph_format.first_line_indent is None  # 스타일에 있음
    assert p.style.paragraph_format.first_line_indent.twips == -567
    numbered = next(p for p in doc.paragraphs if p.text.startswith("[1]"))
    assert numbered.paragraph_format.first_line_indent.twips == 0
    head = para_by_text(doc, "참고문헌")
    assert eff_p(head, flag("w:pageBreakBefore")) and eff_p(head, attr("w:jc", "w:val")) == "center"


def test_ac23_title_not_in_body():
    doc = Document(io.BytesIO(thesis_docx()))
    sections = docx_sections(doc)
    body_texts = [p.text for _, ps in sections[3:] for p in ps]
    assert "스마트 제조 연구" not in body_texts and "국문 초록: 첫 장 앞의 본문 문단입니다." in body_texts


THESIS_FRONT = ["공학석사학위 논문", "스마트 제조 연구", "A Study", "2027년 2월", "인하대학교 제조혁신전문대학원",
                "스마트제조공학과", "김 인 하"]
THESIS_INNER = THESIS_FRONT[:4] + ["지도교수 홍 길 동", "이 논문을 공학석사학위 논문으로 제출함"] + THESIS_FRONT[4:]


def _approval_ok(lines, titles, statement, date):
    assert lines[0] == statement and lines[1] == date and len(lines) == 2 + len(titles)
    for line, title in zip(lines[2:], titles):
        assert line.startswith(title) and line.endswith("(인)")
        assert line == f"{title}{FULL * 2}{FULL * 6}{FULL * 2}(인)"


def test_ac24_ac25_thesis_docx_cover_pages():
    doc = Document(io.BytesIO(thesis_docx()))
    sections = docx_sections(doc)
    front, inner, approval = ([p for p in ps] for _, ps in sections[:3])
    assert [p.text for p in front] == THESIS_FRONT
    assert [p.text for p in inner] == THESIS_INNER
    _approval_ok([p.text for p in approval], ["주심", "부심", "위원"],
                 "이 논문을 김인하의 공학석사학위논문으로 인정함.", "2026년 12월")
    for p in front + inner + approval:
        r = p.runs[0]
        assert eff_r(p, r, attr("w:rFonts", "w:eastAsia")) == "HY신명조" and eff_r(p, r, flag("w:b"))
        assert eff_p(p, attr("w:jc", "w:val")) == "center"
    assert [p.runs[0].font.size.pt for p in front] == [14, 16, 16, 14, 14, 14, 14]
    assert para_by_text(doc, "지도교수 홍 길 동").runs[0].font.size.pt == 16


def test_ac26_ac27_cover_variants():
    def cover_texts(**kw):
        doc = Document(io.BytesIO(thesis_docx(cover=dict(COVER, **kw))))
        return [[p.text for p in ps] for _, ps in docx_sections(doc)[:3]]

    front, _, approval = cover_texts(degree="doctor", committee=[""] * 5)
    assert front[0] == "공학박사학위 논문"
    _approval_ok(approval, ["주심", "부심", "위원", "위원", "위원"], "이 논문을 김인하의 공학박사학위논문으로 인정함.",
                 "2026년 12월")
    front, _, approval = cover_texts(graduation="2027-08")
    assert front[3] == "2027년 8월" and approval[1] == "2027년 6월"
    assert cover_texts(approval="2027-05")[2][1] == "2027년 5월"
    _, inner, _ = cover_texts(advisors=["이몽룡", "성춘향"])
    assert "공동지도교수 이 몽 룡" in inner and "공동지도교수 성 춘 향" in inner
    front, _, approval = cover_texts(spaced_name=False, committee=["박심사", "", ""])
    assert front[-1] == "김인하" and approval[2] == f"주심{FULL * 2}박심사{FULL * 2}(인)"
    assert approval[3] == f"부심{FULL * 2}{FULL * 6}{FULL * 2}(인)"


def test_ac28_thesis_docx_cover_positions():
    doc = Document(io.BytesIO(thesis_docx()))
    sections = docx_sections(doc)
    for (sect, ps), top in zip(sections[:3], (40, 40, 50)):
        margin = tw_mm(sect.find(qn("w:pgMar")).get(qn("w:top")))
        before = tw_mm(ps[0]._p.pPr.find(qn("w:spacing")).get(qn("w:before")))
        assert abs(margin + before - top) <= 1
    # '아래 40mm': 마지막 줄 아래 = 용지 높이 − 40 (계산값 확인)
    page = df.cover_pages("thesis", COVER, "스마트 제조 연구")[0]
    lay = df.layout_cover_page(page, 188, 257)
    bottom = lay["top_mm"] + sum(ln["space_before_mm"] + ln["height_mm"] for ln in lay["lines"])
    assert abs(bottom - (257 - 40)) < 0.01 and not lay["overflow"]


def test_b1_section_break_on_separate_empty_paragraph():
    """구역 나눔이 양쪽 정렬 본문 문단에 붙지 않는다(워드가 끝줄을 벌리지 않게)."""
    for fid in ("inha-mie-thesis", "inha-mie-report"):
        doc = Document(io.BytesIO(writer.to_docx(BLOCKS, META, fmt(fid), COVER)))
        sections = docx_sections(doc)
        sect, ps = sections[3]  # 첫 장 앞 구역(국문 초록)
        last = ps[-1]
        assert last.text == "" and last._p.pPr.find(qn("w:sectPr")) is not None
        assert last._p.pPr.find(qn("w:spacing")).get(qn("w:line")) == "20"
        assert last._p.pPr.find(qn("w:jc")).get(qn("w:val")) == "left"
        assert ps[-2].text == "국문 초록: 첫 장 앞의 본문 문단입니다." and ps[-2]._p.pPr.find(qn("w:sectPr")) is None
        # 표지 줄(가운데 정렬)은 빈 문단을 더하지 않는다
        assert all(p.text for _, cover in sections[:3] for p in cover)
    for p in Document(io.BytesIO(writer.to_docx(BLOCKS, META))).paragraphs:
        assert p._p.pPr is None or p._p.pPr.find(qn("w:sectPr")) is None


def test_ac29_thesis_docx_page_numbers():
    doc = Document(io.BytesIO(thesis_docx()))
    sections = docx_sections(doc)
    assert len(sections) == 5  # 표지·속표지·인정서·첫 장 앞·본문
    for sect, _ in sections[:4]:
        assert "PAGE" not in footer_xml(doc, sect) and sect.find(qn("w:pgNumType")) is None
    sect, ps = sections[4]
    assert ps[0].text == "1. 서론"
    assert sect.find(qn("w:pgNumType")).get(qn("w:start")) == "1"
    foot = footer_xml(doc, sect)
    assert 'w:instr=" PAGE "' in foot and '<w:jc w:val="center"/>' in foot and "- " not in foot


# =================================================================== E. 학위논문 — 한글
def test_ac31_thesis_hwpx_body_section():
    h = Hwpx(thesis_hwpx())
    body_index = next(i for i, sec in enumerate(h.sections)
                      if any(Hwpx.text(p) == "1. 서론" for p in sec.findall(f"{HP}p")))
    page, m = h.page_margin(body_index)
    assert abs(int(page.get("width")) - 53291) <= 28 and abs(int(page.get("height")) - 72850) <= 28
    expected = {"top": 4252, "header": 1417, "bottom": 5953, "footer": 2268, "left": 6803, "right": 5953}
    assert all(abs(m[k] - v) <= 28 for k, v in expected.items()), m


def test_ac32_ac33_thesis_hwpx_paragraphs():
    h = Hwpx(thesis_hwpx())
    p = h.para_with("둘째 장 본문.")
    assert h.line(p) == ("PERCENT", 180) and h.margin(p, "intent") == 3300
    assert h.para_pr(p).find(f"{HH}align").get("horizontal") == "JUSTIFY"
    char = h.char_of(p)
    assert char.get("height") == "1100"
    assert h.font_names(char) == {"hangul": "휴먼명조", "latin": "휴먼명조", "hanja": "휴먼명조"}
    h1 = h.char_of(h.para_with("2. 본론"))
    assert h1.get("height") == "1600" and h1.find(f"{HH}bold") is not None
    h2 = h.char_of(h.para_with("1.1 연구 배경"))
    assert h2.get("height") == "1300" and h2.find(f"{HH}bold") is not None
    assert h.line(h.para_with("인용문 문단입니다.")) == ("PERCENT", 160)
    assert h.margin(h.para_with("인용문 문단입니다."), "left") == 4000  # Q5: 40pt
    # 각주 내용: '각주' 스타일 → 10pt · 140%
    note_p = next(p for p in h.sections[-1].iter(f"{HP}p")
                  if p.find(f"{HP}run/{HP}t") is not None and "둘째 각주" in Hwpx.text(p))
    note_char = h.chars[note_p.find(f"{HP}run").get("charPrIDRef")]
    assert note_char.get("height") == "1000" and h.line(note_p) == ("PERCENT", 140)
    assert h.font_names(note_char)["hangul"] == "휴먼명조"


def test_ac34_ac35_thesis_hwpx_covers_and_numbers():
    data = thesis_hwpx()
    assert HwpxDocument.open(io.BytesIO(data)).validate().issues == ()
    h = Hwpx(data)
    texts = [[Hwpx.text(p) for p in sec.findall(f"{HP}p")] for sec in h.sections]
    assert texts[0] == THESIS_FRONT and texts[1] == THESIS_INNER
    _approval_ok(texts[2], ["주심", "부심", "위원"], "이 논문을 김인하의 공학석사학위논문으로 인정함.", "2026년 12월")
    for sec in h.sections[:3]:
        for p in sec.findall(f"{HP}p"):
            char = h.char_of(p)
            assert h.font_names(char)["hangul"] == "HY신명조" and char.find(f"{HH}bold") is not None
    # 표지 첫 줄 위치: 위 여백 + 첫 문단 앞 간격 = 40mm / 인정서 50mm
    for index, top in ((0, 40), (1, 40), (2, 50)):
        _, m = h.page_margin(index)
        first = h.sections[index].find(f"{HP}p")
        assert abs(df.hwp_to_mm(m["top"] + h.margin(first, "prev")) - top) <= 1
    assert len(h.sections) == 5
    for sec in h.sections[:4]:
        assert sec.find(f".//{HP}footer") is None and sec.find(f".//{HP}header") is None
    numbered = h.sections[4]
    assert Hwpx.text(numbered.find(f"{HP}p")) == "1. 서론"
    assert numbered.find(f".//{HP}startNum").get("page") == "1"
    footer = numbered.find(f".//{HP}footer")
    assert footer.find(f".//{HP}autoNum").get("numType") == "PAGE"
    assert footer.find(f".//{HP}pageNum").get("pos") == "BOTTOM_CENTER"
    assert "-" not in "".join(t.text or "" for t in footer.iter(f"{HP}t"))


# =================================================================== F. 대체 보고서
def report(fmt_name="docx", **cover):
    c = dict(COVER, subtitle="중소 제조기업 사례", **cover)
    fn = writer.to_docx if fmt_name == "docx" else writer.to_hwpx
    return fn(BLOCKS, META, fmt("inha-mie-report"), c)


def test_ac37_report_docx():
    doc = Document(io.BytesIO(report()))
    s = doc.sections[-1]
    assert round(s.page_width.mm) == 210 and round(s.page_height.mm) == 297
    assert abs(s.left_margin.mm - 30) < 0.1 and abs(s.right_margin.mm - 30) < 0.1
    assert abs(s.top_margin.mm - 55) < 0.1 and abs(s.header_distance.mm - 40) < 0.1
    assert abs(s.bottom_margin.mm - 55) < 0.1 and abs(s.footer_distance.mm - 40) < 0.1
    p = para_by_text(doc, "둘째 장 본문.")
    assert eff_r(p, p.runs[0], attr("w:rFonts", "w:eastAsia")) == "휴먼명조"
    assert eff_r(p, p.runs[0], attr("w:rFonts", "w:ascii")) == "한양신명조"  # Q12
    assert eff_r(p, p.runs[0], attr("w:sz", "w:val")) == "22" and eff_p(p, attr("w:spacing", "w:line")) == "396"
    h1, h2 = para_by_text(doc, "2. 본론"), para_by_text(doc, "1.1 연구 배경")
    assert eff_r(h1, h1.runs[0], attr("w:sz", "w:val")) == "32" and eff_r(h1, h1.runs[0], flag("w:b"))
    assert eff_r(h2, h2.runs[0], attr("w:sz", "w:val")) == "28" and eff_r(h2, h2.runs[0], flag("w:b"))
    notes = zread(report(), "word/footnotes.xml").decode()
    assert '<w:sz w:val="18"/>' in notes and 'w:line="252" w:lineRule="exact"' in notes  # 9pt × 140%
    quote = para_by_text(doc, "인용문 문단입니다.")
    assert eff_p(quote, attr("w:ind", "w:left")) == "800"  # 인용문 한 탭 40pt


REPORT_FRONT = [("석사학위 연구보고서", 16), ("스마트 제조 연구", 22), ("– 중소 제조기업 사례 –", 14),
                ("A Study", 16), ("2027년 2월", 16), ("인하대학교 제조혁신전문대학원", 16), ("스마트제조공학과", 16),
                ("김 인 하", 16)]


def test_ac38_report_hwpx():
    h = Hwpx(report("hwpx"))
    page, m = h.page_margin(len(h.sections) - 1)
    assert abs(int(page.get("width")) - 59528) <= 28 and abs(int(page.get("height")) - 84189) <= 28
    expected = {"top": 11339, "header": 4252, "bottom": 11339, "footer": 4252, "left": 8504, "right": 8504}
    assert all(abs(m[k] - v) <= 28 for k, v in expected.items()), m
    char = h.char_of(h.para_with("둘째 장 본문."))
    assert h.font_names(char) == {"hangul": "휴먼명조", "latin": "한양신명조", "hanja": "한양신명조"}
    assert h.margin(h.para_with("인용문 문단입니다."), "left") == 4000


def test_ac39_ac40_report_covers():
    for data in (report(), report(degree="doctor", committee=[""] * 5)):
        doc = Document(io.BytesIO(data))
        sections = docx_sections(doc)
        front, inner, approval = (ps for _, ps in sections[:3])
        assert [(p.text, p.runs[0].font.size.pt) for p in front] == REPORT_FRONT
        texts = [p.text for p in inner]
        assert "지도교수 홍길동" in texts and "이 보고서를 석사학위 연구보고서로 제출함" in texts
        assert texts.index("지도교수 홍길동") < texts.index("이 보고서를 석사학위 연구보고서로 제출함")
        assert approval[0].text == "이 보고서를 김인하의 석사학위 연구보고서로 인정함"
        assert approval[0].runs[0].font.size.pt == 22
        _approval_ok([p.text for p in approval], ["주심", "부심", "위원"],
                     "이 보고서를 김인하의 석사학위 연구보고서로 인정함", "2026년 12월")
        for (sect, ps), top in zip(sections[:3], (55, 55, 70)):
            margin = tw_mm(sect.find(qn("w:pgMar")).get(qn("w:top")))
            before = tw_mm(ps[0]._p.pPr.find(qn("w:spacing")).get(qn("w:before")))
            assert abs(margin + before - top) <= 1
        sect, ps = sections[-1]
        assert ps[0].text == "1. 서론" and sect.find(qn("w:pgNumType")).get(qn("w:start")) == "1"
        foot = footer_xml(doc, sect)
        assert 'w:instr=" PAGE "' in foot and '<w:jc w:val="center"/>' in foot and "- " not in foot
    h = Hwpx(report("hwpx"))
    assert [Hwpx.text(p) for p in h.sections[0].findall(f"{HP}p")] == [t for t, _ in REPORT_FRONT]
    assert [int(h.char_of(p).get("height")) // 100 for p in h.sections[0].findall(f"{HP}p")] == [
        s for _, s in REPORT_FRONT]


# =================================================================== G. APA 7
def test_ac41_apa_docx():
    data = writer.to_docx(BLOCKS, META, fmt("apa7-student"))
    doc = Document(io.BytesIO(data))
    sections = docx_sections(doc)
    assert len(sections) == 1
    s = doc.sections[0]
    assert abs(s.top_margin.mm - 25.4) < 0.1 and abs(s.header_distance.mm - 12.7) < 0.1
    assert abs(s.bottom_margin.mm - 25.4) < 0.1 and abs(s.footer_distance.mm - 12.7) < 0.1
    assert abs(s.left_margin.mm - 25.4) < 0.1 and abs(s.right_margin.mm - 25.4) < 0.1
    p = para_by_text(doc, "둘째 장 본문.")
    assert eff_r(p, p.runs[0], attr("w:rFonts", "w:ascii")) == "Times New Roman"
    assert eff_r(p, p.runs[0], attr("w:sz", "w:val")) == "24"
    assert eff_p(p, attr("w:spacing", "w:line")) == "480" and eff_p(p, attr("w:spacing", "w:lineRule")) == "exact"
    assert eff_p(p, attr("w:ind", "w:firstLine")) == "720" and eff_p(p, attr("w:jc", "w:val")) == "left"
    title = para_by_text(doc, "스마트 제조 연구")
    assert eff_r(title, title.runs[0], attr("w:sz", "w:val")) == "24" and eff_r(title, title.runs[0], flag("w:b"))
    assert eff_p(title, attr("w:jc", "w:val")) == "center"
    h1, h2, h3 = (para_by_text(doc, t) for t in ("2. 본론", "1.1 연구 배경", "1.1.1 세부 항목"))
    assert eff_p(h1, attr("w:jc", "w:val")) == "center" and eff_r(h1, h1.runs[0], flag("w:b"))
    assert eff_p(h2, attr("w:jc", "w:val")) == "left" and eff_r(h2, h2.runs[0], flag("w:b"))
    assert eff_r(h3, h3.runs[0], flag("w:b")) and eff_r(h3, h3.runs[0], flag("w:i"))
    head = para_by_text(doc, "참고문헌")
    assert eff_p(head, flag("w:pageBreakBefore")) and eff_p(head, attr("w:jc", "w:val")) == "center"
    bib = next(p for p in doc.paragraphs if p.text.startswith("Vaswani"))
    assert eff_p(bib, attr("w:ind", "w:hanging")) == "720" and eff_p(bib, attr("w:spacing", "w:line")) == "480"
    sect = sections[0][0]
    assert sect.find(qn("w:pgNumType")).get(qn("w:start")) == "1"
    head_xml = footer_xml(doc, sect, "header")
    assert 'w:instr=" PAGE "' in head_xml and '<w:jc w:val="right"/>' in head_xml


def test_ac42_apa_hwpx():
    h = Hwpx(writer.to_hwpx(BLOCKS, META, fmt("apa7-student")))
    assert len(h.sections) == 1
    _, m = h.page_margin(0)
    assert {k: m[k] for k in ("top", "header", "bottom", "footer", "left", "right")} == {
        "top": 3600, "header": 3600, "bottom": 3600, "footer": 3600, "left": 7200, "right": 7200}
    p = h.para_with("둘째 장 본문.")
    char = h.char_of(p)
    assert h.font_names(char)["latin"] == "Times New Roman" and char.get("height") == "1200"
    assert h.line(p) == ("PERCENT", 200)
    assert h.sections[0].find(f".//{HP}header//{HP}pageNum").get("pos") == "TOP_RIGHT"


# =================================================================== H. 내 양식 적용
def test_ac43_user_copy_applies(client):
    made = client.post("/api/doc-formats", json={"base": "inha-mie-thesis", "name": "12pt"}).json()
    data = made["data"]
    data["body"].update(size_pt=12, line_spacing_pct=200)
    client.patch(f"/api/doc-formats/{made['id']}", json={"data": data})
    req = {"blocks": BLOCKS, "meta": META, "doc_format": made["id"], "cover": COVER}
    doc = Document(io.BytesIO(client.post("/api/export-document", json=dict(req, format="docx")).content))
    p = para_by_text(doc, "둘째 장 본문.")
    assert eff_r(p, p.runs[0], attr("w:sz", "w:val")) == "24" and eff_p(p, attr("w:spacing", "w:line")) == "480"
    h = Hwpx(client.post("/api/export-document", json=dict(req, format="hwpx")).content)
    p = h.para_with("둘째 장 본문.")
    assert h.char_of(p).get("height") == "1200" and h.line(p) == ("PERCENT", 200)
    assert client.post("/api/export-document", json=dict(req, format="docx", doc_format="user-999")).status_code == 400


def test_ac44_no_cover_and_document_numbering():
    f = fmt("inha-mie-thesis", cover__kind="none", page_number__start="document")
    doc = Document(io.BytesIO(writer.to_docx(BLOCKS, META, f, COVER)))
    sections = docx_sections(doc)
    assert len(sections) == 1 and sections[0][1][0].text == "국문 초록: 첫 장 앞의 본문 문단입니다."
    assert sections[0][0].find(qn("w:pgNumType")).get(qn("w:start")) == "1"
    h = Hwpx(writer.to_hwpx(BLOCKS, META, f, COVER))
    assert len(h.sections) == 1 and h.sections[0].find(f".//{HP}footer") is not None


def test_ac45_cover_warnings(client):
    req = {"format": "docx", "blocks": BLOCKS, "meta": META, "doc_format": "inha-mie-thesis"}
    r = client.post("/api/export-document", json=req)
    assert r.status_code == 200
    warnings = json.loads(unquote(r.headers["X-PaperLab-Warnings"]))
    assert warnings == ["cover-missing:name,department,graduation,title_en,degree_field"]
    texts = [p.text for p in Document(io.BytesIO(r.content)).paragraphs]
    assert "○○○" in texts and "○○○학과" in texts and "○○○○년 ○월" in texts
    assert "○○석사학위 논문" in texts and "이 논문을 ○○석사학위 논문으로 제출함" in texts
    r = client.post("/api/export-document", json=dict(req, cover=dict(COVER, name="", department="")))
    assert json.loads(unquote(r.headers["X-PaperLab-Warnings"])) == ["cover-missing:name,department"]
    long_title = "스마트 제조 공정의 품질 예측을 위한 딥러닝 기반 이상 탐지 모델 연구 " * 7
    r = client.post("/api/export-document", json=dict(req, format="hwpx", cover=dict(COVER, title_ko=long_title)))
    warnings = json.loads(unquote(r.headers["X-PaperLab-Warnings"]))
    assert "cover-overflow:front" in warnings and "cover-overflow:inner" in warnings
    r = client.post("/api/export-document", json=dict(req, cover=dict(COVER, title_en="")))
    assert json.loads(unquote(r.headers["X-PaperLab-Warnings"])) == ["cover-missing:title_en"]
    front = [p.text for p in docx_sections(Document(io.BytesIO(r.content)))[0][1]]
    assert front[2] == "○○○"  # 영문 제목 자리표시
    ok = client.post("/api/export-document", json=dict(req, cover=COVER))
    assert "X-PaperLab-Warnings" not in ok.headers
    bad = client.post("/api/export-document", json=dict(req, cover=dict(COVER, graduation="2027-13")))
    assert bad.status_code == 400


def test_ac45a_page_number_distance(client):
    made = client.post("/api/doc-formats", json={"base": "inha-mie-report", "name": "11mm"}).json()
    data = made["data"]
    data["page_number"]["distance_mm"] = 11
    client.patch(f"/api/doc-formats/{made['id']}", json={"data": data})
    req = {"blocks": BLOCKS, "meta": META, "doc_format": made["id"], "cover": COVER}
    doc = Document(io.BytesIO(client.post("/api/export-document", json=dict(req, format="docx")).content))
    sections = docx_sections(doc)
    s = doc.sections[-1]
    assert abs(s.footer_distance.mm - 11) < 0.1 and abs(s.bottom_margin.mm - 55) < 0.1
    assert abs(s.top_margin.mm - 55) < 0.1 and abs(s.header_distance.mm - 40) < 0.1
    # 쪽 번호 없는 '첫 장 앞' 구역은 바꾸지 않는다
    pre = sections[-2][0].find(qn("w:pgMar"))
    assert abs(tw_mm(pre.get(qn("w:footer"))) - 40) < 0.1
    # 한글: 아래쪽 = 11mm, 꼬리말 = 나머지 44mm → 본문 영역(55mm)은 그대로, 꼬리말 글이 아래 끝에서 11mm
    h = Hwpx(client.post("/api/export-document", json=dict(req, format="hwpx")).content)
    _, m = h.page_margin(len(h.sections) - 1)
    assert abs(m["bottom"] - df.mm_to_hwp(11)) <= 1 and abs(m["footer"] - df.mm_to_hwp(44)) <= 1
    assert abs(m["top"] - 11339) <= 28 and abs(m["header"] - 4252) <= 28
    data["page_number"]["distance_mm"] = None
    client.patch(f"/api/doc-formats/{made['id']}", json={"data": data})
    doc = Document(io.BytesIO(client.post("/api/export-document", json=dict(req, format="docx")).content))
    assert abs(doc.sections[-1].footer_distance.mm - 40) < 0.1
    # 머리말 쪽 번호(APA 복사본)는 머리글 거리를 바꾼다
    apa = fmt("apa7-student", page_number__distance_mm=8)
    doc = Document(io.BytesIO(writer.to_docx(BLOCKS, META, apa)))
    assert abs(doc.sections[-1].header_distance.mm - 8) < 0.1 and abs(doc.sections[-1].top_margin.mm - 25.4) < 0.1


# =================================================================== I. 양식 파일 가져오기
def _docx_fixture(headings: bool = True, line_rule: str = "auto") -> bytes:
    from docx.oxml import OxmlElement
    from docx.shared import Mm, Pt

    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(188), Mm(257)
    s.top_margin, s.header_distance, s.bottom_margin, s.footer_distance = Mm(20), Mm(15), Mm(29), Mm(21)
    s.left_margin, s.right_margin = Mm(24), Mm(21)
    normal = d.styles["Normal"]
    normal.font.name = "휴먼명조"
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "휴먼명조")
    normal.font.size = Pt(11)
    if line_rule == "auto":
        normal.paragraph_format.line_spacing = 1.8  # line=432 lineRule=auto
    else:
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:line"), "396")  # 11pt × 180% = 19.8pt
        sp.set(qn("w:lineRule"), line_rule)
        normal.element.get_or_add_pPr().append(sp)
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    ind = OxmlElement("w:ind")
    ind.set(qn("w:firstLineChars"), "300")
    normal.element.get_or_add_pPr().append(ind)
    if headings:
        for name, size in (("Heading 1", 16), ("Heading 2", 13)):
            d.styles[name].font.size = Pt(size)
            d.styles[name].font.bold = True
    else:
        for st in list(d.styles):
            if st.name.lower().startswith("heading"):
                st.element.getparent().remove(st.element)
    d.add_paragraph("본문")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _as_dotx(docx: bytes) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(docx))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for item in src.infolist():
            data = src.read(item)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"document.main+xml", b"template.main+xml")
            z.writestr(item, data)
    return out.getvalue()


EXPECTED_PAGE = {"width_mm": 188, "height_mm": 257, "margin_mm": {"top": 15, "bottom": 21, "left": 24, "right": 21},
                 "header_mm": 5, "footer_mm": 8, "gutter_mm": 0}


def _check_imported(res):
    data = res["data"]
    assert data["page"] == EXPECTED_PAGE
    assert data["fonts"]["hangul"] == "휴먼명조"
    body = data["body"]
    assert (body["size_pt"], body["line_spacing_pct"], body["first_line_indent"], body["align"]) == (
        11, 180, {"value": 3, "unit": "ch"}, "justify")
    for path in ("page.width_mm", "page.height_mm", "page.margin_mm.top", "page.header_mm", "fonts.hangul",
                 "body.size_pt", "body.line_spacing_pct", "body.first_line_indent", "body.align"):
        assert path in res["found"], path


def test_ac46_ac47_import_docx_and_dotx(client):
    for name, raw in (("학위논문양식.docx", _docx_fixture()), ("학위논문양식.dotx", _as_dotx(_docx_fixture()))):
        r = client.post("/api/doc-formats/import", files={"file": (name, raw, "application/octet-stream")})
        assert r.status_code == 200, r.text
        res = r.json()
        assert res["source"] == {"kind": "docx", "filename": name} and res["suggested_name"] == "학위논문양식"
        assert res["base"] == "default"
        _check_imported(res)
        h = res["data"]["headings"]
        assert (h["1"]["size_pt"], h["1"]["bold"], h["2"]["size_pt"], h["2"]["bold"]) == (16, True, 13, True)
        assert {"headings.1.size_pt", "headings.1.bold", "headings.2.size_pt", "headings.2.bold"} <= set(res["found"])


def _hwpx_fixture() -> bytes:
    from hwpx._document import layout

    d = HwpxDocument.new()
    layout.set_page_setup(d, width_mm=188, height_mm=257, margins_mm={
        "top": 15, "header": 5, "bottom": 21, "footer": 8, "left": 24, "right": 21})
    header = d.oxml.headers[0]
    char = d.oxml.ensure_run_style(font="휴먼명조", size=11)
    para = header.ensure_paragraph_format(base_para_pr_id="0", line_spacing_percent=180, margins={"intent": 3300})
    for style in header.element.iter(f"{HH}style"):
        if style.get("engName") == "Normal":
            style.set("charPrIDRef", char)
            style.set("paraPrIDRef", para)
    header.mark_dirty()
    d.add_paragraph("본문")
    return d.to_bytes()


def test_ac46a_line_spacing_rules():
    auto = format_import.import_format("a.docx", _docx_fixture())
    assert auto["data"]["body"]["line_spacing_pct"] == 180
    assert any("Normal" in w and "'배수'라 %로 바꿔 근사" in w for w in auto["warnings"])
    exact = format_import.import_format("a.docx", _docx_fixture(line_rule="exact"))
    assert exact["data"]["body"]["line_spacing_pct"] == 180
    assert not any("Normal" in w for w in exact["warnings"])
    least = format_import.import_format("a.docx", _docx_fixture(line_rule="atLeast"))
    assert least["data"]["body"]["line_spacing_pct"] == 180
    assert any("Normal" in w and "'최소'라" in w for w in least["warnings"])
    # 우리가 내보낸 파일(고정값)은 경고 없이 역변환된다
    assert not any("줄간격" in w for w in format_import.import_format("t.docx", thesis_docx())["warnings"])


def test_ac48_import_hwpx():
    res = format_import.import_format("양식.hwpx", _hwpx_fixture())
    assert res["source"]["kind"] == "hwpx"
    data = res["data"]
    assert data["page"] == EXPECTED_PAGE
    assert data["fonts"] == {"hangul": "휴먼명조", "latin": "휴먼명조", "hanja": "휴먼명조"}
    assert (data["body"]["size_pt"], data["body"]["line_spacing_pct"], data["body"]["first_line_indent"]) == (
        11, 180, {"value": 3, "unit": "ch"})
    assert "title" in res["missing"] and "quote" in res["missing"]


def test_import_roundtrip_of_own_export():
    res = format_import.import_format("t.docx", thesis_docx())
    d = res["data"]
    assert d["page"] == fmt("inha-mie-thesis")["page"]
    assert d["body"]["line_spacing_pct"] == 180 and d["body"]["first_line_indent"] == {"value": 3, "unit": "ch"}
    assert d["headings"]["1"]["size_pt"] == 16 and d["headings"]["1"]["page_break_before"]
    assert d["quote"]["indent_left"] == {"value": 14.1, "unit": "mm"}


def test_ac49_no_heading_styles_keep_base():
    res = format_import.import_format("x.docx", _docx_fixture(headings=False), base="apa7-student")
    assert res["base"] == "apa7-student"
    assert res["data"]["headings"] == fmt("apa7-student")["headings"]
    assert {"headings.1", "headings.2", "headings.3"} <= set(res["missing"])


def test_ac50_import_does_not_save(client):
    before = len(client.get("/api/doc-formats").json())
    res = client.post("/api/doc-formats/import", files={"file": ("a.docx", _docx_fixture(), "x")},
                      data={"base": "inha-mie-thesis"}).json()
    assert res["base"] == "inha-mie-thesis" and len(client.get("/api/doc-formats").json()) == before
    made = client.post("/api/doc-formats", json={"base": res["base"], "name": res["suggested_name"],
                                                 "data": res["data"]}).json()
    assert made["data"]["page"] == EXPECTED_PAGE
    assert len(client.get("/api/doc-formats").json()) == before + 1


@pytest.mark.parametrize("name, raw", [
    ("a.doc", b"\xd0\xcf\x11\xe0" + b"0" * 100),
    ("a.hwp", b"\xd0\xcf\x11\xe0" + b"0" * 100),
    ("a.txt", b"hello"),
    ("a.docx", b"PK\x03\x04 broken zip"),
    ("a.hwpx", b"not a zip at all"),
])
def test_ac51_bad_files(client, name, raw):
    r = client.post("/api/doc-formats/import", files={"file": (name, raw, "x")})
    assert r.status_code == 400 and re.search(r"[가-힣]", r.json()["detail"])


def test_l5_size_checked_before_reading_everything(client):
    big = b"0" * (21 * 1024 * 1024)
    r = client.post("/api/doc-formats/import", content=big,
                    headers={"Content-Type": "multipart/form-data; boundary=x"})
    assert r.status_code == 400 and "20MB" in r.json()["detail"]


def test_l5_zip_bomb_rejected_before_unpacking(client, monkeypatch):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", b"\0" * (60 * 1024 * 1024))
    raw = buf.getvalue()
    assert len(raw) < 1024 * 1024
    called = []
    monkeypatch.setattr(zipfile.ZipFile, "testzip", lambda self: called.append(1))
    r = client.post("/api/doc-formats/import", files={"file": ("bomb.docx", raw, "x")})
    assert r.status_code == 400 and "너무 커요" in r.json()["detail"] and not called


def test_ac51_too_big(client):
    r = client.post("/api/doc-formats/import", files={"file": ("a.docx", b"0" * (20 * 1024 * 1024 + 10), "x")})
    assert r.status_code == 400 and "20MB" in r.json()["detail"]
    assert client.post("/api/doc-formats/import", files={"file": ("a.docx", _docx_fixture(), "x")},
                       data={"base": "user-1"}).status_code == 400


# =================================================================== 표지 문구 규칙 (6.1)
def test_cover_text_rules():
    assert df.spaced("김인하") == "김 인 하" and df.spaced("John Kim") == "John Kim" and df.spaced("남궁민수아이") == "남궁민수아이"
    assert df.ym_text("2027-02") == "2027년 2월" and df.ym_text("") == "○○○○년 ○월"
    assert df.approval_ym({"graduation": "2027-02"}) == "2026-12"
    assert df.approval_ym({"graduation": "2027-08"}) == "2027-06"
    assert df.approval_ym({"graduation": "2027-05"}) == ""
    pages = df.cover_pages("thesis", {"degree_field": "", "include": {"front": True, "inner": False,
                                                                       "approval": False}}, "")
    assert [p["key"] for p in pages] == ["front"] and pages[0]["lines"][0]["text"] == "○○석사학위 논문"
    assert pages[0]["lines"][1]["text"] == "○○○"
    assert df.missing_cover_fields(df.cover_with_defaults({}), "report") == [
        "name", "department", "graduation", "title_en"]  # 보고서는 학위명을 쓰지 않는다
    assert df.cover_pages("none", COVER) == []


@pytest.mark.parametrize("bad", [{"id": "default"}, ["user-1"], 3, True])
def test_non_string_format_id_is_400(client, bad):
    mid = client.post("/api/manuscripts", json={}).json()["id"]
    assert client.put("/api/settings", json={"doc_format_default": bad}).status_code == 400
    r = client.post("/api/export-document", json={"format": "docx", "blocks": BLOCKS, "meta": META, "doc_format": bad})
    assert r.status_code == 400
    assert client.patch(f"/api/manuscripts/{mid}", json={"doc_format": bad}).status_code == 400
    assert client.post("/api/manuscripts", json={"doc_format": bad}).status_code == 400
    assert client.post("/api/doc-formats", json={"base": bad, "name": "x"}).status_code == 400
    assert client.get("/api/settings").json()["doc_format_default"] == "default"
