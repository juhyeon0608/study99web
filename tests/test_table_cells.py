"""표 칸 안 줄바꿈 · 머리글 가운데 (2026-10-26 사용자 요청) — 워드 · 한글 · 마크다운 내보내기."""

import io
import zipfile

from paperlab import writer

BLOCK = {"type": "table", "header": True, "align": ["left", "right"], "caption": [{"text": "표 1. 시험"}],
         "rows": [[[{"text": "구분"}], [{"text": "AS-IS"}]],
                  [[{"text": "데이터 보안"}], [{"text": "1) 개인 PC 보관\n2) 공유 폴더"}]]]}


def test_cell_align_header_center_body_column():
    assert writer._cell_align(BLOCK, 0, 0) == "center" and writer._cell_align(BLOCK, 0, 1) == "center"
    assert writer._cell_align(BLOCK, 1, 0) == "left" and writer._cell_align(BLOCK, 1, 1) == "right"
    assert writer._cell_align(dict(BLOCK, header=False), 0, 1) == "right"


def test_docx_cell_line_break_and_header_center():
    from docx import Document

    doc = Document(io.BytesIO(writer.to_docx([BLOCK], {"title": "t"})))
    t = doc.tables[0]
    head = t.rows[0].cells[0].paragraphs[0]
    body = t.rows[1].cells[1].paragraphs[0]
    assert head.alignment is not None and int(head.alignment) == 1  # CENTER
    assert int(body.alignment) == 2  # RIGHT
    assert "<w:br/>" in body._p.xml and body.text == "1) 개인 PC 보관\n2) 공유 폴더"


def test_hwpx_cell_line_break():
    data = writer.to_hwpx([BLOCK], {"title": "t"})
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read("Contents/section0.xml").decode("utf-8")
    assert "lineBreak" in xml and "1) 개인 PC 보관" in xml and "2) 공유 폴더" in xml


def test_markdown_cell_line_break():
    text = writer.to_markdown([BLOCK])
    assert "1) 개인 PC 보관<br>2) 공유 폴더" in text
