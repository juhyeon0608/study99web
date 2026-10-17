"""내보내기 목차 · 그림 목차 · 표 목차 · 각주 (docs/specs/export-lists-footnotes.md 7장 수용 기준)

- 표지 양식(학위논문 · 대체 보고서)이면 인정서(인준서) 뒤, 본문 앞에 '목 차' · '그림 목차' · '표 목차' 쪽
- 워드: 진짜 TOC 필드(\\o "1-2" · \\t "Figure Caption,1") + 어림 쪽 번호로 채운 결과 + 열 때 필드 업데이트
- 한글: 정적 목록(오른쪽 끝 점선 탭 + 쪽 번호), 문서 검증 통과
- 표지 정보 include.toc · lof · lot · toc_h4로 켜고 끄기 / 표지 없는 양식에는 넣지 않음
- 각주: 워드 word/footnotes.xml · 한글 hp:footNote · 마크다운 [^n]
- 쪽 번호 어림값이 LibreOffice로 그린 실제 쪽과 맞는지 (soffice가 있을 때만)
"""
import io
import json
import os
import re
import shutil
import subprocess
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from hwpx import HwpxDocument

from paperlab import doc_formats as df
from paperlab import doc_lists, writer

GOLDEN = Path(__file__).parent / "golden"
META = {"title": "스마트 제조 연구"}
COVER = {"degree": "master", "degree_field": "공학", "title_en": "A Study", "graduation": "2027-02",
         "department": "스마트제조공학과", "name": "김인하", "spaced_name": True, "advisors": ["홍길동"],
         "committee": ["", "", ""]}
HP = "{http://www.hancom.co.kr/hwpml/2011/paragraph}"
HH = "{http://www.hancom.co.kr/hwpml/2011/head}"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
SOFFICE = next((p for p in (shutil.which("soffice"), r"C:\Program Files\LibreOffice\program\soffice.exe")
                if p and os.path.isfile(p)), None)


def fmt(fid: str) -> dict:
    return df.normalize(df.builtin_data(fid), fid)


def para(text, **kw):
    return {"type": "paragraph", "runs": [{"text": text}], **kw}


def heading(level, text, *extra):
    return {"type": "heading", "level": level, "runs": [{"text": text}, *extra]}


def blocks(chapters: int = 2, sections: int = 2, paras: int = 2) -> list[dict]:
    """장 · 절 · 항 · 본문 · 그림(그림 없음) · 표 · 각주 · 참고문헌이 든 원고 블록"""
    out = [{"type": "title", "runs": [{"text": "스마트 제조 연구"}]}, para("국문 초록: 첫 장 앞의 본문 문단입니다.")]
    for c in range(1, chapters + 1):
        out.append(heading(1, f"제{c}장 장 제목 {c}"))
        for s in range(1, sections + 1):
            out.append(heading(2, f"제{s}절 절 제목 {c}-{s}"))
            out.append(heading(3, f"{s}. 항 제목 {c}-{s}"))
            for k in range(paras):
                runs = [{"text": "본문 문장입니다. " * 30}]
                if k == 0:
                    runs.append({"footnote": [{"text": f"각주 {c}-{s} "}, {"text": "기울임", "i": True}]})
                out.append({"type": "paragraph", "runs": runs})
            out.append({"type": "figure", "src": "0123456789abcdef.png", "alt": "", "number": 1,
                        "caption": [{"text": f"그림 {c}.{s}.", "b": True}, {"text": " 실험 장치"}]})
            out.append({"type": "table", "header": True, "align": [], "number": 1,
                        "caption": [{"text": f"표 {c}.{s}.", "b": True}, {"text": " 결과"}],
                        "rows": [[[{"text": "항목"}], [{"text": "값"}]], [[{"text": "정확도"}], [{"text": "0.9"}]]]})
    out += [{"type": "page_break"}, {"type": "bib_heading", "text": "참고문헌"},
            {"type": "bib_entry", "hanging": True, "runs": [{"text": "Vaswani, A. (2017). Attention."}]}]
    return out


def zread(data: bytes, name: str) -> bytes:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return z.read(name)


def docx_lists(data: bytes) -> dict:
    """{제목: [(스타일 이름, 글, 쪽)…]} — 목록 제목 문단 다음의 toc · table of figures 스타일 문단들"""
    doc = Document(io.BytesIO(data))
    out, cur = {}, None
    for p in doc.paragraphs:
        name = p.style.name
        if p.text in doc_lists.TITLES.values():
            cur = out.setdefault(p.text, [])
        elif cur is not None and (name.startswith("toc ") or name == "table of figures"):
            text, _, page = p.text.partition("\t")
            cur.append((name, text, page))
        else:
            cur = None
    return out


def docx_body_xml(data: bytes) -> str:
    return zread(data, "word/document.xml").decode("utf-8")


# =================================================================== 워드
def test_docx_toc_field_entries_and_pages():
    data = writer.to_docx(blocks(), META, fmt("inha-mie-thesis"), COVER)
    lists = docx_lists(data)
    assert list(lists) == ["목 차", "그림 목차", "표 목차"]
    toc = lists["목 차"]
    assert [(s, t) for s, t, _ in toc] == [
        ("toc 1", "제1장 장 제목 1"), ("toc 2", "제1절 절 제목 1-1"), ("toc 2", "제2절 절 제목 1-2"),
        ("toc 1", "제2장 장 제목 2"), ("toc 2", "제1절 절 제목 2-1"), ("toc 2", "제2절 절 제목 2-2"),
        ("toc 1", "참고문헌")]  # 기본은 장 · 절까지(항 빠짐), 참고문헌 제목 포함
    pages = [int(p) for _, _, p in toc]
    assert pages[0] == 1 and pages == sorted(pages) and pages[-1] > pages[3] > 1
    assert [t for _, t, _ in lists["그림 목차"]] == ["그림 1.1. 실험 장치", "그림 1.2. 실험 장치",
                                                   "그림 2.1. 실험 장치", "그림 2.2. 실험 장치"]
    assert all(s == "table of figures" for s, _, _ in lists["표 목차"])
    xml = docx_body_xml(data)
    # 필드 세 개: 시작 · 명령 · 구분 → 결과 → 끝 (각 목록 하나씩)
    assert xml.count('w:fldCharType="begin"') == 3 and xml.count('w:fldCharType="end"') == 3
    assert ' TOC \\o "1-2" \\h \\z \\u ' in xml
    assert ' TOC \\h \\z \\t "Figure Caption,1" ' in xml and ' TOC \\h \\z \\t "Table Caption,1" ' in xml
    # 캡션 문단은 그림 · 표 목차 필드가 모으는 스타일
    doc = Document(io.BytesIO(data))
    caps = {p.text: p.style.name for p in doc.paragraphs if p.text.endswith(("실험 장치", " 결과"))
            and "\t" not in p.text}
    assert caps["그림 1.1. 실험 장치"] == "Figure Caption" and caps["표 2.2. 결과"] == "Table Caption"
    # 목차 제목은 제목 스타일이 아니다(워드가 목차를 새로 만들 때 '목 차'가 항목이 되지 않게)
    title = next(p for p in doc.paragraphs if p.text == "목 차")
    assert title.style.name == "Normal" and title.paragraph_format.keep_with_next
    assert doc.styles["Heading 1"].element.pPr.find(qn("w:outlineLvl")) is not None
    # 항목 스타일: 오른쪽 끝 점선 탭
    tab = doc.styles["toc 2"].element.pPr.find(qn("w:tabs")).find(qn("w:tab"))
    assert tab.get(qn("w:val")) == "right" and tab.get(qn("w:leader")) == "dot"
    page = fmt("inha-mie-thesis")["page"]
    assert int(tab.get(qn("w:pos"))) == df.mm_to_twip(page["width_mm"] - 24 - 21)


def test_docx_update_fields_setting_in_schema_order():
    settings = zread(writer.to_docx(blocks(1, 1), META, fmt("inha-mie-thesis"), COVER), "word/settings.xml").decode()
    assert '<w:updateFields w:val="true"/>' in settings
    assert settings.index("<w:updateFields") < settings.index("<w:compat")
    off = writer.to_docx(blocks(1, 1), META, fmt("inha-mie-thesis"),
                         dict(COVER, include={"toc": False, "lof": False, "lot": False}))
    assert b"updateFields" not in zread(off, "word/settings.xml")
    assert docx_lists(off) == {} and b"TOC \\" not in zread(off, "word/document.xml")


def test_lists_sections_and_page_numbers():
    """목록은 표지류 뒤 쪽 번호 없는 구역(첫 장 앞 부분과 함께), 본문 구역은 그대로 1쪽부터"""
    from tests.test_doc_formats import docx_sections, footer_xml

    doc = Document(io.BytesIO(writer.to_docx(blocks(1, 1), META, fmt("inha-mie-report"), COVER)))
    sections = docx_sections(doc)
    assert len(sections) == 5
    sect, ps = sections[3]
    texts = [p.text for p in ps]
    assert texts[0] == "목 차" and "국문 초록: 첫 장 앞의 본문 문단입니다." in texts
    assert texts.index("그림 목차") < texts.index("표 목차") < texts.index("국문 초록: 첫 장 앞의 본문 문단입니다.")
    for name in ("그림 목차", "표 목차"):  # 목록마다 새 쪽
        assert next(p for p in ps if p.text == name).paragraph_format.page_break_before
    assert "PAGE" not in footer_xml(doc, sect)
    sect, ps = sections[4]
    assert ps[0].text == "제1장 장 제목 1" and sect.find(qn("w:pgNumType")).get(qn("w:start")) == "1"


def test_lists_options_and_no_cover_formats():
    b = blocks(1, 1)
    deep = docx_lists(writer.to_docx(b, META, fmt("inha-mie-thesis"), dict(COVER, include={"toc_h4": True})))
    assert ("toc 3", "1. 항 제목 1-1") in [(s, t) for s, t, _ in deep["목 차"]]
    assert ' TOC \\o "1-3" ' in docx_body_xml(writer.to_docx(b, META, fmt("inha-mie-thesis"),
                                                               dict(COVER, include={"toc_h4": True})))
    only_toc = docx_lists(writer.to_docx(b, META, fmt("inha-mie-thesis"), dict(COVER, include={"lof": False,
                                                                                               "lot": False})))
    assert list(only_toc) == ["목 차"]
    # 그림 · 표가 없으면 그 목록은 빠진다
    plain = [x for x in b if x["type"] not in ("figure", "table")]
    assert list(docx_lists(writer.to_docx(plain, META, fmt("inha-mie-thesis"), COVER))) == ["목 차"]
    # 표지 없는 양식(APA · 표지 '없음')과 기본 (A4) 경로에는 넣지 않는다
    for data in (writer.to_docx(b, META, fmt("apa7-student"), COVER), writer.to_docx(b, META)):
        assert docx_lists(data) == {} and b"TOC \\" not in zread(data, "word/document.xml")
    none = df.normalize(dict(df.builtin_data("inha-mie-thesis"), cover={"kind": "none"}), "inha-mie-thesis")
    assert docx_lists(writer.to_docx(b, META, none, COVER)) == {}
    # 표지 정보를 보내지 않아도(None) 기본값으로 목차를 넣는다
    assert "목 차" in docx_lists(writer.to_docx(b, META, fmt("inha-mie-thesis"), None))


def test_docx_footnotes_real():
    data = writer.to_docx(blocks(1, 2), META, fmt("inha-mie-thesis"), COVER)
    notes = zread(data, "word/footnotes.xml").decode()
    assert notes.count("<w:footnote ") == 4  # 구분선 2 + 각주 2
    assert "각주 1-1 " in notes and "<w:i/>" in notes
    assert docx_body_xml(data).count("<w:footnoteReference ") == 2
    # 목차 항목(제목 글)에는 각주 번호가 따라 들어가지 않는다
    b = blocks(1, 1)
    b[2]["runs"].append({"footnote": [{"text": "제목 각주"}]})
    toc = docx_lists(writer.to_docx(b, META, fmt("inha-mie-thesis"), COVER))["목 차"]
    assert toc[0][1] == "제1장 장 제목 1"


# =================================================================== 한글
def hwpx_lists(data: bytes) -> dict:
    z = zipfile.ZipFile(io.BytesIO(data))
    names = sorted((n for n in z.namelist() if re.fullmatch(r"Contents/section\d+\.xml", n)),
                   key=lambda n: int(re.findall(r"\d+", n)[0]))
    out, cur = {}, None
    for n in names:
        for p in ET.fromstring(z.read(n)).findall(f"{HP}p"):
            ts = [t for r in p.findall(f"{HP}run") for t in r.findall(f"{HP}t")]
            text = "".join(t.text or "" for t in ts)
            tab = next((x for t in ts for x in t.findall(f"{HP}tab")), None)
            if text in doc_lists.TITLES.values():
                cur = out.setdefault(text, [])
            elif cur is not None and tab is not None:
                cur.append((text, tab.tail or "", tab.get("leader"), tab.get("type"), p.get("paraPrIDRef")))
            else:
                cur = None
    return out


def test_hwpx_static_lists_with_leader_tab():
    data = writer.to_hwpx(blocks(), META, fmt("inha-mie-thesis"), COVER)
    assert HwpxDocument.open(io.BytesIO(data)).validate().issues == ()
    lists = hwpx_lists(data)
    assert list(lists) == ["목 차", "그림 목차", "표 목차"]
    toc = lists["목 차"]
    full = df.FULL_SPACE
    assert [t for t, *_ in toc] == ["제1장 장 제목 1", f"{full * 2}제1절 절 제목 1-1", f"{full * 2}제2절 절 제목 1-2",
                                    "제2장 장 제목 2", f"{full * 2}제1절 절 제목 2-1", f"{full * 2}제2절 절 제목 2-2",
                                    "참고문헌"]
    assert all(leader == "3" and kind == "2" for _, _, leader, kind, _ in toc)
    # 쪽 번호는 워드와 같은 어림값
    docx_pages = [p for _, _, p in docx_lists(writer.to_docx(blocks(), META, fmt("inha-mie-thesis"), COVER))["목 차"]]
    assert [p for _, p, *_ in toc] == docx_pages
    # 문단 모양: 오른쪽 끝(본문 폭) 탭 + 점선(DASH — 한컴 차례와 같음)
    header = ET.fromstring(zread(data, "Contents/header.xml"))
    para_pr = next(x for x in header.iter(f"{HH}paraPr") if x.get("id") == toc[0][4])
    tab_pr = next(x for x in header.iter(f"{HH}tabPr") if x.get("id") == para_pr.get("tabPrIDRef"))
    item = next(tab_pr.iter(f"{HH}tabItem"))
    page = fmt("inha-mie-thesis")["page"]
    assert item.get("type") == "RIGHT" and item.get("leader") == "DASH"
    assert abs(int(item.get("pos")) - df.mm_to_hwp(page["width_mm"] - 24 - 21)) <= 2


def test_hwpx_lists_off_and_footnotes():
    off = writer.to_hwpx(blocks(1, 1), META, fmt("inha-mie-report"), dict(COVER, include={"toc": False, "lof": False,
                                                                                            "lot": False}))
    assert hwpx_lists(off) == {}
    data = writer.to_hwpx(blocks(1, 2), META, fmt("inha-mie-report"), COVER)
    sec = zread(data, "Contents/section4.xml").decode()
    assert len(re.findall(r"<hp:footNote[ >]", sec)) == 2 and "각주 1-2 기울임" in sec


# =================================================================== 골든 (목록 부분만)
def _slice(xml: str, open_re: str, close: str, first: str, last: str) -> str:
    """first가 든 문단 시작부터 마지막 last가 든 문단 끝까지"""
    at = xml.index(first)
    start = [m.start() for m in re.finditer(open_re, xml[:at])][-1]
    end = xml.index(close, xml.rindex(last)) + len(close)
    return xml[start:end]


def _docx_lists_xml(data: bytes) -> str:
    """'목 차' 문단부터 마지막 목록 필드 끝 문단까지 document.xml 조각"""
    return _slice(docx_body_xml(data), r"<w:p[ >]", "</w:p>", ">목 차<", 'w:fldCharType="end"')


def _hwpx_lists_xml(data: bytes) -> str:
    """'목 차' 문단부터 마지막 목록 줄까지 section3.xml 조각 (문단 id는 지움 — 실행마다 다름)"""
    sec = zread(data, "Contents/section3.xml").decode("utf-8")
    sec = re.sub(r'(<hp:p\b[^>]*?)\sid="[^"]*"', r"\1", sec)
    return _slice(sec, r"<hp:p[ >]", "</hp:p>", ">목 차<", "<hp:tab ")


@pytest.mark.parametrize("kind", ["docx", "hwpx"])
def test_lists_golden(kind):
    b = blocks()
    if kind == "docx":
        got = _docx_lists_xml(writer.to_docx(b, META, fmt("inha-mie-thesis"), COVER))
    else:
        got = _hwpx_lists_xml(writer.to_hwpx(b, META, fmt("inha-mie-thesis"), COVER))
    path = GOLDEN / f"lists_{kind}.xml"
    if os.environ.get("PAPERLAB_UPDATE_GOLDEN") == "1":
        path.write_text(got, encoding="utf-8")
    assert got == path.read_text(encoding="utf-8")


# =================================================================== 마크다운 · API
def test_markdown_keeps_footnotes():
    md = writer.to_markdown([para("본문"), {"type": "paragraph", "runs": [
        {"text": "가"}, {"footnote": [{"text": "첫 "}, {"text": "굵게", "b": True}]}, {"text": "나"},
        {"footnote": [{"text": "둘째"}]}]}])
    assert "가[^1]나[^2]" in md and md.rstrip().endswith("[^1]: 첫 **굵게**\n[^2]: 둘째")


def test_export_api_respects_cover_include(local_client):
    req = {"format": "docx", "blocks": blocks(1, 1), "meta": META, "doc_format": "inha-mie-thesis"}
    on = local_client.post("/api/export-document", json=dict(req, cover=COVER))
    assert on.status_code == 200 and list(docx_lists(on.content)) == ["목 차", "그림 목차", "표 목차"]
    off = local_client.post("/api/export-document", json=dict(req, cover=dict(COVER, include={
        "front": True, "inner": True, "approval": True, "toc": False, "lof": True, "lot": False, "toc_h4": False})))
    assert list(docx_lists(off.content)) == ["그림 목차"]
    bad = local_client.post("/api/export-document", json=dict(req, cover=dict(COVER, include={"toc": "yes"})))
    assert bad.status_code == 400 and "cover.include.toc" in bad.json()["detail"]
    md = local_client.post("/api/export-document", json=dict(req, format="md", cover=COVER))
    assert md.status_code == 200 and b"[^1]: " in md.content and "목 차".encode() not in md.content


def test_validate_cover_include_keys():
    out = df.validate_cover({"include": {"toc": False, "lof": True, "lot": False, "toc_h4": True, "zzz": True}})
    assert out == {"include": {"toc": False, "lof": True, "lot": False, "toc_h4": True}}
    assert df.cover_with_defaults({})["include"] == {"front": True, "inner": True, "approval": True, "toc": True,
                                                     "lof": True, "lot": True, "toc_h4": False}
    with pytest.raises(df.FormatError):
        df.validate_cover({"include": {"toc_h4": 1}})


def test_estimate_pages_basic_rules():
    f = fmt("inha-mie-thesis")
    b = [heading(1, "제1장 가"), para("짧은 글"), heading(1, "제2장 나"), {"type": "page_break"}, para("다")]
    pages = doc_lists.estimate_pages([(b, True)], f)
    # 장 제목(제목 1수준)은 새 쪽에서 시작, page_break는 다음 글을 새 쪽으로
    assert [pages[id(x)] for x in b if x["type"] != "page_break"] == [1, 1, 2, 3]
    assert doc_lists.estimate_pages([(b, False)], f) == {}
    tail = para("끝")
    got = doc_lists.estimate_pages([([para("가" * 4000), tail], True)], f)
    assert got[id(tail)] >= 3  # 긴 문단은 여러 쪽에 걸친다


# =================================================================== LibreOffice로 실제 쪽 확인
def _pdf_texts(pdf: Path) -> list[str]:
    import pymupdf

    with pymupdf.open(pdf) as doc:
        return [page.get_text() for page in doc]


@pytest.mark.skipif(SOFFICE is None, reason="LibreOffice(soffice)가 없어 실제 쪽 비교를 건너뜀")
def test_estimated_pages_match_libreoffice(tmp_path):
    """워드 파일을 LibreOffice로 PDF로 그려 목차 · 그림 목차 · 표 목차의 쪽 번호(어림값)가 실제 쪽과 같은지"""
    pytest.importorskip("pymupdf")
    data = writer.to_docx(blocks(3, 3, 4), META, fmt("inha-mie-thesis"), COVER)
    src = tmp_path / "lists.docx"
    src.write_bytes(data)
    profile = (tmp_path / "lo").as_uri()
    subprocess.run([SOFFICE, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf",
                    "--outdir", str(tmp_path), str(src)], check=True, capture_output=True, timeout=240)
    texts = _pdf_texts(tmp_path / "lists.pdf")
    # 본문 1쪽 = '제1장' 줄만 있는 첫 쪽 (목차 줄은 점선 · 쪽 번호가 붙어 같은 줄이 아님)
    first = next(i for i, t in enumerate(texts) if any(line.strip() == "제1장 장 제목 1" for line in t.splitlines()))
    expected = {t: int(p) for entries in docx_lists(data).values() for _, t, p in entries}
    actual = {}
    for t in expected:
        actual[t] = next((i - first + 1 for i in range(first, len(texts))
                          if any(line.strip() == t for line in texts[i].splitlines())), None)
    wrong = {t: (expected[t], actual[t]) for t in expected if expected[t] != actual[t]}
    print("어림 쪽 ≠ LibreOffice 쪽:", json.dumps(wrong, ensure_ascii=False))
    # 어림값이다: 글꼴(휴먼명조가 없으면 다른 글꼴) · 한글 줄 나눔(LibreOffice는 글자, 워드 · 한글은 어절)이 달라
    # 한 쪽 안팎으로 어긋날 수 있다 — 모두 ±1쪽 안, 80% 이상은 같은 쪽
    assert all(a is not None and abs(e - a) <= 1 for e, a in wrong.values()), wrong
    assert len(wrong) <= len(expected) // 5, json.dumps(wrong, ensure_ascii=False)


@pytest.mark.db
def test_manuscript_cover_include_saved(client):
    """표지 정보 창의 목차 체크가 원고 cover.include에 저장되고 그대로 돌아온다(마이그레이션 없음 — 있던 jsonb)"""
    mid = client.post("/api/manuscripts", json={"doc_format": "inha-mie-thesis"}).json()["id"]
    include = {"front": True, "inner": True, "approval": True, "toc": False, "lof": True, "lot": False, "toc_h4": True}
    assert client.patch(f"/api/manuscripts/{mid}", json={"cover": dict(COVER, include=include)}).status_code == 200
    assert client.get(f"/api/manuscripts/{mid}").json()["cover"]["include"] == include
    r = client.patch(f"/api/manuscripts/{mid}", json={"cover": dict(COVER, include={"lof": "no"})})
    assert r.status_code == 400 and r.json()["detail"].startswith("cover.include.lof")
