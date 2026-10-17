"""논문 쓰기 그림 · 표 · 캡션 · 목록 (docs/specs/writing-figures-tables.md)

- 그림 검사(매직 바이트 · 크기 · 머리글), 그림 키 규칙(users/{uid}/figures/{원고}/{hex16}.{png|jpg|gif})
- 워드 · 한글 · 마크다운 내보내기에 그림(실제 이미지) · 표 · 캡션 · 목록이 들어가는지 (기본 · 양식 경로)
- 내보낸 파일을 다시 읽기(인용 넣기 compose · 양식 가져오기 format_import)가 깨지지 않는지
- API: 그림 올리기 · 받기 · 원고 삭제 때 정리 · 내보내기 때 서버가 그림을 읽어 넣음 (DB 테스트)
"""

import io
import struct
import uuid
import zipfile
import zlib

import pytest
from docx import Document
from hwpx import HwpxDocument

from paperlab import compose, doc_formats as df, figures, format_import, storage as st, writer


def png(w: int = 40, h: int = 20) -> bytes:
    raw = b"".join(b"\x00" + b"\xff\x80\x00" * w for _ in range(h))

    def chunk(tag, data):
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)

    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def jpeg_header(w: int, h: int) -> bytes:
    # SOI · APP0(JFIF) · SOF0 — 크기만 읽는 시험용 (그리기 데이터 없음)
    app0 = b"\xff\xe0" + struct.pack(">H", 16) + b"JFIF\x00\x01\x01\x00\x00\x01\x00\x01\x00\x00"
    sof = b"\xff\xc0" + struct.pack(">HBHHB", 11, 8, h, w, 1) + b"\x01\x11\x00"
    return b"\xff\xd8" + app0 + sof + b"\xff\xd9"


def jpeg_variant(first: bytes, w: int = 64, h: int = 32) -> bytes:
    """JFIF · Exif 머리글 없이 시작하는 JPEG (python-docx가 못 읽음): first = SOI 바로 뒤 첫 세그먼트"""
    sof = b"\xff\xc0" + struct.pack(">HBHHB", 11, 8, h, w, 1) + b"\x01\x11\x00"
    return b"\xff\xd8" + first + sof + b"\xff\xd9"


def segment(marker: int, payload: bytes) -> bytes:
    return bytes([0xFF, marker]) + struct.pack(">H", len(payload) + 2) + payload


XMP_FIRST = jpeg_variant(segment(0xE1, b"http://ns.adobe.com/xap/1.0/\x00<x:xmpmeta/>"))  # 어도비 XMP 먼저
ADOBE_APP14 = jpeg_variant(segment(0xEE, b"Adobe\x00\x64\x00\x00\x00\x00\x02"))          # CMYK 등
DQT_FIRST = jpeg_variant(segment(0xDB, b"\x00" + bytes(range(1, 65))))                     # APP 없이 바로 DQT
ODD_JPEGS = [XMP_FIRST, ADOBE_APP14, DQT_FIRST]

GIF = b"GIF89a" + struct.pack("<HH", 3, 2) + b"\x00\x00\x00;"
IMG = png(400, 200)
NAME = "0123456789abcdef.png"
BLOCKS = [
    {"type": "title", "runs": [{"text": "그림과 표"}]},
    {"type": "heading", "level": 1, "runs": [{"text": "1. 서론"}]},
    {"type": "paragraph", "runs": [{"text": "그림 1과 표 1을 보세요."}]},
    {"type": "list_item", "ordered": False, "depth": 0, "number": 1, "runs": [{"text": "첫 항목"}]},
    {"type": "list_item", "ordered": False, "depth": 1, "number": 1, "runs": [{"text": "하위 항목"}]},
    {"type": "list_item", "ordered": True, "depth": 0, "number": 2, "runs": [{"text": "번호 항목"}]},
    {"type": "figure", "src": NAME, "alt": "실험 장치",
     "caption": [{"text": "그림 1.", "b": True}, {"text": " 실험 장치"}],
     "image": {"data": IMG, "format": "png", "width": 400, "height": 200}},
    {"type": "figure", "src": "0123456789abcdee.png", "alt": "없음", "caption": [{"text": "그림 2.", "b": True}]},
    {"type": "table", "header": True, "align": ["left", "center", "right"],
     "caption": [{"text": "표 1.", "b": True}, {"text": " 결과 요약"}, {"footnote": [{"text": "각주는 글자로"}]}],
     "rows": [[[{"text": "항목"}], [{"text": "값"}], [{"text": "비고"}]],
              [[{"text": "정확도"}], [{"text": "0.91"}], [{"text": "좋음", "i": True}]],
              [[{"text": "재현율 | 단위"}], [{"text": "0.85"}]]]},
]
THESIS = df.normalize(df.builtin_data("inha-mie-thesis"))
COVER = {"name": "홍길동", "department": "기계공학과", "graduation": "2027-02", "title_en": "T", "degree_field": "공학"}


# =================================================================== 그림 검사 · 키 규칙
def test_sniff_and_size_from_magic_bytes():
    assert figures.validate(IMG) == {"ext": "png", "content_type": "image/png", "width": 400, "height": 200,
                                     "size": len(IMG)}
    assert figures.sniff(jpeg_header(640, 480)) == "jpg"
    assert figures.image_size(jpeg_header(640, 480), "jpg") == (640, 480)
    assert figures.image_size(GIF, "gif") == (3, 2)
    for bad, msg in ((b"", "빈 파일"), (b"<svg xmlns='http://www.w3.org/2000/svg'/>", "PNG"),
                     (b"%PDF-1.7 ...", "PNG"), (b"\x89PNG\r\n\x1a\n" + b"\x00" * 30, "읽지 못했어요"),
                     (b"RIFF\x00\x00\x00\x00WEBPVP8 ", "PNG")):
        with pytest.raises(figures.FigureError, match=msg):
            figures.validate(bad)
    with pytest.raises(figures.FigureError, match="10MB"):
        figures.validate(IMG + b"\x00" * figures.MAX_FIGURE_BYTES)


def test_fit_mm_keeps_ratio_and_never_enlarges():
    assert figures.fit_mm(96, 48, 150, 200) == pytest.approx((25.4, 12.7))
    w, h = figures.fit_mm(4000, 1000, 150, 200)
    assert w == pytest.approx(150) and h == pytest.approx(37.5)
    w, h = figures.fit_mm(1000, 8000, 150, 200)
    assert h == pytest.approx(200) and w == pytest.approx(25)


def test_figure_key_rules():
    uid, other = str(uuid.uuid4()), str(uuid.uuid4())
    key = st.figure_key(uid, 7, NAME)
    assert key == f"users/{uid}/figures/7/{NAME}"
    assert st.check_user_key(uid, key) == key
    with pytest.raises(st.StorageKeyError):
        st.check_user_key(other, key)
    for bad in ("../x.png", "0123456789ABCDEF.png", "0123456789abcdef.svg", "0123456789abcdef.png\n", "a/b.png", ""):
        with pytest.raises(st.StorageKeyError):
            st.figure_key(uid, 7, bad)
    for mid in (0, -1, True, "7", 10 ** 19):
        with pytest.raises(st.StorageKeyError):
            st.figure_prefix(uid, mid)
    us = st.UserStorage(st.FakeStorage(), uid)
    us.put(us.figure_key(7, NAME), IMG, "image/png")
    us.storage.put(f"users/{uid}/figures/7/evil.html", b"x")  # 규칙 밖 키는 목록에서 빠진다
    assert us.list_figures(7) == [(key, len(IMG))]
    with pytest.raises(st.StorageKeyError):  # 그림은 서명 주소를 만들지 않는다(서버가 내려줌)
        us.sign_get(key)


def test_attach_images_reads_only_valid_names_and_ignores_client_bytes():
    seen = []

    def load(name):
        seen.append(name)
        return IMG if name == NAME else None

    blocks = [{"type": "figure", "src": NAME, "image": {"data": b"evil"}},
              {"type": "figure", "src": "../../etc/passwd", "image": {"data": b"evil"}},
              {"type": "figure", "src": "0123456789abcdee.png"},
              {"type": "paragraph", "runs": []}]
    out = figures.attach_images(blocks, load)
    assert seen == [NAME, "0123456789abcdee.png"]
    assert out[0]["image"]["data"] == IMG and out[0]["image"]["width"] == 400
    assert "image" not in out[1] and "image" not in out[2]
    assert blocks[0]["image"] == {"data": b"evil"}  # 원래 목록은 그대로


# =================================================================== 내보내기
def _docx_check(data: bytes) -> None:
    d = Document(io.BytesIO(data))
    texts = [p.text for p in d.paragraphs]
    assert "그림 1. 실험 장치" in texts and "그림 2." in texts and writer.FIGURE_MISSING in texts
    assert "• 첫 항목" in texts and "◦ 하위 항목" in texts and "2. 번호 항목" in texts
    cap = texts.index("표 1. 결과 요약각주는 글자로")
    assert d.paragraphs[cap].paragraph_format.keep_with_next or "keepNext" in d.paragraphs[cap]._p.xml
    (t,) = d.tables
    assert [[c.text for c in r.cells] for r in t.rows] == [["항목", "값", "비고"], ["정확도", "0.91", "좋음"],
                                                          ["재현율 | 단위", "0.85", ""]]
    assert all(r.bold for r in t.rows[0].cells[0].paragraphs[0].runs)
    assert "w:tblHeader" in t.rows[0]._tr.xml
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
        assert len(media) == 1 and z.read(media[0]) == IMG
        body = z.read("word/document.xml").decode()
    assert body.count("<pic:pic") == 1
    # 그림 문단은 고정 줄간격이 아니어야 잘리지 않는다 · 가운데 · 다음(캡션)과 함께
    (pic_p,) = [p for p in d.paragraphs if "pic:pic" in p._p.xml]
    ppr = pic_p._p.pPr.xml
    assert 'w:lineRule="auto"' in ppr and 'w:lineRule="exact"' not in ppr
    assert 'w:jc w:val="center"' in ppr and "w:keepNext" in ppr


def test_docx_default_export_has_figures_tables_lists():
    data = writer.to_docx(BLOCKS, {"title": "t"})
    _docx_check(data)
    d = Document(io.BytesIO(data))
    assert d.element.body[-2].tag.endswith("}p")  # 표로 끝나면 빈 문단을 붙인다


def test_docx_thesis_export_has_figures_tables_lists():
    warnings = []
    data = writer.to_docx(BLOCKS, {"title": "t"}, THESIS, COVER, warnings)
    _docx_check(data)


@pytest.mark.parametrize("fmt", [None, THESIS])
def test_hwpx_export_has_picture_table_captions(fmt):
    data = writer.to_hwpx(BLOCKS, {"title": "t"}, fmt, COVER if fmt else None)
    d = HwpxDocument.open(io.BytesIO(data))
    assert d.validate().issues == ()
    text = d.text.plain()
    for s in ("그림 1. 실험 장치", "표 1. 결과 요약", "• 첫 항목", "◦ 하위 항목", writer.FIGURE_MISSING, "정확도", "재현율 | 단위"):
        assert s in text, s
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        bins = [n for n in z.namelist() if n.startswith("BinData/")]
        assert len(bins) == 1 and z.read(bins[0]) == IMG
        section = "".join(z.read(n).decode() for n in z.namelist() if n.startswith("Contents/section"))
    assert section.count("<hp:pic") == 1 and section.count("<hp:tbl") == 1
    assert 'repeatHeader="1"' in section and 'rowCnt="3"' in section and 'colCnt="3"' in section


def test_markdown_export_keeps_figure_ref_and_table():
    md = writer.to_markdown(BLOCKS)
    assert f"![실험 장치](fig:{NAME})\n\n**그림 1.** 실험 장치" in md
    assert "**표 1.** 결과 요약" in md
    assert "| 항목 | 값 | 비고 |\n| :-- | :-: | --: |\n| 정확도 | 0.91 | *좋음* |\n| 재현율 \\| 단위 | 0.85 |  |" in md
    assert "- 첫 항목\n\n  - 하위 항목" in md
    bad = writer.to_markdown([{"type": "figure", "src": "../x.png", "alt": "a", "caption": [{"text": "그림 1."}]}])
    assert "fig:" not in bad


def test_table_rows_are_padded_and_capped():
    rows = writer._table_rows({"rows": [[[{"text": "a"}]], [], "bad", [[{"text": "b"}], [{"text": "c"}]]]})
    assert [[len(c) for c in r] for r in rows] == [[1, 0], [0, 0], [1, 1]]
    big = writer._table_rows({"rows": [[[]] * 99] * 900})
    assert len(big) == writer.MAX_TABLE_ROWS and len(big[0]) == writer.MAX_TABLE_COLS
    assert writer._table_rows({}) == [[[]]]
    # 빈 표 · 이상한 정렬 값도 내보내기가 깨지지 않음
    odd = [{"type": "table", "rows": [], "align": ["diagonal"], "caption": "x"}]
    Document(io.BytesIO(writer.to_docx(odd)))
    assert HwpxDocument.open(io.BytesIO(writer.to_hwpx(odd))).validate().issues == ()


def test_exported_files_round_trip_through_compose_and_format_import():
    blocks = BLOCKS + [{"type": "paragraph", "runs": [{"text": "인용 [@vaswani2017attention] 자리"}]},
                       {"type": "paragraph", "runs": [{"text": "[참고문헌]"}]}]
    docx = writer.to_docx(blocks, {"title": "t"})
    hwpx = writer.to_hwpx(blocks, {"title": "t"})
    scan = compose.docx_scan(docx)
    assert [c.raw for c in scan["citations"]] == ["[@vaswani2017attention]"] and scan["has_bib_marker"]
    out = Document(io.BytesIO(compose.docx_apply(docx, [{"runs": [{"text": "(V, 2017)"}]}], [[{"text": "Ref"}]],
                                                 "참고문헌", note_style=False)))
    assert len(out.tables) == 1 and any(p.text == "인용 (V, 2017) 자리" for p in out.paragraphs)
    hscan = compose.hwpx_scan(hwpx)
    assert len(hscan["citations"]) == 1
    hout = HwpxDocument.open(io.BytesIO(compose.hwpx_apply(hwpx, [{"runs": [{"text": "(V, 2017)"}]}],
                                                           [[{"text": "Ref"}]], "참고문헌")))
    assert hout.validate().issues == () and "인용 (V, 2017) 자리" in hout.text.plain()
    for name, data in (("t.docx", docx), ("t.hwpx", hwpx), ("thesis.docx", writer.to_docx(BLOCKS, {}, THESIS, COVER)),
                       ("thesis.hwpx", writer.to_hwpx(BLOCKS, {}, THESIS, COVER))):
        res = format_import.import_format(name, data)
        assert res["data"]["page"]["width_mm"] > 0


def test_export_api_without_manuscript_writes_placeholder(local_client):
    blocks = [dict(b, image={"data": "AAAA", "format": "png", "width": 1, "height": 1}) if b["type"] == "figure" else b
              for b in BLOCKS if "image" not in b] + [{"type": "figure", "src": NAME, "caption": [{"text": "그림 3."}],
                                                       "image": {"data": "evil"}}]
    r = local_client.post("/api/export-document", json={"format": "docx", "blocks": blocks})
    assert r.status_code == 200, r.text
    texts = [p.text for p in Document(io.BytesIO(r.content)).paragraphs]
    assert texts.count(writer.FIGURE_MISSING) == 2  # 화면이 보낸 image 값은 쓰지 않음
    r = local_client.post("/api/export-document", json={"format": "md", "blocks": BLOCKS[:3] + BLOCKS[7:]})
    assert r.status_code == 200 and "**표 1.** 결과 요약" in r.text


# =================================================================== QA 수정: JPEG 변형 · 같은 그림 반복 · 하위 목록 · 사용량
def test_validate_rejects_jpeg_word_cannot_read():
    assert figures.validate(jpeg_header(64, 32))["ext"] == "jpg"  # JFIF 머리글이면 받는다
    assert figures.word_readable(jpeg_header(64, 32)) and figures.word_readable(IMG) and figures.word_readable(GIF)
    for data in ODD_JPEGS:
        assert figures.sniff(data) == "jpg" and figures.image_size(data, "jpg") == (64, 32)
        assert not figures.word_readable(data)
        with pytest.raises(figures.FigureError, match="워드에서 읽을 수 없는"):
            figures.validate(data)


@pytest.mark.parametrize("fmt", [None, THESIS], ids=["default", "thesis"])
@pytest.mark.parametrize("data", ODD_JPEGS, ids=["xmp", "app14", "dqt"])
def test_docx_export_survives_unreadable_jpeg(fmt, data):
    """이미 R2에 있는(검사 전에 올린) 이상한 JPEG도 워드 내보내기 전체를 깨뜨리지 않는다 → 자리 글 + 경고"""
    blocks = [{"type": "paragraph", "runs": [{"text": "본문"}]},
              {"type": "figure", "src": NAME, "caption": [{"text": "그림 1.", "b": True}],
               "image": {"data": data, "format": "jpg", "width": 64, "height": 32}}]
    warnings = []
    out = writer.to_docx(blocks, {"title": "t"}, fmt, COVER if fmt else None, warnings)
    d = Document(io.BytesIO(out))
    texts = [p.text for p in d.paragraphs]
    assert writer.FIGURE_MISSING in texts and "그림 1." in texts and "본문" in texts
    assert writer.FIGURE_UNREADABLE_WARNING in warnings
    with zipfile.ZipFile(io.BytesIO(out)) as z:
        assert not [n for n in z.namelist() if n.startswith("word/media/")]
        assert "<w:drawing" not in z.read("word/document.xml").decode()
    # 한글은 그대로 넣는다
    h = HwpxDocument.open(io.BytesIO(writer.to_hwpx(blocks, {"title": "t"}, fmt, COVER if fmt else None)))
    assert h.validate().issues == ()


def test_attach_images_loads_each_name_once_and_caps_total(monkeypatch):
    calls = []

    def load(name):
        calls.append(name)
        return IMG

    blocks = [{"type": "figure", "src": NAME} for _ in range(figures.MAX_EXPORT_FIGURES)]
    out = figures.attach_images(blocks, load)
    assert calls == [NAME]
    assert all(b["image"] is out[0]["image"] for b in out)
    # 읽은 합계가 원고 상한을 넘으면 나머지는 그림 없이(자리 글)
    monkeypatch.setattr(figures, "MAX_FIGURE_BYTES_PER_MANUSCRIPT", len(IMG) * 2)
    names = [f"{i:016x}.png" for i in range(4)]
    calls.clear()
    out = figures.attach_images([{"type": "figure", "src": n} for n in names], load)
    assert ["image" in b for b in out] == [True, True, False, False]
    assert calls == names[:2]


def test_repeated_figure_is_stored_once_in_exports():
    image = {"data": IMG, "format": "png", "width": 400, "height": 200}
    blocks = [{"type": "figure", "src": NAME, "caption": [], "image": image} for _ in range(3)]
    for data, prefix in ((writer.to_hwpx(blocks), "BinData/"), (writer.to_hwpx(blocks, {}, THESIS, COVER), "BinData/"),
                         (writer.to_docx(blocks), "word/media/")):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            assert len([n for n in z.namelist() if n.startswith(prefix)]) == 1


def test_markdown_nested_lists_follow_parent_marker_width():
    """하위 항목은 부모 기호 폭만큼 들여써야 CommonMark · marked에서 하위 목록 ("10. x" 아래는 4칸)"""
    def li(depth, n, ordered=True, t="x"):
        return {"type": "list_item", "ordered": ordered, "depth": depth, "number": n, "runs": [{"text": t}]}

    md = writer.to_markdown([li(0, 10, t="열"), li(1, 1, False, "하위"), li(2, 1, t="손자"), li(0, 11, t="열하나"),
                             {"type": "paragraph", "runs": [{"text": "문단"}]}, li(1, 1, False, "부모 없음")])
    assert "10. 열\n\n    - 하위\n\n      1. 손자\n\n11. 열하나" in md
    assert "\n\n- 부모 없음" in md  # 부모 없는 하위 항목은 맨 바깥으로 (코드 블록이 되지 않게)


def test_figure_size_cache_counts_only_figure_keys():
    fake = st.FakeStorage()
    uid = str(uuid.uuid4())
    fake.put(st.figure_key(uid, 3, NAME), IMG, "image/png")
    fake.put(f"users/{uid}/papers/1.pdf", b"x" * 500)
    fake.put("backups/db/x.dump", b"y" * 70)
    cache = st.FigureSizeCache(fake, ttl=300)
    assert cache.get() == len(IMG)
    cache.add(100)  # 올린 직후 캐시에 바로 더함
    assert cache.get() == len(IMG) + 100
    assert st.BackupSizeCache(fake).get() == 70


# =================================================================== API (DB)
@pytest.mark.db
def test_figure_upload_refused_when_storage_full_and_counted_in_usage(cloud):
    """R2 95% 규칙: 사용량이 'full'이면 그림도 못 올린다. 그림 바이트는 사용량에 더해진다"""
    c = cloud.client(cloud.user())
    mid = c.post("/api/manuscripts", json={"content": "# 원고"}).json()["id"]
    cloud.app.state.figure_sizes.ttl = -1
    before = c.get("/api/storage/usage").json()["used_bytes"]
    assert c.post(f"/api/manuscripts/{mid}/figures", files={"file": ("a.png", IMG, "image/png")}).status_code == 200
    assert c.get("/api/storage/usage").json()["used_bytes"] >= before + len(IMG)
    limit = cloud.app.state.config.storage_limit_bytes
    cloud.app.state.config.storage_limit_bytes = 1
    try:
        r = c.post(f"/api/manuscripts/{mid}/figures", files={"file": ("a.png", IMG, "image/png")})
        assert r.status_code == 400 and "저장 공간이 거의 찼어요" in r.json()["detail"]
    finally:
        cloud.app.state.config.storage_limit_bytes = limit
    # 워드가 못 읽는 JPEG는 올릴 때 막는다 (화면은 그 전에 JPEG로 다시 저장해 보냄)
    bad = c.post(f"/api/manuscripts/{mid}/figures", files={"file": ("x.jpg", XMP_FIRST, "image/jpeg")})
    assert bad.status_code == 400 and "워드에서 읽을 수 없는" in bad.json()["detail"]


@pytest.mark.db
def test_figure_upload_get_export_and_cleanup(cloud):
    user = cloud.user()
    c = cloud.client(user)
    mid = c.post("/api/manuscripts", json={"content": "# 그림 원고\n\n본문"}).json()["id"]
    r = c.post(f"/api/manuscripts/{mid}/figures", files={"file": ("shot.webp", IMG, "image/webp")})
    assert r.status_code == 200, r.text
    fig = r.json()
    assert figures.NAME_RE.fullmatch(fig["name"]) and fig["name"].endswith(".png") and fig["ref"] == f"fig:{fig['name']}"
    assert (fig["width"], fig["height"]) == (400, 200)
    key = f"users/{user.id}/figures/{mid}/{fig['name']}"
    assert cloud.storage.objects[key] == IMG

    got = c.get(f"/api/manuscripts/{mid}/figures/{fig['name']}")
    assert got.status_code == 200 and got.content == IMG and got.headers["content-type"] == "image/png"
    assert got.headers["x-content-type-options"] == "nosniff"
    assert c.get(f"/api/manuscripts/{mid}/figures/0123456789abcdee.png").status_code == 404
    assert c.get(f"/api/manuscripts/{mid}/figures/..%2Fx.png").status_code == 404
    # 형식 · 크기 검사 (매직 바이트)
    bad = c.post(f"/api/manuscripts/{mid}/figures", files={"file": ("a.png", b"<svg/>", "image/png")})
    assert bad.status_code == 400 and "PNG" in bad.json()["detail"]
    # 남의 원고 · 없는 원고
    other = cloud.client(cloud.user())
    assert other.get(f"/api/manuscripts/{mid}/figures/{fig['name']}").status_code == 404
    assert other.post(f"/api/manuscripts/{mid}/figures", files={"file": ("a.png", IMG, "image/png")}).status_code == 404

    blocks = [{"type": "figure", "src": fig["name"], "caption": [{"text": "그림 1.", "b": True}]}]
    out = c.post("/api/export-document", json={"format": "docx", "blocks": blocks, "manuscript_id": mid})
    assert out.status_code == 200
    with zipfile.ZipFile(io.BytesIO(out.content)) as z:
        assert [z.read(n) for n in z.namelist() if n.startswith("word/media/")] == [IMG]
    out = c.post("/api/export-document", json={"format": "hwpx", "blocks": blocks, "manuscript_id": mid})
    with zipfile.ZipFile(io.BytesIO(out.content)) as z:
        assert [z.read(n) for n in z.namelist() if n.startswith("BinData/")] == [IMG]
    # 남의 원고 번호로는 내 그림을 못 읽는다
    out = other.post("/api/export-document", json={"format": "docx", "blocks": blocks, "manuscript_id": mid})
    assert writer.FIGURE_MISSING in [p.text for p in Document(io.BytesIO(out.content)).paragraphs]

    assert c.delete(f"/api/manuscripts/{mid}").status_code == 200
    assert key not in cloud.storage.objects
