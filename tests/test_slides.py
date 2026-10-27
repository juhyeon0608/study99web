"""발표 자료 (2026-10-27): AI 슬라이드 JSON → pptx (제목 · 글머리 · 발표자 노트 · 예상 질문) · API"""
import io

import pytest

from paperlab import slides

DATA = ('```json\n{"title": "청소년 우울", "slides": [{"title": "연구 배경", "bullets": ["우울 증가", ""], "notes": "설명"},'
        ' {"title": "", "bullets": []}], "questions": [{"q": "표본은?", "a": "검정력"}, {"q": "한계는?"}, {"q": ""}]}\n```')


def test_parse_and_build_pptx():
    from pptx import Presentation
    d = slides.parse(DATA)
    assert d["slides"] == [{"title": "연구 배경", "bullets": ["우울 증가"], "notes": "설명"}]
    assert [q["q"] for q in d["questions"]] == ["표본은?", "한계는?"]
    prs = Presentation(io.BytesIO(slides.to_pptx(d)))
    titles = [s.shapes.title.text for s in prs.slides]
    assert titles == ["청소년 우울", "연구 배경", "예상 질문"]
    assert prs.slides[1].notes_slide.notes_text_frame.text == "설명"
    xml = prs.slides[1].shapes.placeholders[1]._element.xml
    assert 'typeface="맑은 고딕"' in xml and "<a:ea " in xml
    for bad in ("JSON 아님", {"slides": []}, None):
        with pytest.raises(slides.SlidesError):
            slides.parse(bad)


def test_slides_api(local_client):
    r = local_client.post("/api/slides", json={"data": DATA, "filename": "발표"})
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert "presentationml" in r.headers["content-type"]
    assert local_client.post("/api/slides", json={"data": "없음"}).status_code == 400
