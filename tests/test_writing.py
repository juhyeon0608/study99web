"""원고 내보내기(워드·한글)와 워드·한글 문서의 인용 표시 치환."""
import io
import zipfile

import pytest
from docx import Document
from hwpx import HwpxDocument

from paperlab import compose, writer

from .conftest import SAMPLE

BLOCKS = [
    {"type": "title", "runs": [{"text": "그래프 신경망 연구"}]},
    {"type": "heading", "level": 1, "runs": [{"text": "1. 서론"}]},
    {"type": "paragraph", "runs": [{"text": "어텐션만 쓴다 "}, {"text": "(Vaswani et al., 2017)"}, {"text": ". "},
                                   {"text": "굵게", "b": True}, {"text": " x"}, {"text": "2", "sup": True},
                                   {"footnote": [{"text": "각주 "}, {"text": "기울임", "i": True}]}, {"text": " 끝."}]},
    {"type": "list_item", "ordered": True, "number": 1, "depth": 0, "runs": [{"text": "항목"}]},
    {"type": "bib_heading", "text": "참고문헌"},
    {"type": "bib_entry", "hanging": True, "runs": [{"text": "Vaswani, A. (2017). Attention. "},
                                                    {"text": "NeurIPS", "i": True}, {"text": "."}]},
]


def test_docx_export_with_footnote():
    data = writer.to_docx(BLOCKS, {"title": "t"})
    d = Document(io.BytesIO(data))
    texts = [p.text for p in d.paragraphs]
    assert texts[0] == "그래프 신경망 연구" and "1. 서론" in texts and "1. 항목" in texts
    body = next(p for p in d.paragraphs if p.text.startswith("어텐션"))
    assert any(r.bold for r in body.runs if r.text == "굵게")
    assert any(r.font.superscript for r in body.runs if r.text == "2")
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        notes = z.read("word/footnotes.xml").decode()
        assert "각주 " in notes and "<w:i/>" in notes
        assert 'w:footnoteReference w:id="1"' in z.read("word/document.xml").decode()
        assert "footnotes+xml" in z.read("[Content_Types].xml").decode()
    bib = next(p for p in d.paragraphs if p.text.startswith("Vaswani"))
    assert bib.paragraph_format.first_line_indent < 0


def test_hwpx_export_is_valid():
    data = writer.to_hwpx(BLOCKS)
    d = HwpxDocument.open(io.BytesIO(data))
    assert d.validate().issues == ()
    text = d.text.plain()
    assert "그래프 신경망 연구" in text and "Vaswani, A. (2017)" in text
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        section = z.read("Contents/section0.xml").decode()
    assert "footNote" in section and "각주 기울임" in section


def test_markdown_export():
    md = writer.to_markdown(BLOCKS)
    assert md.startswith("# 그래프 신경망 연구") and "[^1]" in md and "*NeurIPS*" in md


def test_parse_citation_syntax():
    assert compose.parse_citation("@a") == [{"key": "a", "suppress_author": False}]
    items = compose.parse_citation("see @a, p. 12; -@b, 34쪽")
    assert items[0] == {"key": "a", "suppress_author": False, "prefix": "see", "locator": "12", "label": "page"}
    assert items[1]["key"] == "b" and items[1]["suppress_author"] and items[1]["locator"] == "34"
    assert compose.parse_citation("이메일 a@b.com") is not None  # 키 형태면 인용으로 본다
    cites = compose.find_citations("문장 [@a; @b]. 그리고 [참고]와 [@c, pp. 3-5].")
    assert [c.raw for c in cites] == ["[@a; @b]", "[@c, pp. 3-5]"]
    assert cites[1].items[0]["locator"] == "3-5"
    # 품질팀 F2: 앞에 \가 붙은 [는 인용이 아님 — 참고 패널이 이스케이프해 넣은 PDF 글 속 [@…]로 키를 주입하지 못함
    # (refquote.js의 CITE_RE와 같은 규칙 — tests/js/refquote.test.mjs)
    text = "“as in \\[@evil, p. 9\\] and \\[@evil2\\]” [@k, p. 3] 그리고 [@a]"
    assert [c.raw for c in compose.find_citations(text)] == ["[@k, p. 3]", "[@a]"]


def _docx_with_markers() -> bytes:
    d = Document()
    p = d.add_paragraph("트랜스포머는 ")
    r = p.add_run("[@vaswani")  # 표시가 서식이 다른 두 런에 걸쳐 있다
    r.bold = True
    p.add_run("2017attention] 를 제안했다.")
    d.add_paragraph("두 번째 [@vaswani2017attention, p. 3] 인용과 [@없는키].")
    d.add_paragraph("[참고문헌]")
    d.add_paragraph("부록")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_docx_compose_author_date():
    data = _docx_with_markers()
    scan = compose.docx_scan(data)
    assert [c.raw for c in scan["citations"]] == ["[@vaswani2017attention]", "[@vaswani2017attention, p. 3]", "[@없는키]"]
    assert scan["has_bib_marker"]
    rendered = [{"runs": [{"text": "(Vaswani et al., 2017)"}]}, {"runs": [{"text": "(Vaswani et al., 2017, p. 3)"}]},
                {"runs": [{"text": "[@없는키?]"}]}]
    bib = [[{"text": "Vaswani, A. (2017). "}, {"text": "Attention", "i": True}, {"text": "."}]]
    out = Document(io.BytesIO(compose.docx_apply(data, rendered, bib, "참고문헌", note_style=False)))
    texts = [p.text for p in out.paragraphs]
    assert texts[0] == "트랜스포머는 (Vaswani et al., 2017) 를 제안했다."
    assert texts[1] == "두 번째 (Vaswani et al., 2017, p. 3) 인용과 [@없는키?]."
    assert texts[2:] == ["참고문헌", "Vaswani, A. (2017). Attention.", "부록"]
    assert any(r.bold for r in out.paragraphs[0].runs if "Vaswani" in r.text)  # 원래 굵은 서식 유지
    assert any(r.italic for r in out.paragraphs[3].runs if r.text == "Attention")


def test_docx_compose_footnote_style():
    data = _docx_with_markers()
    rendered = [{"runs": [{"text": "Ashish Vaswani, “Attention,” 2017."}]}, {"runs": [{"text": "Ibid., 3."}]},
                {"runs": [{"text": "?"}]}]
    raw = compose.docx_apply(data, rendered, [], "참고문헌", note_style=True)
    out = Document(io.BytesIO(raw))
    assert out.paragraphs[0].text == "트랜스포머는  를 제안했다."
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        doc_xml = z.read("word/document.xml").decode()
        notes = z.read("word/footnotes.xml").decode()
    assert doc_xml.count("w:footnoteReference") == 3 and "Ibid., 3." in notes


def test_hwpx_compose():
    d = HwpxDocument.new()
    d.add_paragraph("트랜스포머는 [@vaswani2017attention] 를 제안했다.")
    d.add_paragraph("[참고문헌]")
    d.add_paragraph("부록")
    data = d.to_bytes()
    scan = compose.hwpx_scan(data)
    assert [c.raw for c in scan["citations"]] == ["[@vaswani2017attention]"] and scan["has_bib_marker"]
    out = compose.hwpx_apply(data, [{"runs": [{"text": "(Vaswani et al., 2017)"}]}],
                             [[{"text": "Vaswani, A. (2017). "}, {"text": "Attention", "i": True}]], "참고문헌")
    doc = HwpxDocument.open(io.BytesIO(out))
    assert doc.validate().issues == ()
    lines = [ln for ln in doc.text.plain().splitlines() if ln.strip()]
    assert lines == ["트랜스포머는 (Vaswani et al., 2017) 를 제안했다.", "참고문헌", "Vaswani, A. (2017). Attention", "부록"]


def test_detect_kind_rejects_legacy():
    with pytest.raises(ValueError, match=".hwpx"):
        compose.detect_kind("a.hwp", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"0" * 100)


@pytest.mark.db
def test_manuscript_api_and_compose_flow(client):
    tpls = client.get("/api/manuscript-templates").json()
    assert {"blank", "kr_journal", "thesis", "imrad"} <= {t["id"] for t in tpls}
    m = client.post("/api/manuscripts", json={"template": "kr_journal"}).json()
    assert m["title"] == "논문 제목" and "[참고문헌]" in m["content"]
    client.patch(f"/api/manuscripts/{m['id']}", json={"content": "# 새 제목\n\n본문"})
    assert client.get("/api/manuscripts").json()[0]["title"] == "새 제목"

    pid = client.post("/api/papers", json=SAMPLE).json()["paper"]["id"]
    key = client.get(f"/api/papers/{pid}").json()["citekey"]
    ko = client.post("/api/papers", json={"title": "그래프 연구", "year": 2023,
                                          "authors": [{"family": "홍", "given": "길동"}]}).json()["paper"]
    assert ko["citekey"] == "hong2023"
    look = client.post("/api/citekeys", json={"keys": [key, "nope"]}).json()
    assert look["items"][key]["title"] == SAMPLE["title"] and look["items"]["nope"] is None

    r = client.post("/api/export-document", json={"format": "hwpx", "blocks": BLOCKS, "filename": "내 원고"})
    assert r.status_code == 200 and "UTF-8''%EB%82%B4%20%EC%9B%90%EA%B3%A0.hwpx" in r.headers["content-disposition"]
    assert client.post("/api/export-document", json={"format": "pdf", "blocks": []}).status_code == 400

    scan = client.post("/api/compose/scan", files={"file": ("draft.docx", _docx_with_markers(), "application/octet-stream")}).json()
    assert scan["kind"] == "docx" and scan["items"]["없는키"] is None and len(scan["citations"]) == 3
    out = client.post("/api/compose/apply", json={"token": scan["token"], "rendered": [{"runs": [{"text": "(V)"}]}] * 3,
                                                  "bibliography": [[{"text": "Ref"}]], "bib_title": "References"})
    assert out.status_code == 200 and "_%EC%9D%B8%EC%9A%A9%EC%99%84%EB%A3%8C.docx" in out.headers["content-disposition"]
    assert Document(io.BytesIO(out.content)).paragraphs[2].text == "References"
    bad = client.post("/api/compose/scan", files={"file": ("x.hwp", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"0" * 64, "x")})
    assert bad.status_code == 400


def test_docx_compose_into_document_with_existing_footnotes():
    blocks = [{"type": "paragraph", "runs": [{"text": "기존 각주"}, {"footnote": [{"text": "원래 각주"}]},
                                             {"text": " 다음 [@k] 인용."}]}]
    data = writer.to_docx(blocks)
    raw = compose.docx_apply(data, [{"runs": [{"text": "새 각주"}]}], [], "참고문헌", note_style=True)
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        notes = z.read("word/footnotes.xml").decode()
        doc_xml = z.read("word/document.xml").decode()
    assert "원래 각주" in notes and "새 각주" in notes and 'w:id="2"' in notes
    assert doc_xml.count("w:footnoteReference") == 2


def test_docx_compose_footnotes_follow_document_order():
    d = Document()
    d.add_paragraph("가 [@a] 나 [@b] 다 [@c].")
    buf = io.BytesIO()
    d.save(buf)
    raw = compose.docx_apply(buf.getvalue(), [{"runs": [{"text": f"주{i}"}]} for i in (1, 2, 3)], [], "", True)
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        doc_xml = z.read("word/document.xml").decode()
        notes = z.read("word/footnotes.xml").decode()
    import re
    order = re.findall(r'footnoteReference w:id="(\d+)"', doc_xml)
    contents = {i: re.search(rf'w:id="{i}">.*?주(\d)', notes, re.S).group(1) for i in order}
    assert [contents[i] for i in order] == ["1", "2", "3"]


@pytest.mark.db
def test_ai_write_endpoint_passes_sources(project, session_db, users):
    from .conftest import Cloud
    import json as _json

    class FakeAI:
        def status(self):
            return {"ready": True, "message": ""}

        def write(self, mode, text, instruction="", context="", sources=None):
            self.seen = (mode, text, sources)
            yield {"type": "delta", "text": "초안"}
            yield {"type": "done", "text": f"초안 [@{sources[0]['key']}]"}

    fake = FakeAI()
    cloud = Cloud(project, session_db, users, ai=fake)
    c = cloud.client(cloud.user())
    pid = c.post("/api/papers", json=SAMPLE).json()["paper"]["id"]
    key = c.get(f"/api/papers/{pid}").json()["citekey"]
    c.put(f"/api/papers/{pid}/note", json={"content": "중요한 논문"})
    with c.stream("POST", "/api/ai/write", json={"mode": "draft", "text": "", "keys": [key]}) as r:
        evs = [_json.loads(line[6:]) for line in r.iter_lines() if line.startswith("data: ")]
    assert evs[-1]["text"] == f"초안 [@{key}]"
    assert fake.seen[2][0]["note"] == "중요한 논문" and fake.seen[2][0]["title"] == SAMPLE["title"]
    assert c.post("/api/ai/write", json={"mode": "polish", "text": ""}).status_code == 400
